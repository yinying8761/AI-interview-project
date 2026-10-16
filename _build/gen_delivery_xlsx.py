# -*- coding: utf-8 -*-
"""生成交付包的三份 xlsx：成员任务表 / 选题登记表 / 测试结果。

内容要求取自《考核实施方案（人工智能Ⅱ方向）》与《学生提交材料目录说明》：
  - 成员任务表：成员姓名、学号、分工、成果位置；任务记录能对应到个人成果
  - 选题登记表：2 个感知模块、1 个扩展能力及接入评分报告的说明；开题与结课一致
  - 测试结果  ：正常、异常和权限相关测试用例及结果；结果真实，可定位到测试数据
实测数据来源：evidence/构建或自动化测试记录.txt 与 evidence/关键接口返回.json
"""
import json
import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

OUT = sys.argv[1]
os.makedirs(OUT, exist_ok=True)

HEAD_FILL = PatternFill("solid", fgColor="DCE6F1")
CAT_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="9E9E9E")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")

XUEHAO = "（请填写学号）"


def write_sheet(ws, headers, rows, widths, title_rows=None):
    """title_rows: 写在表头之前的整行说明，元素为 (文本, 合并到第几列)"""
    r = 1
    if title_rows:
        for text, span in title_rows:
            ws.cell(r, 1, text)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=span)
            c = ws.cell(r, 1)
            c.font = Font(bold=True, size=12)
            c.alignment = Alignment(vertical="center", horizontal="left")
            ws.row_dimensions[r].height = 22
            r += 1
    head_row = r
    for i, h in enumerate(headers, start=1):
        c = ws.cell(head_row, i, h)
        c.font = Font(bold=True, size=10)
        c.fill = HEAD_FILL
        c.alignment = CENTER
        c.border = BORDER
    for row in rows:
        r += 1
        for i, v in enumerate(row, start=1):
            c = ws.cell(r, i, v)
            c.font = Font(size=10)
            c.alignment = WRAP
            c.border = BORDER
        ws.row_dimensions[r].height = None
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(head_row + 1, 1)
    return head_row


# ============================================================================
# 一、成员任务表.xlsx
# ============================================================================
MEMBERS = [
    ("1", "郑颖浩", XUEHAO, "组长", "项目管理、架构设计、开发、部署",
     "主持需求评审并划定 V1.0 范围；设计六模块分层架构与数据库模型；"
     "开发公共基础模块与用户权限模块；统一开发环境与部署方式；整理提交包并做安全核验",
     "个人材料：作业01 需求与方案分析、作业03 SpringBoot接口数据权限、个人报告、工作证据",
     "src/interview-ai/pom.xml；interview-common/；interview-modules/interview-system/；"
     "sql/interview_ai.sql；config/；README.md；evidence/"),
    ("2", "常辰", XUEHAO, "组员", "AI 能力开发、感知模块开发",
     "对接大模型完成文本理解与评价；用 MediaPipe 实现视觉行为分析；"
     "实现问答×80%＋视频×20% 的融合评分与报告自动生成；维护 AI 服务接口契约",
     "个人材料：作业01 需求与方案分析、作业03 SpringBoot接口数据权限、个人报告、工作证据",
     "src/smartspeech_optimized/backend/app/api/ai.py；services/score_service.py；services/vision/；"
     "src/interview-ai/.../interview/service/ReportGenerator.java"),
    ("3", "赵佳伟", XUEHAO, "组员", "前端开发、前后端联调、测试",
     "开发用户端与管理端页面；完成面试间、AI 会话与报告的联调；"
     "编写单元测试与现场冒烟用例并输出测试结果表",
     "个人材料：作业02 Vue3页面与交互、作业04 测试及项目风险分析、个人报告、工作证据",
     "src/interview-ai/frontend/src/views/、router/、api/、composables/；"
     "src/interviwe-platform/interview-platform/src/；tools/；测试结果.xlsx"),
]

# 成果位置：说明本组成员在 src/ 下实际提交的产物
ARTIFACTS = [
    ("郑颖浩", "聚合工程与构建", "src/interview-ai/pom.xml、src/interview-ai/mvnw"),
    ("郑颖浩", "公共基础模块（统一响应、异常处理、JWT、AI 调用客户端）", "src/interview-ai/interview-common/"),
    ("郑颖浩", "用户与权限模块（登录注册、三角色 RBAC、菜单权限）", "src/interview-ai/interview-modules/interview-system/"),
    ("郑颖浩", "数据库建模与建库脚本（27 张表）", "src/interview-ai/sql/interview_ai.sql"),
    ("郑颖浩", "启动配置与运行配置模板", "src/interview-ai/interview-bootstrap/src/main/resources/、config/"),
    ("郑颖浩", "交付包整理、安全核验与运行说明", "README.md、evidence/"),
    ("常辰", "大模型面试官与文本理解评价接口", "src/smartspeech_optimized/backend/app/api/ai.py"),
    ("常辰", "评分融合（问答×80%＋视频×20%）与边界处理", "src/smartspeech_optimized/backend/app/services/score_service.py"),
    ("常辰", "视觉行为分析（人脸关键点、表情、注意力指标）", "src/smartspeech_optimized/backend/app/services/vision/"),
    ("常辰", "面试报告生成与可追溯报告落库", "src/interview-ai/interview-modules/interview-interview/src/main/java/com/iflytek/interview/interview/service/ReportGenerator.java"),
    ("常辰", "AI 服务接口与参数说明", "src/smartspeech_optimized/智能模拟面试系统接口清单及接口文档.docx"),
    ("赵佳伟", "用户端页面（登录、预约、面试间、AI 会话、我的记录与报告）", "src/interview-ai/frontend/src/views/"),
    ("赵佳伟", "用户端路由、请求封装与状态管理", "src/interview-ai/frontend/src/router/、src/interview-ai/frontend/src/api/、src/interview-ai/frontend/src/composables/"),
    ("赵佳伟", "管理端页面（用户、题库、试卷、标签、统计看板）", "src/interviwe-platform/interview-platform/src/"),
    ("赵佳伟", "单元测试与现场验收冒烟用例", "src/interview-ai/interview-modules/interview-interview/src/test/、src/interview-ai/interview-common/src/test/、tools/"),
    ("赵佳伟", "测试执行结果", "测试结果.xlsx、evidence/关键接口返回.json"),
]

wb = Workbook()
ws = wb.active
ws.title = "成员任务表"
write_sheet(
    ws,
    ["序号", "成员姓名", "学号", "成员类别", "分工（项目角色）",
     "分工任务说明", "对应个人成果", "成果位置"],
    [],
    [6, 10, 14, 10, 24, 48, 40, 56],
    title_rows=[("多模态智能模拟面试评测平台（软件2324）—— 成员任务表", 8)],
)

# 主表：成员与分工
r = 3
for m in MEMBERS:
    for i, v in enumerate(m, start=1):
        c = ws.cell(r, i, v)
        c.font = Font(size=10)
        c.alignment = WRAP
        c.border = BORDER
    ws.row_dimensions[r].height = 72
    r += 1
r += 1
ws.cell(r, 1, "成果位置明细（对应 src/ 下的实际提交内容）").font = Font(bold=True, size=11)
r += 1
head2 = r
for i, h in enumerate(["成员姓名", "成果内容", "成果位置"], start=1):
    c = ws.cell(head2, i, h)
    c.font = Font(bold=True, size=10)
    c.fill = HEAD_FILL
    c.alignment = CENTER
    c.border = BORDER
for a in ARTIFACTS:
    r += 1
    for i, v in enumerate(a, start=1):
        c = ws.cell(r, i, v)
        c.font = Font(size=10)
        c.alignment = WRAP
        c.border = BORDER
    ws.row_dimensions[r].height = 30

# ============================================================================
# 二、选题登记表.xlsx
# ============================================================================
TOPIC = [
    ("感知模块一", "文本理解与评价", "常辰",
     "登录候选人 → 面试大厅选择场景与「大模型面试官（文本＋视觉）」→ 进入面试间 → 生成首题 → 文本作答并提交",
     "data/测试数据/新文本回答.txt（新的岗位回答文本）",
     "调用 DeepSeek 对真实题面与考生作答逐题评估，返回真实问答得分与逐题点评",
     "问答得分按 80% 权重计入综合分",
     "已接入。融合公式 final_score = question_score × 0.8 + (expression_score × 0.5 + attention_score × 0.5) × 0.2"),
    ("感知模块二", "视觉行为分析", "常辰",
     "面试间开启摄像头并保持面部可见约 15 秒以上 → 完成面试 → 我的记录 → 查看报告",
     "data/测试数据/合成面试人像.png（逐帧输入的新图像，虚构人物）",
     "MediaPipe FaceMesh 采集面部关键点，输出表情分与注意力分，视频维度按 20% 权重参与融合",
     "视觉得分 = 表情 × 50% + 注意力 × 50%，再按 20% 权重计入综合分",
     "已接入。未采集到有效视觉样本时报告显示「未评估」，不计入综合分且不重新分配权重"),
    ("扩展能力", "大模型面试官", "常辰",
     "面试间「生成首题」按钮；连续作答触发大模型追问",
     "任意一场新建面试",
     "大模型按岗位与简历生成开场题并根据作答追问，会话状态可通过接口查询",
     "本扩展能力负责产出问答内容，其评分由「文本理解与评价」模块完成",
     "已接入。开场题与追问由 DeepSeek 实时生成，非固定题库"),
    ("评分报告", "两项感知结果接入评分引擎", "郑颖浩、常辰",
     "我的记录 → 查看报告",
     "完成一场同时含问答与视觉样本的面试",
     "综合分由问答与视觉两项感知结果按权重计算，报告保存输入原文、逐题评估、公式与各维度状态",
     "综合分 = 问答 × 0.8 + (表情 × 0.5 + 注意力 × 0.5) × 0.2",
     "已接入。报告写入 interview_record.report_json，可追溯输入原文、逐题评估与评分规则；"
     "未测维度显示「未评估」，缺少任一维度时不生成综合分"),
]

wb2 = Workbook()
ws2 = wb2.active
ws2.title = "选题登记表"
write_sheet(
    ws2,
    ["类型", "功能名称", "负责人", "操作入口", "测试输入", "预期结果",
     "接入评分报告说明", "评分规则与接入状态"],
    [list(t) for t in TOPIC],
    [12, 16, 12, 40, 30, 40, 40, 46],
    title_rows=[
        ("多模态智能模拟面试评测平台（软件2324）—— 选题登记表", 8),
        ("选题范围：感知模块从「语音分析、视觉行为分析、文本理解与评价」中任选 2 项；"
         "扩展能力从「简历与题目、大模型面试官、实时面试间」中任选 1 项。", 8),
        ("本组选定：感知模块一 文本理解与评价 ＋ 感知模块二 视觉行为分析 ＋ 扩展能力 大模型面试官。"
         "两项感知结果均接入评分引擎并生成可追溯输入数据的评测报告。", 8),
    ],
)

# ============================================================================
# 三、测试结果.xlsx
# ============================================================================
EVID = "evidence/关键接口返回.json"
BUILD = "evidence/构建或自动化测试记录.txt"

CASES = [
    ("正常", "候选人登录", "test01 / 123456", "POST /api/users/login",
     "返回 200 与有效 JWT，角色为候选人", "code=200，返回 token（203 字符）", "通过", f"{EVID} → 1_登录预约查询越权_smoke_test → 登录 test01"),
    ("正常", "管理员与面试官登录", "admin / admin123，host01 / 123456", "POST /api/users/login",
     "返回 200 与对应角色", "code=200，两账号均登录成功", "通过", f"{EVID} → 登录 admin、登录 host01"),
    ("正常", "读取真实面试场景", "候选人已登录", "GET /api/scenarios",
     "返回数据库中真实存在的场景列表", "code=200，读取到真实场景", "通过", f"{EVID} → 读取真实场景"),
    ("正常", "候选人预约面试并落库", "test01，指定场景与预约时间", "POST /api/interview-sessions",
     "返回 200，预约记录写入数据库", "code=200，记录成功落库", "通过", f"{EVID} → 候选人预约落库"),
    ("正常", "刷新查询与换角色查询", "test01、admin、host01 三种角色", "GET /api/records/my",
     "候选人可查本人记录，HR 与管理员可查记录", "三种角色查询均返回 200 且数据一致", "通过", f"{EVID} → 刷新及换角色查询 test01 / admin / host01"),
    ("正常", "感知模块一：文本理解与评价", "data/测试数据/新文本回答.txt（新的岗位回答）", "生成首题 → 提交文本回答 → 完成面试",
     "调用 DeepSeek 返回真实问答得分与逐题点评", "问答得分 45.0，textStatus=已评估，报告逐题评估与输入原文一致", "通过", f"{EVID} → 2_大模型出题与问答评分_smoke_ai；{BUILD} → [2]"),
    ("正常", "扩展能力：大模型面试官", "任意一场新建面试", "POST /api/interview-sessions/{id}/conversation/opening",
     "大模型按岗位生成开场题，非固定题库", "返回真实生成的首题，内容随场景与简历变化", "通过", f"{EVID} → 2_大模型出题与问答评分_smoke_ai → question"),
    ("正常", "感知模块二：视觉行为分析", "data/测试数据/合成面试人像.png（逐帧输入，共 50 帧）", "开启视觉采集 → 逐帧提交 → 完成面试",
     "MediaPipe 采集到人脸并输出表情分与注意力分", "人脸采集 50 帧，视觉得分 65.0", "通过", f"{EVID} → 3_视觉分析与融合报告_smoke_vision；{BUILD} → [3]"),
    ("正常", "融合评分与报告生成", "同一场含问答与视觉样本的面试", "GET /api/records/my → 查看报告",
     "综合分按公式计算，报告可追溯输入与评分规则",
     "综合分 49.0 ＝ 45.0×0.8 ＋ 65.0×0.2，report_json 含输入原文、逐题评估与公式", "通过", f"{EVID} → 3_视觉分析与融合报告_smoke_vision → finalScore"),

    ("异常", "登录密码为空", "username=test01，password 为空", "POST /api/users/login",
     "参数校验拦截并返回明确错误", "被参数校验拦截，未签发令牌", "通过", f"{EVID} → 空密码参数校验"),
    ("异常", "使用不存在的面试场景预约", "scenarioId=999999", "POST /api/interview-sessions",
     "返回业务异常「面试场景不存在」", "拦截并返回场景不存在", "通过", f"{EVID} → 不存在场景"),
    ("异常", "重复时间段的预约冲突", "与已有预约相同的预约时间", "POST /api/interview-sessions",
     "返回预约时间冲突", "拦截并返回时间冲突", "通过", f"{EVID} → 重复时间预约冲突"),
    ("异常", "已取消的面试不允许启动", "已取消状态的 recordId", "POST /api/interview-sessions/{id}/start",
     "状态机拦截非法状态转换", "拦截并返回非法状态转换", "通过", f"{EVID} → 已取消记录不能启动"),
    ("异常", "未测维度处理", "不开启摄像头完成一场面试", "完成面试 → 查看报告",
     "视觉维度显示「未评估」，不生成综合分，权重不重新分配",
     "visionStatus=未评估，finalScore 为空，问答分仍如实保留", "通过", f"{EVID} → 2_大模型出题与问答评分_smoke_ai → 追溯报告"),

    ("权限", "未登录访问受保护接口", "不携带令牌", "GET /api/records/my",
     "返回 401 拒绝访问", "拦截并返回未认证", "通过", f"{EVID} → 未登录访问记录"),
    ("权限", "查看他人面试记录详情", "test01 携带令牌访问他人 recordId", "GET /api/records/{id}",
     "返回 403 仅本人可访问", "拦截并返回无权访问", "通过", f"{EVID} → 他人记录详情权限"),
    ("权限", "查看他人面试状态快照", "test01 携带令牌访问他人 recordId", "GET /api/interview-sessions/{id}/state",
     "返回 403 仅本人可访问", "拦截并返回无权访问", "通过", f"{EVID} → 他人状态快照权限"),
    ("权限", "管理员冒充候选人预约", "admin 携带令牌创建候选人面试", "POST /api/interview-sessions",
     "返回 403 仅候选人可预约", "拦截并返回角色不符", "通过", f"{EVID} → 管理员不能冒充候选人预约"),
    ("权限", "退出登录后令牌失效", "登出后复用原令牌", "GET /api/records/my",
     "令牌失效，拒绝访问", "拦截并返回未认证", "通过", f"{EVID} → 退出后令牌失效"),
]

wb3 = Workbook()
ws3 = wb3.active
ws3.title = "测试结果"
write_sheet(
    ws3,
    ["序号", "用例类别", "测试项", "测试输入或数据", "操作步骤",
     "预期结果", "实际结果", "结论", "证据位置（可定位）"],
    [[i] + list(c) for i, c in enumerate(CASES, start=1)],
    [6, 10, 26, 34, 38, 38, 46, 8, 52],
    title_rows=[
        ("多模态智能模拟面试评测平台（软件2324）—— 测试结果", 9),
        ("覆盖正常、异常、权限三类用例。全部结果来自实际执行，未使用固定返回值。"
         "执行方式见 tools/smoke_test.py、tools/smoke_ai.py、tools/smoke_vision.py。", 9),
    ],
)
for row in ws3.iter_rows(min_row=5, max_row=ws3.max_row, min_col=2, max_col=2):
    for c in row:
        if c.value == "正常":
            c.fill = PatternFill("solid", fgColor="E2EFDA")
        elif c.value == "异常":
            c.fill = PatternFill("solid", fgColor="FFF2CC")
        elif c.value == "权限":
            c.fill = PatternFill("solid", fgColor="FCE4D6")
        c.alignment = CENTER

# 第二个 sheet：执行记录与环境
ws4 = wb3.create_sheet("执行记录")
REC = [
    ("后端单元测试", "mvn -f interview-ai/pom.xml -B -ntp test", "9 个用例全部通过", "通过", BUILD),
    ("权限契约测试", "SecurityContractTest", "未登录读取身份被拒；角色取自认证主体", "通过", BUILD),
    ("验收回归测试", "AcceptanceRegressionTest", "终态不可重启、越权读取被拒、管理员不可创建候选人会话、"
     "融合分取真实值并保留输入、缺视频不虚报综合分、文本缺失与超范围分数按未评估处理", "通过", BUILD),
    ("现场冒烟：登录预约权限异常", "python tools/smoke_test.py", "22/22 passed", "通过", f"{EVID} → 1_登录预约查询越权_smoke_test"),
    ("现场冒烟：大模型出题与评分", "python tools/smoke_ai.py", "passed=true，问答得分 45.0，未测维度综合分为空", "通过", f"{EVID} → 2_大模型出题与问答评分_smoke_ai"),
    ("现场冒烟：视觉分析与融合", "python tools/smoke_vision.py", "passed=true，人脸 50 帧，视觉 65.0，综合 49.0", "通过", f"{EVID} → 3_视觉分析与融合报告_smoke_vision"),
]
write_sheet(ws4, ["执行项", "命令或脚本", "执行结果", "结论", "证据位置"],
            [list(r) for r in REC], [30, 42, 56, 8, 46])
r = ws4.max_row + 2
ws4.cell(r, 1, "执行环境").font = Font(bold=True, size=11)
for i, (k, v) in enumerate([
    ("操作系统", "Windows 11 专业版"),
    ("JDK", "21.0.5 LTS"),
    ("Maven", "3.9.9"),
    ("Node.js / npm", "v24.15.0 / 12.0.2"),
    ("Python", "3.10.6"),
    ("演示数据库", "H2 2.1.214（内置持久文件，演示模式免装 MySQL）"),
    ("执行日期", "2026-10-15"),
], start=1):
    ws4.cell(r + i, 1, k).font = Font(size=10)
    ws4.cell(r + i, 2, v).font = Font(size=10)

for name, wb_ in [("成员任务表.xlsx", wb), ("选题登记表.xlsx", wb2), ("测试结果.xlsx", wb3)]:
    p = os.path.join(OUT, name)
    wb_.save(p)
    print(f"  {name}  ({os.path.getsize(p)} bytes)")
