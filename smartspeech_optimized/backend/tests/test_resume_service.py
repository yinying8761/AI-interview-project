from __future__ import annotations

import io
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

from docx import Document

from backend.app import create_app
from backend.app.services.resume_service import ResumeError, ResumeService


def make_docx(*paragraphs: str, table_text: str = "") -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table_text:
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = table_text
        table.cell(0, 1).text = "Java 17"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


class ResumeServiceTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        root = Path(self.temp_dir.name)
        self.service = ResumeService(
            data_dir=root / "data",
            temp_dir=root / "temp",
            max_size=1024 * 1024,
        )

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_docx_resume_is_parsed_and_only_latest_text_is_saved(self) -> None:
        content = make_docx(
            "Candidate profile and backend engineering experience.",
            "Built a concurrent order service and improved throughput by 35 percent.",
            "Owned database indexing, Redis caching, and production incident response.",
            table_text="Project: distributed order platform",
        )

        item = self.service.save("user-a", "candidate.docx", content)

        self.assertEqual(item.extension, "docx")
        self.assertIn("concurrent order service", item.content)
        self.assertIn("distributed order platform", item.content)
        self.assertEqual(self.service.latest("user-a").id, item.id)
        saved_files = [
            path for path in (Path(self.temp_dir.name) / "data").rglob("*")
            if path.is_file()
        ]
        self.assertEqual([path.name for path in saved_files], ["latest.json"])

    def test_latest_resume_is_isolated_by_user_and_can_be_deleted(self) -> None:
        item = self.service.save(
            "user-a",
            "resume.docx",
            make_docx(
                "Backend engineer with Java, concurrency, database, Redis, "
                "RAG, algorithm, and project delivery experience."
            ),
        )

        self.assertIsNone(self.service.latest("user-b"))
        self.assertIsNone(self.service.get("user-b", item.id))
        self.assertTrue(self.service.delete("user-a"))
        self.assertIsNone(self.service.latest("user-a"))
        self.assertFalse(self.service.delete("user-a"))

    def test_rejects_unsupported_empty_oversized_and_textless_resumes(self) -> None:
        with self.assertRaises(ResumeError):
            self.service.save("user-a", "resume.txt", b"plain text" * 10)

        with self.assertRaises(ResumeError):
            self.service.save("user-a", "resume.pdf", b"")

        tiny_limit_service = ResumeService(
            data_dir=Path(self.temp_dir.name) / "small",
            temp_dir=Path(self.temp_dir.name) / "small-temp",
            max_size=1024,
        )
        with self.assertRaises(ResumeError):
            tiny_limit_service.save("user-a", "resume.pdf", b"x" * 1025)

        with self.assertRaises(ResumeError):
            self.service.save("user-a", "resume.docx", make_docx("short"))


class ResumeApiTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        root = Path(self.temp_dir.name)
        self.app = create_app()
        self.app.config.update(TESTING=True)
        self.resume_service = ResumeService(
            data_dir=root / "data",
            temp_dir=root / "temp",
        )
        self.app.extensions["services"] = SimpleNamespace(
            auth=SimpleNamespace(get_user=lambda user_id: {"id": user_id}),
            resumes=self.resume_service,
        )
        self.client = self.app.test_client()
        with self.client.session_transaction() as flask_session:
            flask_session["user_id"] = "user-a"

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_upload_latest_restore_isolation_and_delete(self) -> None:
        response = self.client.post(
            "/api/resumes",
            data={
                "resume": (
                    io.BytesIO(
                        make_docx(
                            "Java backend engineer with concurrency, database, "
                            "Redis, algorithm, RAG, and delivery experience."
                        )
                    ),
                    "candidate.docx",
                )
            },
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 201)
        resume = response.get_json()["resume"]
        self.assertEqual(resume["filename"], "candidate.docx")
        self.assertNotIn("content", resume)

        latest = self.client.get("/api/resumes/latest")
        self.assertEqual(latest.status_code, 200)
        self.assertEqual(latest.get_json()["resume"]["id"], resume["id"])

        with self.client.session_transaction() as flask_session:
            flask_session["user_id"] = "user-b"
        self.assertIsNone(self.client.get("/api/resumes/latest").get_json()["resume"])

        with self.client.session_transaction() as flask_session:
            flask_session["user_id"] = "user-a"
        deleted = self.client.delete("/api/resumes/latest")
        self.assertEqual(deleted.status_code, 200)
        self.assertIsNone(self.client.get("/api/resumes/latest").get_json()["resume"])

    def test_upload_rejects_unsupported_extension(self) -> None:
        response = self.client.post(
            "/api/resumes",
            data={"resume": (io.BytesIO(b"plain text"), "resume.txt")},
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("PDF", response.get_json()["error"])


if __name__ == "__main__":
    unittest.main()
