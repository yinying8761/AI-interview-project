from __future__ import annotations

from datetime import date
import json
from pathlib import Path
from typing import Any, Iterable, Sequence

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = PROJECT_ROOT / '智能模拟面试系统接口清单及接口文档.docx'

BLUE = '1F4E79'
TEAL = '147D73'
LIGHT_BLUE = 'DCE6F1'
LIGHT_TEAL = 'E2F0EF'
LIGHT_GRAY = 'F2F2F2'
WHITE = 'FFFFFF'
BLUE_RGB = RGBColor.from_string(BLUE)
TEAL_RGB = RGBColor.from_string(TEAL)
TEXT = RGBColor(48, 48, 48)
MUTED = RGBColor(100, 100, 100)

API_ENDPOINTS = [
    (1, '系统健康检查', 'GET', '/api/health', '否'),
    (2, '用户注册', 'POST', '/api/auth/register', '否'),
    (3, '密码登录', 'POST', '/api/auth/login', '否'),
    (4, '人脸登录', 'POST', '/api/auth/face-login', '否'),
    (5, '绑定或更新人脸', 'POST', '/api/auth/face-register', '是'),
    (6, '退出登录', 'POST', '/api/auth/logout', '否'),
    (7, '获取当前用户', 'GET', '/api/auth/me', '是'),
    (8, '创建面试会话', 'POST', '/api/interviews', '是'),
    (9, '生成开场与首题', 'POST', '/api/interviews/{session_id}/opening', '是'),
    (10, '获取面试会话', 'GET', '/api/interviews/{session_id}', '是'),
    (11, '获取实时指标', 'GET', '/api/interviews/{session_id}/metrics', '是'),
    (12, '放弃面试会话', 'DELETE', '/api/interviews/{session_id}', '是'),
    (13, '提交视频帧', 'POST', '/api/interviews/{session_id}/frames', '是'),
    (14, '提交回答并获取下一题', 'POST', '/api/interviews/{session_id}/messages', '是'),
    (15, '结束面试并生成报告', 'POST', '/api/interviews/{session_id}/finish', '是'),
    (16, '获取报告列表', 'GET', '/api/reports', '是'),
    (17, '获取报告详情', 'GET', '/api/reports/{session_id}', '是'),
    (18, '语音转文字', 'POST', '/api/media/transcribe', '是'),
    (19, '文本转语音', 'POST', '/api/media/tts', '是'),
    (20, '获取可用音色', 'GET', '/api/media/voices', '是'),
]


def example_text(value):
    if isinstance(value, str):
        return value.strip()
    return json.dumps(value, ensure_ascii=False, indent=2)


def add_label(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(10.5)
    run.font.color.rgb = TEAL_RGB
    set_east_asia_font(run, '微软雅黑')


def add_endpoint(
    document,
    number,
    name,
    method,
    path,
    auth,
    description,
    params,
    request_example,
    response_fields,
    response_example,
    errors,
):
    add_heading(document, f'4.{number} {name}', 2)
    add_table(
        document,
        ['请求方法', '接口路径', '登录要求'],
        [[method, path, auth]],
        [2.2, 9.6, 2.0],
    )
    add_body(document, description)
    add_label(document, '请求参数')
    if params:
        add_table(
            document,
            ['参数', '类型', '必填', '说明'],
            params,
            [3.2, 2.0, 1.4, 7.2],
        )
    else:
        add_body(document, '无请求参数。')
    add_label(document, '请求示例')
    add_code_block(document, example_text(request_example))
    add_label(document, '响应字段')
    add_table(
        document,
        ['字段', '类型', '说明'],
        response_fields,
        [4.0, 2.3, 7.5],
    )
    add_label(document, '响应示例')
    add_code_block(document, example_text(response_example))
    add_label(document, '主要错误')
    add_table(
        document,
        ['状态码', '场景', '处理建议'],
        errors,
        [1.8, 5.8, 6.2],
    )


def build_endpoint_details(document):
    add_heading(document, '4 接口详情', 1)
    add_endpoint(
        document, 1, '系统健康检查', 'GET', '/api/health', '否',
        '检查后端服务状态，并返回当前登录状态与核心依赖能力。',
        [],
        'curl http://127.0.0.1:8765/api/health',
        [
            ('ok', 'boolean', '服务是否可正常处理请求。'),
            ('name', 'string', '应用名称。'),
            ('authenticated', 'boolean', '当前 Cookie Session 是否已登录。'),
            ('capabilities', 'object', 'DeepSeek、Whisper、Edge-TTS、MediaPipe、人脸识别及前端构建能力。'),
        ],
        {
            'ok': True,
            'name': '智能体模拟面试系统',
            'authenticated': False,
            'capabilities': {
                'deepseek': True,
                'whisper': True,
                'edge_tts': True,
                'mediapipe': True,
                'face_recognition': True,
                'frontend_built': True,
            },
        },
        [
            ('500', '未捕获的服务端异常。', '查看服务日志并检查运行环境依赖。'),
        ],
    )
    add_endpoint(
        document, 2, '用户注册', 'POST', '/api/auth/register', '否',
        '创建本地账号并自动建立登录会话。请求可只提交账号信息，也可同时上传人脸图像完成人脸绑定。',
        [
            ('username', 'string', '是', '用户名，规范化后长度为 3-32 个字符，不能包含空格。'),
            ('password', 'string', '是', '密码，长度为 6-128 个字符。'),
            ('display_name', 'string', '否', '显示名称，最长 40 个字符；为空时使用用户名。'),
            ('image', 'file', '否', 'multipart 表单中的人脸图像字段。'),
            ('image_data', 'string', '否', 'JSON 模式下的 Base64 或 Data URL 图像。'),
        ],
        {
            'username': 'zhangsan',
            'password': 'secret123',
            'display_name': '张三',
            'image_data': 'data:image/jpeg;base64,/9j/4AAQSkZJRg...',
        },
        [
            ('user', 'object', '当前用户公开信息。'),
            ('user.id', 'string', '用户唯一标识。'),
            ('user.username', 'string', '用户名。'),
            ('user.display_name', 'string', '显示名称。'),
            ('user.created_at', 'string', 'ISO 8601 UTC 创建时间。'),
            ('user.face_registered', 'boolean', '是否已绑定人脸特征。'),
        ],
        {
            'user': {
                'id': '3f2b7f027f5e4df6a2b0f9e3d9f4c8a1',
                'username': 'zhangsan',
                'display_name': '张三',
                'created_at': '2026-09-14T02:20:30.123456+00:00',
                'face_registered': True,
            },
        },
        [
            ('400', '字段长度或格式不符合要求，或图像无法解析。', '按参数规则修正后重试。'),
            ('409', '用户名已存在。', '更换用户名或直接登录。'),
            ('413', '请求体超过 12 MiB。', '压缩图片后重试。'),
            ('500', '本地账号存储失败。', '检查 data/auth 目录写权限。'),
        ],
    )
    add_endpoint(
        document, 3, '密码登录', 'POST', '/api/auth/login', '否',
        '使用用户名和密码登录，成功后返回用户公开信息并写入会话 Cookie。',
        [
            ('username', 'string', '是', '注册时使用的用户名。'),
            ('password', 'string', '是', '账号密码。'),
        ],
        {'username': 'zhangsan', 'password': 'secret123'},
        [
            ('user', 'object', '当前用户公开信息，字段与注册接口一致。'),
        ],
        {
            'user': {
                'id': '3f2b7f027f5e4df6a2b0f9e3d9f4c8a1',
                'username': 'zhangsan',
                'display_name': '张三',
                'created_at': '2026-09-14T02:20:30.123456+00:00',
                'face_registered': True,
            },
        },
        [
            ('400', '请求体不是 JSON 对象。', '提交 application/json 请求体。'),
            ('401', '用户名或密码不正确。', '核对账号信息后重试。'),
            ('500', '账号存储读取失败。', '检查 data/auth/users.json。'),
        ],
    )
    add_endpoint(
        document, 4, '人脸登录', 'POST', '/api/auth/face-login', '否',
        '提取请求图像中的人脸特征，并与已注册的人脸特征库比对。匹配成功后建立登录会话。',
        [
            ('image', 'file', '二选一', 'multipart 表单中的人脸图像字段。'),
            ('image_data', 'string', '二选一', 'JSON 模式下的 Base64 或 Data URL 图像。'),
        ],
        {'image_data': 'data:image/jpeg;base64,/9j/4AAQSkZJRg...'},
        [
            ('user', 'object', '匹配到的用户公开信息。'),
            ('match.similarity', 'number', '人脸相似度，保留 4 位小数。'),
            ('match.distance', 'number', '人脸特征距离，保留 4 位小数。'),
        ],
        {
            'user': {
                'id': '3f2b7f027f5e4df6a2b0f9e3d9f4c8a1',
                'username': 'zhangsan',
                'display_name': '张三',
                'created_at': '2026-09-14T02:20:30.123456+00:00',
                'face_registered': True,
            },
            'match': {'similarity': 0.91, 'distance': 0.42},
        },
        [
            ('400', '图像缺失、编码无效或未检测到有效人脸。', '重新拍摄清晰正面照。'),
            ('401', '未匹配到已注册人脸或账号不存在。', '改用密码登录。'),
            ('500', '人脸特征库读取失败。', '检查 data/auth/faces 目录。'),
        ],
    )
    add_endpoint(
        document, 5, '绑定或更新人脸', 'POST', '/api/auth/face-register', '是',
        '为当前登录账号绑定人脸特征；重复调用会替换原有人脸特征。',
        [
            ('image', 'file', '二选一', 'multipart 表单中的人脸图像字段。'),
            ('image_data', 'string', '二选一', 'JSON 模式下的 Base64 或 Data URL 图像。'),
        ],
        'curl -b cookies.txt -F image=@face.jpg http://127.0.0.1:8765/api/auth/face-register',
        [
            ('user', 'object', '更新后的用户公开信息，face_registered 为 true。'),
        ],
        {
            'user': {
                'id': '3f2b7f027f5e4df6a2b0f9e3d9f4c8a1',
                'username': 'zhangsan',
                'display_name': '张三',
                'created_at': '2026-09-14T02:20:30.123456+00:00',
                'face_registered': True,
            },
        },
        [
            ('400', '图像无效、未检测到人脸或账号不存在。', '更换清晰正面照后重试。'),
            ('401', '未登录或会话失效。', '先登录并携带 Cookie。'),
            ('500', '人脸特征保存失败。', '检查人脸目录写权限。'),
        ],
    )
    add_endpoint(
        document, 6, '退出登录', 'POST', '/api/auth/logout', '否',
        '清除当前 Flask Session。接口本身不报未登录错误，重复调用仍返回成功。',
        [],
        'curl -b cookies.txt -c cookies.txt -X POST http://127.0.0.1:8765/api/auth/logout',
        [('ok', 'boolean', '固定为 true，表示会话已清除。')],
        {'ok': True},
        [('500', '服务端未捕获异常。', '检查服务日志。')],
    )
    add_endpoint(
        document, 7, '获取当前用户', 'GET', '/api/auth/me', '是',
        '根据会话中的 user_id 返回当前用户公开信息。前端通常在页面初始化时调用。',
        [],
        'curl -b cookies.txt http://127.0.0.1:8765/api/auth/me',
        [('user', 'object', '当前用户公开信息。')],
        {
            'user': {
                'id': '3f2b7f027f5e4df6a2b0f9e3d9f4c8a1',
                'username': 'zhangsan',
                'display_name': '张三',
                'created_at': '2026-09-14T02:20:30.123456+00:00',
                'face_registered': True,
            },
        },
        [
            ('401', '未登录或用户已不存在；服务端同时清除 Session。', '跳转登录页并重新认证。'),
        ],
    )
    add_endpoint(
        document, 8, '创建面试会话', 'POST', '/api/interviews', '是',
        '创建内存中的面试会话，并启动视觉分析工作线程。首题由 opening 接口异步生成，因此本接口返回较快。',
        [
            ('job_role', 'string', '否', '目标岗位，1-100 个字符，默认 Python 后端开发工程师。'),
            ('difficulty', 'string', '否', '面试难度，最长 20 个字符，默认中级。'),
            ('voice', 'string', '否', 'TTS 音色，默认 zh-CN-YunxiNeural。'),
            ('auto_tts', 'boolean', '否', '是否自动朗读面试官问题，默认 true。'),
            ('analysis_fps', 'number', '否', '目标视觉采样帧率，默认 3，服务端限制在 1-8。'),
            ('snapshot_interval', 'number', '否', '报告快照间隔秒数，默认 2，服务端限制在 0.5-10。'),
        ],
        {
            'job_role': 'Python 后端开发工程师',
            'difficulty': '中级',
            'voice': 'zh-CN-YunxiNeural',
            'auto_tts': True,
            'analysis_fps': 3,
            'snapshot_interval': 2,
        },
        [
            ('id', 'string', '会话 ID，32 位十六进制字符串。'),
            ('job_role/difficulty/voice', 'string', '创建后的会话配置。'),
            ('auto_tts', 'boolean', '自动朗读设置。'),
            ('analysis_fps', 'number', '实际生效的采样帧率。'),
            ('snapshot_interval', 'number', '实际生效的快照间隔。'),
            ('status', 'string', '会话状态：active、evaluating、stopped、finished。'),
            ('created_at', 'string', '创建时间，ISO 8601 UTC。'),
            ('finished_at', 'string/null', '结束时间，未结束时为 null。'),
            ('transcript', 'array', '消息数组，创建时为 []。'),
            ('live_metrics', 'object', '初始实时视觉指标。'),
            ('opening_required', 'boolean', '固定为 true，提示客户端调用 opening。'),
        ],
        {
            'id': '8cf12d6aa74d4a2d95d845fdef14a77c',
            'job_role': 'Python 后端开发工程师',
            'difficulty': '中级',
            'voice': 'zh-CN-YunxiNeural',
            'auto_tts': True,
            'analysis_fps': 3.0,
            'snapshot_interval': 2.0,
            'status': 'active',
            'created_at': '2026-09-14T02:25:00.123456+00:00',
            'finished_at': None,
            'transcript': [],
            'live_metrics': {'session_active': True, 'face_detected': False},
            'opening_required': True,
        },
        [
            ('400', '岗位、难度或数值参数无效。', '按参数范围修正后重试。'),
            ('401', '未登录或会话失效。', '先登录。'),
            ('500', '会话或视觉线程创建失败。', '检查 MediaPipe 与运行环境。'),
            ('413', '请求体超过 12 MiB。', '缩小请求体。'),
        ],
    )
    add_endpoint(
        document, 9, '生成开场与首题', 'POST', '/api/interviews/{session_id}/opening', '是',
        '调用 DeepSeek 生成开场白和第一题，并写入会话 transcript。接口幂等：若已有面试官消息，直接返回原消息。',
        [('session_id', 'path string', '是', '创建面试会话时返回的 id。')],
        'curl -b cookies.txt -X POST http://127.0.0.1:8765/api/interviews/8cf12d6aa74d4a2d95d845fdef14a77c/opening',
        [
            ('message', 'object', '首题消息；重复调用时返回会话中的第一条面试官消息。'),
            ('warning', 'string/null', 'DeepSeek 失败并采用备用首题时返回警告，正常时为 null。'),
            ('session', 'object', '更新后的会话公开状态。'),
        ],
        {
            'message': {
                'id': '9a39e4dd6e5f48bea98921ef2cb9c10e',
                'role': 'assistant',
                'content': '请先做一个自我介绍，并重点说明与目标岗位最匹配的一段经历。',
                'type': 'question',
                'created_at': '2026-09-14T02:25:02.123456+00:00',
            },
            'warning': None,
            'session': {
                'id': '8cf12d6aa74d4a2d95d845fdef14a77c',
                'status': 'active',
                'transcript': [{'role': 'assistant', 'type': 'question'}],
            },
        },
        [
            ('401', '未登录。', '先登录。'),
            ('404', '会话不存在或不属于当前账号。', '检查 session_id 和账号归属。'),
            ('500', '会话处理异常。', '检查服务日志。'),
        ],
    )
    add_endpoint(
        document, 10, '获取面试会话', 'GET', '/api/interviews/{session_id}', '是',
        '获取指定面试会话的最新状态、完整 transcript 和实时视觉指标。',
        [('session_id', 'path string', '是', '面试会话 ID。')],
        'curl -b cookies.txt http://127.0.0.1:8765/api/interviews/8cf12d6aa74d4a2d95d845fdef14a77c',
        [
            ('id', 'string', '会话 ID。'),
            ('job_role', 'string', '目标岗位。'),
            ('difficulty', 'string', '面试难度。'),
            ('voice', 'string', 'TTS 音色。'),
            ('auto_tts', 'boolean', '自动朗读设置。'),
            ('analysis_fps', 'number', '目标采样帧率。'),
            ('snapshot_interval', 'number', '快照间隔秒数。'),
            ('status', 'string', '当前会话状态。'),
            ('created_at/finished_at', 'string/null', '创建与结束时间。'),
            ('transcript', 'array', '完整消息数组。'),
            ('live_metrics', 'object', '实时视觉指标。'),
        ],
        {
            'id': '8cf12d6aa74d4a2d95d845fdef14a77c',
            'job_role': 'Python 后端开发工程师',
            'difficulty': '中级',
            'voice': 'zh-CN-YunxiNeural',
            'auto_tts': True,
            'analysis_fps': 3.0,
            'snapshot_interval': 2.0,
            'status': 'active',
            'created_at': '2026-09-14T02:25:00.123456+00:00',
            'finished_at': None,
            'transcript': [],
            'live_metrics': {'session_active': True, 'face_detected': True},
        },
        [
            ('401', '未登录。', '先登录。'),
            ('404', '会话不存在或不属于当前账号。', '检查 session_id。'),
        ],
    )
    add_endpoint(
        document, 11, '获取实时指标', 'GET', '/api/interviews/{session_id}/metrics', '是',
        '前端轮询该接口获取表情、注意力、采样量和视觉工作线程状态。该接口不消费视频帧。',
        [('session_id', 'path string', '是', '面试会话 ID。')],
        'curl -b cookies.txt http://127.0.0.1:8765/api/interviews/8cf12d6aa74d4a2d95d845fdef14a77c/metrics',
        [
            ('metrics', 'object', '实时指标；分析器不可用时为 {}。'),
            ('metrics.face_detected', 'boolean', '最新处理帧是否检测到人脸。'),
            ('metrics.expression', 'string', '当前表情标签。'),
            ('metrics.expression_score', 'number/null', '表情分，校准完成前为 null。'),
            ('metrics.attention_state', 'string', '注意力状态。'),
            ('metrics.attention_score', 'number/null', '覆盖度修正后的注意力分。'),
            ('metrics.calibration_progress', 'number', '表情校准进度，范围 0-1。'),
            ('metrics.processed_frames', 'integer', '已分析帧数。'),
            ('metrics.face_detection_rate', 'number', '人脸检测覆盖率百分比。'),
            ('metrics.trend', 'array', '最近最多 60 个轻量趋势点。'),
            ('vision', 'object', '视觉工作线程状态。'),
            ('vision.state', 'string', 'warming_up、ready、processing、degraded 或 stopped。'),
            ('vision.effective_fps', 'number', '最近完成帧的有效处理帧率。'),
            ('vision.last_latency_ms', 'number/null', '最近一次处理耗时。'),
        ],
        {
            'metrics': {
                'session_active': True,
                'face_detected': True,
                'expression': '平静',
                'expression_score': 82.5,
                'attention_state': 'attentive',
                'attention_score': 88.0,
                'calibration_progress': 1.0,
                'processed_frames': 126,
                'face_detection_rate': 98.4,
            },
            'vision': {
                'state': 'ready',
                'healthy': True,
                'processing': False,
                'pending': False,
                'processed_frames': 126,
                'dropped_frames': 2,
                'effective_fps': 2.8,
                'last_latency_ms': 142.7,
                'last_error': '',
            },
        },
        [
            ('401', '未登录。', '先登录。'),
            ('404', '会话不存在或不属于当前账号。', '检查 session_id。'),
        ],
    )
    add_endpoint(
        document, 12, '放弃面试会话', 'DELETE', '/api/interviews/{session_id}', '是',
        '删除当前账号的面试会话，关闭视觉线程并丢弃本轮采样指标和对话，不生成报告。',
        [('session_id', 'path string', '是', '要删除的面试会话 ID。')],
        'curl -b cookies.txt -X DELETE http://127.0.0.1:8765/api/interviews/8cf12d6aa74d4a2d95d845fdef14a77c',
        [('ok', 'boolean', '固定为 true。')],
        {'ok': True},
        [
            ('401', '未登录。', '先登录。'),
            ('404', '会话不存在或不属于当前账号。', '检查 session_id。'),
        ],
    )
    add_endpoint(
        document, 13, '提交视频帧', 'POST', '/api/interviews/{session_id}/frames', '是',
        '提交浏览器摄像头的最新画面用于表情和注意力分析。服务端只保留最新待处理帧，较旧待处理帧会被替换。',
        [
            ('session_id', 'path string', '是', '面试会话 ID。'),
            ('image', 'file', '二选一', 'multipart 表单中的 JPEG/PNG 图像。'),
            ('image_data', 'string', '二选一', 'JSON 模式下的 Base64 或 Data URL 图像。'),
        ],
        'curl -b cookies.txt -F image=@frame.jpg http://127.0.0.1:8765/api/interviews/8cf12d6aa74d4a2d95d845fdef14a77c/frames',
        [
            ('accepted', 'boolean', '是否已接收；成功时固定为 true。'),
            ('replaced_pending', 'boolean', '提交时是否替换了尚未处理的旧帧。'),
            ('metrics', 'object', '提交完成后的最新实时指标。'),
            ('vision', 'object', '视觉工作线程状态。'),
        ],
        {
            'accepted': True,
            'replaced_pending': False,
            'metrics': {
                'face_detected': True,
                'expression': '平静',
                'attention_state': 'attentive',
            },
            'vision': {
                'state': 'ready',
                'healthy': True,
                'processing': False,
                'pending': False,
                'dropped_frames': 2,
                'effective_fps': 2.8,
            },
        },
        [
            ('400', '图像缺失、编码无效或无法解析。', '提交有效 JPEG/PNG 图像。'),
            ('401', '未登录。', '先登录。'),
            ('404', '会话不存在或不属于当前账号。', '检查 session_id。'),
            ('409', '会话不在 active 状态，或视觉服务已停止。', '只在面试进行中提交帧。'),
            ('413', '帧数据超过 12 MiB。', '降低分辨率或 JPEG 质量。'),
            ('500', '帧算法处理异常。', '检查 MediaPipe 运行状态。'),
            ('503', '视觉分析服务尚未就绪。', '稍后重试并查看 vision.state。'),
        ],
    )
    add_endpoint(
        document, 14, '提交回答并获取下一题', 'POST', '/api/interviews/{session_id}/messages', '是',
        '将候选人回答发送给 DeepSeek，生成面试官下一题，并把两条消息写入 transcript。',
        [
            ('session_id', 'path string', '是', '面试会话 ID。'),
            ('message', 'string', '是', '候选人回答，去除首尾空白后不能为空，最长 5000 个字符。'),
        ],
        {
            'message': '我负责过一个订单查询接口优化，通过增加复合索引和缓存，把 P95 从 800ms 降到 120ms。',
        },
        [
            ('user_message', 'object', '已保存的候选人消息，type 为 answer。'),
            ('assistant_message', 'object', '面试官下一题，type 为 question。'),
            ('session', 'object', '更新后的完整会话公开状态。'),
            ('message.id', 'string', '消息唯一标识。'),
            ('message.role', 'string', 'user 或 assistant。'),
            ('message.content', 'string', '消息正文。'),
            ('message.type', 'string', 'answer 或 question。'),
            ('message.created_at', 'string', 'ISO 8601 UTC 时间。'),
        ],
        {
            'user_message': {
                'id': '44977b8512524ce4b3ecf430f12f8890',
                'role': 'user',
                'content': '我负责过一个订单查询接口优化，通过增加复合索引和缓存，把 P95 从 800ms 降到 120ms。',
                'type': 'answer',
                'created_at': '2026-09-14T02:27:10.123456+00:00',
            },
            'assistant_message': {
                'id': 'cba54732a9d74b0bb23cc04740f54d89',
                'role': 'assistant',
                'content': '你如何验证缓存一致性？如果缓存与数据库短暂不一致，你会怎样处理？',
                'type': 'question',
                'created_at': '2026-09-14T02:27:12.123456+00:00',
            },
            'session': {
                'id': '8cf12d6aa74d4a2d95d845fdef14a77c',
                'status': 'active',
                'transcript': [],
            },
        },
        [
            ('400', 'message 为空或超过 5000 个字符。', '修正回答内容。'),
            ('401', '未登录。', '先登录。'),
            ('404', '会话不存在或不属于当前账号。', '检查 session_id。'),
            ('409', '会话已结束，不在 active 状态。', '新建面试会话。'),
            ('502', 'DeepSeek 请求失败或返回无效内容。', '检查 DEEPSEEK_API_KEY、网络和上游状态后重试。'),
        ],
    )
    add_endpoint(
        document, 15, '结束面试并生成报告', 'POST', '/api/interviews/{session_id}/finish', '是',
        '停止视觉采样，调用 DeepSeek 评价完整问答记录，融合表情与注意力分数，并保存报告。已有完成报告时幂等返回。',
        [('session_id', 'path string', '是', '面试会话 ID。')],
        'curl -b cookies.txt -X POST http://127.0.0.1:8765/api/interviews/8cf12d6aa74d4a2d95d845fdef14a77c/finish',
        [
            ('report.version', 'integer', '报告结构版本，当前为 1。'),
            ('report.session_id', 'string', '面试会话 ID。'),
            ('report.created_at/finished_at', 'string', '面试创建和结束时间。'),
            ('report.duration_seconds', 'number', '视觉会话统计时长。'),
            ('report.transcript', 'array', '完整问答记录。'),
            ('report.question_evaluation', 'object', 'DeepSeek 文本评价。'),
            ('report.video_summary', 'object/null', '表情与注意力汇总。'),
            ('report.score_fusion', 'object', '问答与视频分数融合结果。'),
            ('report.privacy', 'object', '不保存原始视频帧、人脸图片和音频的声明。'),
        ],
        {
            'report': {
                'version': 1,
                'session_id': '8cf12d6aa74d4a2d95d845fdef14a77c',
                'created_at': '2026-09-14T02:25:00.123456+00:00',
                'finished_at': '2026-09-14T02:31:20.123456+00:00',
                'job_role': 'Python 后端开发工程师',
                'difficulty': '中级',
                'duration_seconds': 372.4,
                'transcript': [],
                'question_evaluation': {
                    'overall_score': 78,
                    'level': '中等',
                    'dimensions': {'专业能力': 78, '逻辑表达': 80},
                    'summary': '技术基础较扎实，建议补充量化结果。',
                    'strengths': [],
                    'improvements': [],
                    'question_reviews': [],
                    'next_steps': [],
                },
                'video_summary': {'video_available': True},
                'score_fusion': {
                    'available': True,
                    'question_score': 78.0,
                    'expression_score': 82.0,
                    'attention_score': 88.0,
                    'expression_attention_score': 85.0,
                    'final_score': 79.4,
                },
                'privacy': {
                    'stores_video_frames': False,
                    'stores_face_images': False,
                    'stores_audio': False,
                },
            },
        },
        [
            ('400', '尚无候选人回答，或评价输入无效。', '至少完成一次回答后再结束。'),
            ('401', '未登录。', '先登录。'),
            ('404', '会话不存在或不属于当前账号。', '检查 session_id。'),
            ('502', 'DeepSeek 评价失败。', '检查 API Key、网络和上游状态；会话状态会恢复为 stopped。'),
            ('500', '报告写入失败。', '检查 runtime/reports 目录写权限。'),
        ],
    )
    add_endpoint(
        document, 16, '获取报告列表', 'GET', '/api/reports', '是',
        '读取当前账号已保存的报告摘要，按 created_at 倒序排列，最多返回 50 条。',
        [],
        'curl -b cookies.txt http://127.0.0.1:8765/api/reports',
        [
            ('reports', 'array', '报告摘要数组。'),
            ('session_id/job_role/difficulty', 'string', '报告对应的会话基础信息。'),
            ('created_at', 'string', '面试创建时间，ISO 8601 UTC。'),
            ('duration_seconds', 'number', '会话统计时长。'),
            ('question_score', 'number/null', '问答得分。'),
            ('expression_score', 'number/null', '表情得分。'),
            ('attention_score', 'number/null', '注意力得分。'),
            ('final_score', 'number/null', '最终得分；视频样本不足时为 null。'),
            ('level', 'string', '文本评价等级。'),
            ('summary', 'string', '文本评价总结。'),
        ],
        {
            'reports': [
                {
                    'session_id': '8cf12d6aa74d4a2d95d845fdef14a77c',
                    'created_at': '2026-09-14T02:25:00.123456+00:00',
                    'job_role': 'Python 后端开发工程师',
                    'difficulty': '中级',
                    'duration_seconds': 372.4,
                    'question_score': 78.0,
                    'expression_score': 82.0,
                    'attention_score': 88.0,
                    'final_score': 79.4,
                    'level': '中等',
                    'summary': '技术基础较扎实，建议补充量化结果。',
                },
            ],
        },
        [
            ('401', '未登录。', '先登录。'),
        ],
    )
    add_endpoint(
        document, 17, '获取报告详情', 'GET', '/api/reports/{session_id}', '是',
        '读取指定报告。优先从持久化报告目录读取；内存会话仍有报告时也会作为回退来源。',
        [('session_id', 'path string', '是', '面试会话 ID。')],
        'curl -b cookies.txt http://127.0.0.1:8765/api/reports/8cf12d6aa74d4a2d95d845fdef14a77c',
        [('report', 'object', '完整报告对象，字段与结束面试接口返回的 report 一致。')],
        {'report': {'version': 1, 'session_id': '8cf12d6aa74d4a2d95d845fdef14a77c'}},
        [
            ('401', '未登录。', '先登录。'),
            ('404', '报告不存在。', '确认 session_id，或完成面试生成报告。'),
        ],
    )
    add_endpoint(
        document, 18, '语音转文字', 'POST', '/api/media/transcribe', '是',
        '接收录音文件，使用 Whisper 转写文本并返回分段结果。临时音频文件在请求结束后删除。',
        [
            ('audio', 'file', '是', 'multipart 表单音频；支持 wav、mp3、m4a、mp4、ogg、webm、flac。'),
        ],
        'curl -b cookies.txt -F audio=@answer.webm http://127.0.0.1:8765/api/media/transcribe',
        [
            ('text', 'string', '完整转写文本。'),
            ('language', 'string/null', 'Whisper 识别的语言代码。'),
            ('segments', 'array', '转写分段数组。'),
            ('segments[].start', 'number', '分段起始秒数。'),
            ('segments[].end', 'number', '分段结束秒数。'),
            ('segments[].text', 'string', '分段文本。'),
        ],
        {
            'text': '我主要负责订单系统的接口开发与性能优化。',
            'language': 'zh',
            'segments': [
                {'start': 0.0, 'end': 2.8, 'text': '我主要负责订单系统的接口开发与性能优化。'},
            ],
        },
        [
            ('400', '未上传 audio、文件为空或无法读取。', '重新录音并提交非空音频文件。'),
            ('401', '未登录。', '先登录。'),
            ('413', '上传内容超过 12 MiB。', '缩短录音或压缩音频。'),
            ('500', 'Whisper 模型加载或转写失败。', '检查 whisper 依赖、模型文件和运行内存。'),
        ],
    )
    add_endpoint(
        document, 19, '文本转语音', 'POST', '/api/media/tts', '是',
        '把面试官问题转换为语音。优先使用 Edge-TTS；在线合成失败时，Windows 环境会回退到本地 SAPI。',
        [
            ('text', 'string', '是', '待合成文本，去除首尾空白后不能为空，最长 5000 个字符。'),
            ('voice', 'string', '否', 'Edge-TTS 音色，为空时使用服务端默认音色。'),
        ],
        {'text': '请介绍一下你最有挑战性的项目。', 'voice': 'zh-CN-YunxiNeural'},
        [
            ('响应体', 'binary', '音频二进制流，不是 JSON。'),
            ('Content-Type', 'header', '通常为 audio/mpeg；Windows SAPI 回退时可能为 audio/wav。'),
            ('Content-Disposition', 'header', 'inline，并包含服务端临时文件名。'),
        ],
        'HTTP/1.1 200 OK\nContent-Type: audio/mpeg\nContent-Length: 24576\n\n<binary audio data>',
        [
            ('400', 'text 为空、请求体不是 JSON 或文本超过 5000 字符。', '提交有效 JSON 和文本。'),
            ('401', '未登录。', '先登录。'),
            ('500', 'Edge-TTS 与本地语音引擎均合成失败，或输出音频无效。', '检查网络、系统语音和运行环境。'),
            ('413', '请求体超过 12 MiB。', '缩短文本。'),
        ],
    )
    add_endpoint(
        document, 20, '获取可用音色', 'GET', '/api/media/voices', '是',
        '返回 Edge-TTS 中的中文音色。在线查询失败时返回内置中文音色列表，前端仍可正常选择。',
        [],
        'curl -b cookies.txt http://127.0.0.1:8765/api/media/voices',
        [
            ('voices', 'array', '音色数组。'),
            ('voices[].short_name', 'string', '传给 TTS 接口的 voice 值。'),
            ('voices[].label', 'string', '界面显示名称。'),
        ],
        {
            'voices': [
                {'short_name': 'zh-CN-YunxiNeural', 'label': '云希 · 男声'},
                {'short_name': 'zh-CN-XiaoxiaoNeural', 'label': '晓晓 · 女声'},
            ],
        },
        [
            ('401', '未登录。', '先登录。'),
        ],
    )


def build_front_matter(document):
    for _ in range(4):
        document.add_paragraph()
    title = document.add_paragraph(style='Title')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run('智能模拟面试系统')
    title_run.bold = True
    title_run.font.size = Pt(30)
    title_run.font.color.rgb = BLUE_RGB
    set_east_asia_font(title_run, '微软雅黑')
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run('接口清单及接口文档')
    subtitle_run.bold = True
    subtitle_run.font.size = Pt(21)
    subtitle_run.font.color.rgb = TEAL_RGB
    set_east_asia_font(subtitle_run, '微软雅黑')
    document.add_paragraph()
    add_table(
        document,
        ['项目', '内容'],
        [
            ('文档版本', 'V1.0'),
            ('接口版本', 'API v1（当前 Flask 路由实现）'),
            ('后端技术', 'Flask 3.1.3 + Waitress + Flask-CORS'),
            ('默认服务地址', 'http://127.0.0.1:8765'),
            ('接口前缀', '/api'),
            ('生成日期', date.today().isoformat()),
            ('适用项目', 'smartspeech_optimized_new'),
        ],
        [3.4, 10.4],
    )
    document.add_page_break()
    add_heading(document, '文档信息', 1)
    add_table(
        document,
        ['维护项', '说明'],
        [
            ('文档目标', '为前后端联调、测试、部署验收和后续维护提供完整接口依据。'),
            ('接口来源', 'backend/app/api 下的 Flask Blueprint 路由。'),
            ('覆盖范围', '系统、认证、面试会话、报告、语音识别和语音合成，共 20 个接口。'),
            ('数据口径', '请求参数、字段、限制、状态码和评分公式均按当前源码实现整理。'),
            ('阅读建议', '先阅读第 2、3 章，再按第 4 章接口详情进行开发。'),
        ],
        [3.2, 10.6],
    )
    add_heading(document, '目录', 1)
    add_toc(document)
    document.add_page_break()


def build_introduction(document):
    add_heading(document, '1 文档说明', 1)
    add_heading(document, '1.1 编写目的', 2)
    add_body(document, '本文档说明智能模拟面试系统当前版本的 HTTP API，包括接口清单、请求参数、响应字段、错误状态、核心数据模型和典型调用流程。')
    add_heading(document, '1.2 目标读者', 2)
    add_bullets(document, [
        '前端开发人员：按接口契约实现登录、面试工作台和报告页面。',
        '后端开发人员：维护现有接口、扩展字段和执行回归测试。',
        '测试人员：设计正常请求、边界请求和异常链路测试用例。',
        '部署与验收人员：核对接口可用性、依赖能力和隐私数据处理方式。',
    ])
    add_heading(document, '1.3 文档边界', 2)
    add_body(document, '本文档只描述通过 /api 暴露的接口。前端页面路由、静态资源服务、Flask 内部函数和第三方 DeepSeek API 的私有协议不作为对外接口承诺。')


def build_overview(document):
    add_heading(document, '2 接口清单', 1)
    add_body(document, '所有接口均使用 /api 前缀。登录要求中的“是”表示请求必须携带有效会话 Cookie。')
    add_table(
        document,
        ['序号', '模块', '接口名称', '方法', '路径', '登录'],
        [
            (number, path.split('/')[2], name, method, path, auth)
            for number, name, method, path, auth in API_ENDPOINTS
        ],
        [1.1, 1.8, 3.1, 1.5, 5.6, 1.1],
    )


def build_conventions(document):
    add_heading(document, '3 通用约定', 1)
    add_heading(document, '3.1 服务地址与跨域', 2)
    add_body(document, '默认监听地址为 http://127.0.0.1:8765，生产部署时应替换为实际域名。CORS 默认允许 http://127.0.0.1:5173 和 http://localhost:5173，并允许携带凭据。')
    add_heading(document, '3.2 认证与会话', 2)
    add_body(document, '系统使用 Flask Cookie Session，不返回 Bearer Token。Cookie 名为 smart_interview_session，有效期 12 小时，HttpOnly 且 SameSite=Lax。')
    add_bullets(document, [
        '浏览器请求需设置 credentials: include。',
        '服务端启用 HTTPS 后，应将 SESSION_COOKIE_SECURE 配置为 1。',
        '获得 401 后应清除前端登录态，并引导用户重新登录。',
        '会话中的所有面试和报告均按 user_id 校验归属，不能访问其他账号的数据。',
    ])
    add_heading(document, '3.3 数据格式与上传限制', 2)
    add_table(
        document,
        ['项目', '约定'],
        [
            ('JSON', '请求头使用 Content-Type: application/json，响应默认 UTF-8 JSON。'),
            ('multipart', '人脸图像使用 image；音频转写使用 audio；视频帧使用 image。'),
            ('图片 JSON', '使用 image_data 传 Base64 或 data:image/...;base64,...。'),
            ('请求体上限', 'MAX_CONTENT_LENGTH 为 12 MiB；超限返回 413。'),
            ('时间', '服务端时间为 ISO 8601 UTC，例如 2026-09-14T02:25:00+00:00。'),
            ('ID', '用户和面试会话 ID 为 UUID hex 字符串；消息 ID 为 32 位十六进制字符串。'),
        ],
        [3.2, 10.6],
    )
    add_heading(document, '3.4 响应与错误', 2)
    add_body(document, '成功响应可能是 JSON 对象或二进制音频。失败响应统一为 JSON 对象，格式如下：')
    add_code_block(document, json.dumps({'error': '错误说明'}, ensure_ascii=False, indent=2))
    add_table(
        document,
        ['状态码', '含义', '常见场景'],
        [
            ('200', '成功', '查询、更新、登录、语音合成等正常完成。'),
            ('201', '创建成功', '注册账号或创建面试会话。'),
            ('400', '参数错误', '字段缺失、格式错误、图片无法解析、文本超长。'),
            ('401', '未认证', '未登录、会话失效、密码错误或人脸未匹配。'),
            ('404', '不存在', '会话、报告不存在或不属于当前账号。'),
            ('409', '状态冲突', '用户名重复、会话已结束、视觉服务已停止。'),
            ('413', '内容过大', '上传内容超过 12 MiB。'),
            ('500', '服务端错误', '存储失败、模型运行失败或未捕获异常。'),
            ('502', '上游错误', 'DeepSeek 请求失败或返回内容不可用。'),
            ('503', '服务不可用', '视觉分析服务尚未就绪。'),
        ],
        [2.0, 2.5, 9.3],
    )


def build_models(document):
    add_heading(document, '5 核心数据模型', 1)
    add_heading(document, '5.1 User 与 Message', 2)
    add_body(document, 'User 是接口返回的账号公开模型，不包含密码哈希或人脸向量。Message 是面试对话中的统一消息模型。')
    add_table(
        document,
        ['模型', '字段', '类型', '说明'],
        [
            ('User', 'id', 'string', '用户唯一标识。'),
            ('User', 'username', 'string', '用户名。'),
            ('User', 'display_name', 'string', '界面显示名称。'),
            ('User', 'created_at', 'string', '创建时间。'),
            ('User', 'face_registered', 'boolean', '是否已绑定人脸。'),
            ('Message', 'id', 'string', '消息 ID。'),
            ('Message', 'role', 'string', 'user 或 assistant。'),
            ('Message', 'content', 'string', '消息正文。'),
            ('Message', 'type', 'string', 'answer 或 question。'),
            ('Message', 'created_at', 'string', '消息创建时间。'),
        ],
        [2.2, 3.0, 2.0, 6.6],
    )
    add_heading(document, '5.2 InterviewSession', 2)
    add_table(
        document,
        ['字段', '类型', '说明'],
        [
            ('id', 'string', '面试会话 ID。'),
            ('job_role', 'string', '目标岗位。'),
            ('difficulty', 'string', '面试难度。'),
            ('voice', 'string', 'TTS 音色。'),
            ('auto_tts', 'boolean', '是否自动朗读。'),
            ('analysis_fps', 'number', '采样帧率，范围 1-8。'),
            ('snapshot_interval', 'number', '快照间隔，范围 0.5-10 秒。'),
            ('status', 'string', 'active、evaluating、stopped、finished。'),
            ('created_at', 'string', '创建时间。'),
            ('finished_at', 'string/null', '结束时间。'),
            ('transcript', 'Message[]', '完整对话记录。'),
            ('live_metrics', 'object', '实时视觉指标，可选。'),
        ],
        [4.2, 2.4, 7.2],
    )
    add_heading(document, '5.3 实时指标 Metrics', 2)
    add_table(
        document,
        ['字段', '类型', '说明'],
        [
            ('session_active', 'boolean', '分析会话是否仍处于活动状态。'),
            ('face_detected', 'boolean', '最新帧是否检测到人脸。'),
            ('expression', 'string', '当前表情标签，例如平静或微笑。'),
            ('expression_score', 'number/null', '表情评分，校准完成前为 null。'),
            ('attention_state', 'string', '注意力状态。'),
            ('attention_score', 'number/null', '覆盖度修正后的注意力评分。'),
            ('attention_coverage_factor', 'number', '按人脸覆盖率计算的修正系数。'),
            ('calibration_progress', 'number', '表情校准进度，范围 0-1。'),
            ('pitch/yaw/roll', 'number', '头部姿态角。'),
            ('gaze_x', 'number', '水平视线偏移。'),
            ('attention_rate', 'number', '有效视线居中比例。'),
            ('distraction_count', 'integer', '注意力分散次数。'),
            ('blink_count', 'integer', '眨眼次数。'),
            ('expression_changes', 'integer', '表情变化次数。'),
            ('video_score_available', 'boolean', '视频分数是否达到可用条件。'),
            ('processed_frames', 'integer', '累计处理帧数。'),
            ('face_detection_rate', 'number', '人脸检测覆盖率百分比。'),
            ('missing_face_episodes', 'integer', '离开画面次数。'),
            ('snapshot_count', 'integer', '已生成快照数量。'),
            ('trend', 'array', '最近最多 60 个趋势点。'),
        ],
        [4.4, 2.3, 7.1],
    )
    add_heading(document, '5.4 Report', 2)
    add_table(
        document,
        ['字段', '类型', '说明'],
        [
            ('version', 'integer', '报告结构版本，当前为 1。'),
            ('session_id', 'string', '面试会话 ID。'),
            ('created_at/finished_at', 'string', '面试创建和结束时间。'),
            ('job_role/difficulty', 'string', '面试配置。'),
            ('duration_seconds', 'number', '视觉分析会话时长。'),
            ('transcript', 'Message[]', '完整问答记录。'),
            ('question_evaluation', 'object', 'DeepSeek 文本评价。'),
            ('video_summary', 'object/null', '表情、注意力、事件和快照汇总。'),
            ('score_fusion', 'object', '最终评分融合结果。'),
            ('privacy', 'object', '数据留存声明。'),
        ],
        [4.2, 2.4, 7.2],
    )
    add_heading(document, '5.5 ScoreFusion', 2)
    add_table(
        document,
        ['字段', '类型', '说明'],
        [
            ('available', 'boolean', '视频得分是否可用，并据此决定是否生成 final_score。'),
            ('question_score', 'number', '问答分，范围 0-100。'),
            ('expression_score', 'number/null', '表情分，不可用时为 null。'),
            ('attention_score', 'number/null', '注意力分，不可用时为 null。'),
            ('expression_attention_score', 'number/null', '表情分与注意力分各占 50% 的视频综合分。'),
            ('final_score', 'number/null', '最终得分；视频样本不足时为 null。'),
            ('weights', 'object', '评分权重说明。'),
            ('formula', 'string', '评分公式文本。'),
            ('video_sample_count', 'integer', '参与视频评分的有效人脸帧数。'),
            ('note', 'string', '最终分数是否生成的中文说明。'),
        ],
        [5.0, 2.3, 6.5],
    )
    add_body(document, '最终评分公式为：final_score = question_score × 0.8 + (expression_score × 0.5 + attention_score × 0.5) × 0.2。视频样本不足时 final_score 为 null，系统不会把视频权重重新分配给问答分。')


def build_workflows(document):
    add_heading(document, '6 典型调用流程', 1)
    add_heading(document, '6.1 账号认证流程', 2)
    add_bullets(document, [
        '页面启动时调用 GET /api/auth/me 恢复登录状态。',
        '新用户调用 POST /api/auth/register 创建账号；可同时上传人脸。',
        '已有用户调用 POST /api/auth/login，或调用 POST /api/auth/face-login。',
        '登录后如需补充或更换人脸，调用 POST /api/auth/face-register。',
        '退出时调用 POST /api/auth/logout。',
    ])
    add_heading(document, '6.2 完整面试流程', 2)
    add_code_block(document, '\n'.join([
        'GET  /api/media/voices',
        'POST /api/interviews',
        'POST /api/interviews/{session_id}/opening',
        '循环：POST /api/interviews/{session_id}/frames',
        '循环：GET  /api/interviews/{session_id}/metrics',
        'POST /api/media/transcribe  # 录音回答时可选',
        'POST /api/interviews/{session_id}/messages',
        'POST /api/media/tts         # 自动朗读时可选',
        'POST /api/interviews/{session_id}/finish',
        'GET  /api/reports/{session_id}',
    ]))
    add_body(document, '前端当前每 800 毫秒轮询一次指标；视频帧目标间隔由 analysis_fps 决定，并由客户端按请求延迟自适应退避。')
    add_heading(document, '6.3 报告查询流程', 2)
    add_bullets(document, [
        '报告列表页调用 GET /api/reports 获取最多 50 条摘要。',
        '用户打开某条记录时调用 GET /api/reports/{session_id} 获取完整报告。',
        '报告详情包含文本评价、完整 transcript、视频汇总、评分融合和隐私说明。',
        '放弃面试时调用 DELETE /api/interviews/{session_id}，此操作不会生成报告。',
    ])


def build_appendix(document):
    add_heading(document, '7 安全、隐私与维护', 1)
    add_heading(document, '7.1 安全建议', 2)
    add_bullets(document, [
        '生产环境必须设置高强度 FLASK_SECRET_KEY，避免会话 Cookie 被伪造。',
        '生产环境启用 HTTPS，并设置 SESSION_COOKIE_SECURE=1。',
        '将 CORS_ORIGINS 限制为实际前端域名。',
        'DeepSeek API Key 只保存在服务端 .env，不得下发给浏览器。',
        '日志中不要记录密码、Cookie、Base64 图片、人脸向量或完整音频内容。',
    ])
    add_heading(document, '7.2 数据留存与隐私', 2)
    add_table(
        document,
        ['数据', '处理方式'],
        [
            ('账号数据', '保存于 data/auth/users.json；密码使用 PBKDF2-SHA256 加盐哈希。'),
            ('人脸特征', '只保存归一化特征向量到 data/auth/faces，不保存注册照片。'),
            ('视频帧', '仅在内存队列中分析，原始帧不落盘。'),
            ('人脸图片', '只用于提取特征，不写入磁盘。'),
            ('回答录音', '转写结束后删除临时文件。'),
            ('面试报告', '保存文本、指标、事件、快照和评分到 runtime/reports。'),
            ('面试会话', '保存在服务器内存中，默认最长保留 8 小时后清理。'),
        ],
        [3.5, 10.3],
    )
    add_heading(document, '7.3 版本与兼容性', 2)
    add_bullets(document, [
        '当前接口无独立 URL 版本号，基础前缀为 /api；破坏性变更应先引入兼容期。',
        '新增响应字段应保持向后兼容；客户端不应因未知字段而失败。',
        '报告使用 version 字段标注结构版本，当前为 1。',
        '修改评分公式、状态枚举或字段语义时，应同步更新本文档和回归测试。',
    ])
    add_heading(document, '7.4 联调检查清单', 2)
    add_bullets(document, [
        '确认后端 http://127.0.0.1:8765/api/health 返回 ok=true。',
        '确认请求启用 Cookie 凭据，并正确处理 401。',
        '创建会话后先调用 opening，不要假设创建接口已生成首题。',
        '只在 status=active 时提交视频帧和回答。',
        '结束面试前至少提交一次候选人回答。',
        'TTS 响应按 Blob 或二进制流处理，不能调用 response.json()。',
        '报告列表的 final_score 可能为 null，界面必须兼容该状态。',
    ])


def build_document():
    document = Document()
    configure_document(document)
    build_front_matter(document)
    build_introduction(document)
    build_overview(document)
    build_conventions(document)
    build_endpoint_details(document)
    build_models(document)
    build_workflows(document)
    build_appendix(document)
    add_page_number(document)
    document.save(OUTPUT_PATH)
    return OUTPUT_PATH


if False:
    output = build_document()
    print(f'Generated: {output}')
    add_heading(document, '文档信息', 1)
    add_table(
        document,
        ['维护项', '说明'],
        [
            ('文档目标', '为前后端联调、测试、部署验收和后续维护提供完整接口依据。'),
            ('接口来源', 'backend/app/api 下的 Flask Blueprint 路由。'),
            ('覆盖范围', '系统、认证、面试会话、报告、语音识别和语音合成，共 20 个接口。'),
            ('数据口径', '请求参数、字段、限制、状态码和评分公式均按当前源码实现整理。'),
            ('阅读建议', '先阅读第 2、3 章，再按第 4 章接口详情进行开发。'),
        ],
        [3.2, 10.6],
    )
    add_heading(document, '目录', 1)
    add_toc(document)
    document.add_page_break()


def set_east_asia_font(run, name: str) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn('w:shd'))
    if shading is None:
        shading = OxmlElement('w:shd')
        properties.append(shading)
    shading.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in('w:tcMar')
    if margins is None:
        margins = OxmlElement('w:tcMar')
        properties.append(margins)
    for margin_name, value in (
        ('top', top),
        ('start', start),
        ('bottom', bottom),
        ('end', end),
    ):
        node = margins.find(qn(f'w:{margin_name}'))
        if node is None:
            node = OxmlElement(f'w:{margin_name}')
            margins.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row) -> None:
    row_properties = row._tr.get_or_add_trPr()
    marker = OxmlElement('w:tblHeader')
    marker.set(qn('w:val'), 'true')
    row_properties.append(marker)


def prevent_row_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement('w:cantSplit'))


def set_cell_text(cell, value: Any, *, bold=False, color=None, size=9.5) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run("" if value is None else str(value))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color or TEXT
    set_east_asia_font(run, '微软雅黑')
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_table(document, headers, rows, widths=None):
    rows = list(rows)
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    table.autofit = True
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, BLUE)
        set_cell_text(cell, header, bold=True, color=RGBColor(255, 255, 255))
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for index, value in enumerate(values):
            set_cell_text(cells[index], value)
            if row_index % 2 == 1:
                set_cell_shading(cells[index], LIGHT_GRAY)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    document.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_code_block(document, text):
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_GRAY)
    set_cell_margins(cell, top=120, start=160, bottom=120, end=160)
    cell.text = ''
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    for line_index, line in enumerate(text.strip('\n').splitlines()):
        if line_index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        run.font.name = 'Consolas'
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(45, 45, 45)
        set_east_asia_font(run, '微软雅黑')
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_bullets(document, items):
    for item in items:
        paragraph = document.add_paragraph(style='List Bullet')
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(item)
        run.font.size = Pt(10)
        set_east_asia_font(run, '微软雅黑')


def add_body(document, text, bold_prefix=''):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.45
    if bold_prefix:
        prefix = paragraph.add_run(bold_prefix)
        prefix.bold = True
        prefix.font.size = Pt(10.5)
        set_east_asia_font(prefix, '微软雅黑')
    run = paragraph.add_run(text)
    run.font.size = Pt(10.5)
    run.font.color.rgb = TEXT
    set_east_asia_font(run, '微软雅黑')


def add_heading(document, text, level):
    heading = document.add_heading(text, level=level)
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.space_before = Pt(10 if level > 1 else 14)
    heading.paragraph_format.space_after = Pt(6)
    return heading


def add_page_number(document):
    for section in document.sections:
        paragraph = section.footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run('第 ')
        set_east_asia_font(run, '微软雅黑')
        run.font.size = Pt(8)
        begin = OxmlElement('w:fldChar')
        begin.set(qn('w:fldCharType'), 'begin')
        instruction = OxmlElement('w:instrText')
        instruction.set(qn('xml:space'), 'preserve')
        instruction.text = ' PAGE '
        end = OxmlElement('w:fldChar')
        end.set(qn('w:fldCharType'), 'end')
        run._r.extend([begin, instruction, end])
        suffix = paragraph.add_run(' 页')
        set_east_asia_font(suffix, '微软雅黑')
        suffix.font.size = Pt(8)


def add_toc(document):
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instruction = OxmlElement('w:instrText')
    instruction.set(qn('xml:space'), 'preserve')
    quote = chr(34)
    instruction.text = ' TOC \\o ' + quote + '1-3' + quote + ' \\h \\z \\u '
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    placeholder = OxmlElement('w:t')
    placeholder.text = '在 Word 中按 Ctrl+A，再按 F9 更新目录。'
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    run._r.extend([begin, instruction, separate, placeholder, end])


def configure_document(document):
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    normal = document.styles['Normal']
    normal.font.name = '微软雅黑'
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    normal.paragraph_format.line_spacing = 1.35

    heading_specs = {
        'Title': (28, BLUE_RGB),
        'Heading 1': (18, BLUE_RGB),
        'Heading 2': (14, BLUE_RGB),
        'Heading 3': (11.5, TEAL_RGB),
    }
    for style_name, (size, color) in heading_specs.items():
        style = document.styles[style_name]
        style.font.name = '微软雅黑'
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    document.core_properties.title = '智能模拟面试系统接口清单及接口文档'
    document.core_properties.subject = 'Flask REST API 接口说明'
    document.core_properties.author = '智能模拟面试系统项目组'
    document.core_properties.keywords = '智能模拟面试, Flask, REST API, 接口文档'


if __name__ == '__main__':
    output = build_document()
    print(f'Generated: {output}')
