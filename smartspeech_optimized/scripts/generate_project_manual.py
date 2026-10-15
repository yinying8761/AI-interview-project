from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path
from typing import Iterable, Sequence
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = PROJECT_ROOT / "智能模拟面试系统介绍、原理、应用及操作教程.docx"
SCREENSHOT_SOURCE = PROJECT_ROOT / "截图.docx"

INK = RGBColor(23, 45, 54)
MUTED = RGBColor(91, 108, 114)
TEAL = RGBColor(13, 122, 114)
CORAL = RGBColor(180, 82, 56)
LIGHT_TEAL = "E7F4F1"
LIGHT_GRAY = "F3F6F6"
LIGHT_AMBER = "FBF1DE"


def set_run_font(run, name: str, size: float, *, bold: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.2)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.widow_control = True

    for style_name, size, color in (
        ("Title", 26, INK),
        ("Heading 1", 17, TEAL),
        ("Heading 2", 13.5, INK),
        ("Heading 3", 11.5, INK),
    ):
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.keep_with_next = True

    document.styles["Heading 1"].paragraph_format.space_before = Pt(14)
    document.styles["Heading 1"].paragraph_format.space_after = Pt(9)
    document.styles["Heading 2"].paragraph_format.space_before = Pt(10)
    document.styles["Heading 2"].paragraph_format.space_after = Pt(6)
    document.styles["Heading 3"].paragraph_format.space_before = Pt(7)
    document.styles["Heading 3"].paragraph_format.space_after = Pt(4)

    properties = document.core_properties
    properties.title = "智能模拟面试系统介绍、原理、应用及操作教程"
    properties.subject = "系统说明与用户操作手册"
    properties.author = "智能体模拟面试系统项目组"
    properties.keywords = "Vue, Flask, 模拟面试, DeepSeek, Whisper, Edge-TTS, MediaPipe"
    properties.comments = "智能体模拟面试系统完整使用教程"

    header = section.header.paragraphs[0]
    header.text = "智能体模拟面试系统"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_run_font(run, "Microsoft YaHei", 9)
        run.font.color.rgb = MUTED

    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    prefix = footer.add_run("智能体模拟面试系统 · ")
    set_run_font(prefix, "Microsoft YaHei", 8.5)
    prefix.font.color.rgb = MUTED
    page_prefix = footer.add_run("第 ")
    set_run_font(page_prefix, "Microsoft YaHei", 8.5)
    page_prefix.font.color.rgb = MUTED
    add_field(footer, "PAGE")
    page_suffix = footer.add_run(" 页")
    set_run_font(page_suffix, "Microsoft YaHei", 8.5)
    page_suffix.font.color.rgb = MUTED

    settings = document.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def add_cover(document: Document) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(54)

    eyebrow = document.add_paragraph()
    eyebrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = eyebrow.add_run("AI INTERVIEW WORKSPACE")
    set_run_font(run, "Consolas", 10, bold=True)
    run.font.color.rgb = TEAL

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("智能体模拟面试系统")
    set_run_font(run, "Microsoft YaHei", 27, bold=True)
    run.font.color.rgb = INK

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("介绍、原理、应用及操作教程")
    set_run_font(run, "Microsoft YaHei", 15, bold=True)
    run.font.color.rgb = CORAL

    description = document.add_paragraph()
    description.alignment = WD_ALIGN_PARAGRAPH.CENTER
    description.paragraph_format.space_before = Pt(18)
    run = description.add_run(
        "Vue 3 + Flask 架构 · 大模型连续追问 · 语音交互 · 实时视频状态分析"
    )
    set_run_font(run, "Microsoft YaHei", 11)
    run.font.color.rgb = MUTED

    info = document.add_table(rows=4, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info.autofit = False
    info.columns[0].width = Cm(4.2)
    info.columns[1].width = Cm(8.8)
    values = (
        ("文档版本", "V1.0"),
        ("适用系统", "智能体模拟面试系统"),
        ("运行环境", "Windows 10/11 · nlp_env · Python 3.10"),
        ("文档日期", date.today().strftime("%Y年%m月%d日")),
    )
    for row, (label, value) in zip(info.rows, values):
        row.cells[0].text = label
        row.cells[1].text = value
        shade_cell(row.cells[0], LIGHT_TEAL)
        for index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for paragraph_run in paragraph.runs:
                    set_run_font(
                        paragraph_run,
                        "Microsoft YaHei",
                        10,
                        bold=index == 0,
                    )
                    paragraph_run.font.color.rgb = INK

    document.add_page_break()


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def add_field(paragraph, instruction: str) -> None:
    """Append a simple Word field such as PAGE or TOC to a paragraph."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction_node = OxmlElement("w:instrText")
    instruction_node.set(qn("xml:space"), "preserve")
    instruction_node.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "正在更新"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction_node, separate, placeholder, end))
    set_run_font(run, "Microsoft YaHei", 8.5)
    run.font.color.rgb = MUTED


def add_toc(document: Document) -> None:
    document.add_heading("目录", level=1)
    add_body(
        document,
        "本文档已插入自动目录。首次用 Word 打开时，如果出现“是否更新域”的提示，"
        "请选择“是”；如未出现，可按 Ctrl+A 全选后按 F9 手动更新。目录中的标题"
        "均与正文的一、二、三级标题关联，适合直接跳转阅读。",
    )
    paragraph = document.add_paragraph()
    add_field(paragraph, r'TOC \o "1-3" \h \z \u')
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(8)


def add_section_break(document: Document) -> None:
    document.add_page_break()


def add_body(
    document: Document,
    text: str,
    *,
    bold_prefix: str | None = None,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0.74)
    if bold_prefix and text.startswith(bold_prefix):
        prefix = paragraph.add_run(bold_prefix)
        set_run_font(prefix, "Microsoft YaHei", 10.5, bold=True)
        prefix.font.color.rgb = INK
        run = paragraph.add_run(text[len(bold_prefix) :])
        set_run_font(run, "Microsoft YaHei", 10.5)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, "Microsoft YaHei", 10.5)
    paragraph.paragraph_format.space_after = Pt(6)


def add_bullets(document: Document, items: Iterable[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = Cm(0.75)
        paragraph.paragraph_format.first_line_indent = Cm(-0.35)
        run = paragraph.add_run(item)
        set_run_font(run, "Microsoft YaHei", 10.5)


def add_steps(document: Document, items: Sequence[str]) -> None:
    for index, item in enumerate(items, start=1):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.35)
        number = paragraph.add_run(f"{index}. ")
        set_run_font(number, "Microsoft YaHei", 10.5, bold=True)
        number.font.color.rgb = TEAL
        run = paragraph.add_run(item)
        set_run_font(run, "Microsoft YaHei", 10.5)


def add_note(document: Document, title: str, text: str) -> None:
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(3.0)
    table.columns[1].width = Cm(12.5)
    shade_cell(table.cell(0, 0), LIGHT_AMBER)
    shade_cell(table.cell(0, 1), "FFFDF8")
    table.cell(0, 0).text = title
    table.cell(0, 1).text = text
    for cell in table.rows[0].cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph_index, paragraph in enumerate(cell.paragraphs):
            paragraph.paragraph_format.space_after = Pt(2)
            for run in paragraph.runs:
                set_run_font(
                    run,
                    "Microsoft YaHei",
                    9.5,
                    bold=(cell is table.cell(0, 0) and paragraph_index == 0),
                )
                run.font.color.rgb = INK


def add_code(document: Document, code: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    shade_cell(table.cell(0, 0), LIGHT_GRAY)
    paragraph = table.cell(0, 0).paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(code)
    set_run_font(run, "Consolas", 9.5)
    run.font.color.rgb = INK


def add_table(
    document: Document,
    headers: Sequence[str],
    rows: Sequence[Sequence[str]],
    widths: Sequence[float] | None = None,
) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = widths is None

    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        shade_cell(cell, LIGHT_TEAL)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            set_run_font(run, "Microsoft YaHei", 9.5, bold=True)
            run.font.color.rgb = INK

    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_properties.append(repeat_header)

    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cells[index].paragraphs[0].runs:
                set_run_font(run, "Microsoft YaHei", 9.2)
                run.font.color.rgb = INK

    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_screenshot(
    document: Document,
    archive: ZipFile,
    image_name: str,
    figure_number: int,
    title: str,
    *,
    width_cm: float = 16.0,
) -> None:
    """Insert one real screenshot from the source DOCX package."""
    document.add_heading(
        f"8.{figure_number + 1} 图 {figure_number}：{title}",
        level=2,
    )
    image_paragraph = document.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.keep_with_next = True
    image_paragraph.paragraph_format.space_after = Pt(3)
    image_paragraph.add_run().add_picture(
        BytesIO(archive.read(f"word/media/{image_name}")),
        width=Cm(width_cm),
    )

    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = True
    caption.paragraph_format.space_after = Pt(7)
    run = caption.add_run(
        f"图 {figure_number}  智能体模拟面试系统真实运行界面（来自“截图.docx”）"
    )
    set_run_font(run, "Microsoft YaHei", 8.5)
    run.font.color.rgb = MUTED


def build_legacy_manual() -> Path:
    document = Document()
    configure_document(document)
    add_cover(document)

    document.add_heading("1. 系统介绍", level=1)
    add_body(
        document,
        "智能体模拟面试系统是一套面向求职训练、课程实践和面试能力评估的本地化应用。"
        "系统把连续追问、语音交互和视频状态分析放在同一条面试链路中，"
        "让候选人不仅能获得问题反馈，还能看到表达过程、表情变化和注意力状态。",
    )
    add_body(
        document,
        "系统名称：智能体模拟面试系统。系统采用 Vue 3 构建前端工作台，"
        "使用 Flask 提供后端接口，并通过 Waitress 提供稳定的生产运行服务。"
        "所有功能通过浏览器访问，核心数据默认保存在本机。",
        bold_prefix="系统名称：",
    )

    document.add_heading("1.1 设计目标", level=2)
    add_bullets(
        document,
        (
            "让用户在真实摄像头画面常驻的前提下完成整场模拟面试。",
            "让每个问题能够结合岗位、难度和上一轮回答自然追问。",
            "让语音回答、文字回答和实时视觉状态形成统一的数据链路。",
            "让面试结果可量化、可复盘，并能长期查看历史报告。",
            "让系统在普通 Windows 电脑上保持清晰、稳定和流畅。",
        ),
    )

    document.add_heading("1.2 目标用户", level=2)
    add_table(
        document,
        ("用户类型", "主要需求", "系统价值"),
        (
            ("求职者", "模拟真实面试、发现表达短板", "连续追问与综合报告"),
            ("在校学生", "课程实训、面试入门练习", "低门槛、本地运行"),
            ("教师与教练", "设计演示和训练流程", "完整记录与可复盘评分"),
            ("开发学习者", "了解多模态应用架构", "Vue + Flask 模块化示例"),
        ),
        widths=(3.2, 6.2, 6.2),
    )

    document.add_heading("2. 核心功能", level=1)
    add_table(
        document,
        ("功能模块", "实现方式", "用户可见结果"),
        (
            (
                "身份认证",
                "账号密码、PBKDF2 密码哈希、人脸特征向量",
                "密码登录、人脸登录、账号注册",
            ),
            (
                "面试智能体",
                "DeepSeek Chat Completions API",
                "开场白、连续追问、综合评估",
            ),
            (
                "语音听写",
                "浏览器录音、Whisper 转写",
                "语音回答自动转为文字",
            ),
            (
                "语音合成",
                "Edge-TTS，Windows SAPI 回退",
                "面试官自动或手动朗读问题",
            ),
            (
                "实时视频分析",
                "MediaPipe 表情、人脸和头部姿态分析",
                "人脸、表情、视线、注意力、FPS、延迟",
            ),
            (
                "报告系统",
                "本地 JSON 报告存储与 Vue 报告页",
                "评分总览、逐题复盘、历史报告",
            ),
        ),
        widths=(3.2, 6.1, 6.3),
    )

    document.add_heading("3. 系统原理", level=1)
    document.add_heading("3.1 总体架构", level=2)
    add_body(
        document,
        "浏览器负责摄像头采集、录音、页面交互和实时展示；Flask 负责业务接口、"
        "模型调用、会话管理、视觉计算和报告生成。Vue 构建后的静态资源由 Flask "
        "统一托管，因此生产模式只需要访问一个端口。",
    )
    add_code(
        document,
        "浏览器 Vue 工作台\n"
        "  -> Flask API\n"
        "      -> DeepSeek 面试智能体\n"
        "      -> Whisper 语音听写\n"
        "      -> Edge-TTS 语音合成\n"
        "      -> MediaPipe 视频状态分析\n"
        "      -> 本地会话与报告存储",
    )

    document.add_heading("3.2 连续追问原理", level=2)
    add_body(
        document,
        "创建面试会话时，系统记录目标岗位、难度、音色和画面采样频率。"
        "生成第一题后，每一轮用户回答都会和最近对话上下文一起发送给面试智能体。"
        "模型据此生成下一道追问题，问题会逐步贴近候选人的经历、技术细节和回答缺口。",
    )
    add_body(
        document,
        "会话对象只保存在后端内存中，用户身份和会话所有权均经过校验。"
        "结束面试后，系统将问答记录交给模型进行综合评估，并把评估结果与视频聚合指标融合。",
    )

    document.add_heading("3.3 视频分析原理", level=2)
    add_body(
        document,
        "前端以页面设定频率从摄像头视频元素截取最新画面，并通过接口提交给后端。"
        "后端只保留一个待处理帧，新帧到来时替换旧帧，避免低速模型造成请求堆积。"
        "媒体处理完成后，前端读取聚合指标并更新当前状态、趋势和延迟数据。",
    )
    add_table(
        document,
        ("指标", "含义", "用途"),
        (
            ("人脸检测", "当前画面是否识别到人脸", "判断是否正对摄像头"),
            ("表情状态", "微笑、专注、紧张等面部状态", "形成表情得分与变化次数"),
            ("注意力状态", "视线居中、头部姿态和离屏情况", "形成注意力得分与分散次数"),
            ("有效 FPS", "后端实际完成分析的帧率", "判断视觉链路是否流畅"),
            ("分析延迟", "单帧从开始到完成处理的时间", "触发前端自适应降频"),
            ("丢帧数", "被最新帧替换掉的待处理帧", "衡量实时性与负载"),
        ),
        widths=(3.4, 6.1, 6.1),
    )

    document.add_heading("3.4 评分原理", level=2)
    add_body(
        document,
        "问答部分由大模型从逻辑、专业度、岗位匹配度等维度综合评分。"
        "视频部分由表情得分和注意力得分等权合成，再与问答得分按固定权重计算最终得分。",
    )
    add_code(
        document,
        "最终得分 = 问答得分 × 80% + 视频综合得分 × 20%\n"
        "视频综合得分 = 表情得分 × 50% + 注意力得分 × 50%",
    )
    add_note(
        document,
        "评分原则",
        "有效视频样本不足时，最终得分保持为空。系统不会重新分配权重，"
        "也不会用少量偶然画面生成虚高分数。",
    )

    document.add_heading("4. 应用场景", level=1)
    add_table(
        document,
        ("场景", "使用方式", "建议配置"),
        (
            (
                "个人求职训练",
                "按目标岗位创建面试，完成多轮回答并复盘报告",
                "中级难度，2-3 帧/秒",
            ),
            (
                "校园课程实训",
                "演示多模态 AI 应用的完整前后端链路",
                "提前准备好模型和环境",
            ),
            (
                "面试表达训练",
                "重点观察表情、视线和注意力趋势",
                "光线充足，保持正对摄像头",
            ),
            (
                "技术面试演练",
                "选择对应技术岗位并连续追问项目细节",
                "高级或专家难度",
            ),
            (
                "效果演示",
                "展示实时视频、语音问答和报告生成",
                "使用稳定网络与清晰摄像头",
            ),
        ),
        widths=(3.2, 8.0, 4.5),
    )

    document.add_heading("5. 安装与启动", level=1)
    document.add_heading("5.1 环境准备", level=2)
    add_bullets(
        document,
        (
            "Windows 10 或 Windows 11。",
            "Python 3.10，本项目使用 nlp_env 虚拟环境。",
            "Node.js 20 或更高版本，以及 npm。",
            "Chrome 或 Edge 浏览器。",
            "可用的 DeepSeek API 密钥，用于连续追问与综合评估。",
        ),
    )
    add_code(
        document,
        r"D:\Work_Software\Anaconda3\envs\nlp_env\python.exe",
    )

    document.add_heading("5.2 一键启动", level=2)
    add_steps(
        document,
        (
            "打开项目根目录。",
            "确认 .env 中已填写 DEEPSEEK_API_KEY。",
            "双击 start_system.bat。",
            "等待脚本检查依赖并完成 Vue 构建。",
            "浏览器自动打开 http://127.0.0.1:8765。",
        ),
    )
    add_note(
        document,
        "首次启动",
        "首次安装前端依赖或加载 Whisper 模型时耗时较长。"
        "后续启动会复用 node_modules 和本机模型缓存。",
    )

    document.add_heading("5.3 手动启动", level=2)
    add_body(document, "生产模式启动命令：")
    add_code(
        document,
        r"""cd "D:\Project_Code_2026\tyut_comprehensive_pratice\smartspeech_optimized_new"
& "D:\Work_Software\Anaconda3\envs\nlp_env\python.exe" -m pip install -r requirements.txt
cd frontend
npm ci
npm run build
cd ..
& "D:\Work_Software\Anaconda3\envs\nlp_env\python.exe" -m backend.run""",
    )
    add_body(document, "开发模式需要两个终端：")
    add_code(
        document,
        r"""# 终端 1：后端
& "D:\Work_Software\Anaconda3\envs\nlp_env\python.exe" -m backend.run

# 终端 2：前端
cd frontend
npm run dev""",
    )
    add_body(document, "开发页面地址为 http://127.0.0.1:5173。")

    document.add_heading("6. 操作教程", level=1)
    document.add_heading("6.1 注册与登录", level=2)
    add_steps(
        document,
        (
            "进入登录页，选择“注册”。",
            "填写用户名、显示名称、密码和确认密码。",
            "如需使用人脸登录，勾选“注册后录入人脸”，拍照后完成绑定。",
            "已有账号可切换到“密码登录”或“人脸登录”。",
            "登录成功后进入面试工作台。",
        ),
    )
    add_note(
        document,
        "隐私提示",
        "人脸注册只保存特征向量，不保存原始照片。摄像头画面不会写入磁盘。",
    )

    document.add_heading("6.2 创建面试", level=2)
    add_steps(
        document,
        (
            "选择目标岗位，也可以填写自定义岗位。",
            "选择初级、中级、高级或专家难度。",
            "选择面试官音色，并决定是否开启自动朗读。",
            "根据电脑性能选择 2、3 或 4 帧/秒。普通电脑建议选择 2-3 帧/秒。",
            "点击“开始面试并打开摄像头”，允许浏览器访问摄像头。",
        ),
    )

    document.add_heading("6.3 完成面试", level=2)
    add_steps(
        document,
        (
            "等待面试官生成开场白和第一题。",
            "在回答框中输入文字，或点击“语音回答”进行录音。",
            "语音停止后系统自动转写并提交，面试官生成下一道问题。",
            "观察左侧视频画面以及右侧人脸、注意力、FPS、延迟和趋势指标。",
            "完成至少一轮回答后，点击“结束面试”。",
            "等待综合评估完成，系统自动打开报告详情页。",
        ),
    )
    add_note(
        document,
        "流畅度建议",
        "如果页面卡顿，先降低画面采样频率，并关闭其他占用摄像头、CPU "
        "或网络的程序。系统会根据处理延迟自动调整采集间隔。",
    )

    document.add_heading("6.4 查看报告", level=2)
    add_bullets(
        document,
        (
            "报告顶部展示最终得分、问答得分、表情得分、注意力得分和视频综合得分。",
            "综合评估区展示总体总结、维度分数、表现亮点、待改进项和训练建议。",
            "逐题复盘区展示每道问题的回答评估和改进建议。",
            "右侧展示视频样本统计、状态变化时间线和隐私说明。",
            "点击“打印 / 导出 PDF”可使用浏览器打印功能输出报告。",
            "在顶部导航进入“历史报告”，可以查看账号下的历次面试。",
        ),
    )

    document.add_heading("7. 接口与数据结构", level=1)
    add_table(
        document,
        ("接口", "方法", "作用"),
        (
            ("/api/health", "GET", "健康检查与能力探测"),
            ("/api/auth/register", "POST", "注册并创建登录会话"),
            ("/api/auth/login", "POST", "账号密码登录"),
            ("/api/auth/face-register", "POST", "录入或更新人脸特征"),
            ("/api/auth/face-login", "POST", "通过人脸特征匹配登录"),
            ("/api/interviews", "POST", "创建面试会话"),
            ("/api/interviews/<id>/opening", "POST", "生成开场和第一题"),
            ("/api/interviews/<id>/frames", "POST", "提交最新视频帧"),
            ("/api/interviews/<id>/metrics", "GET", "读取实时聚合指标"),
            ("/api/interviews/<id>/messages", "POST", "提交回答并生成追问"),
            ("/api/interviews/<id>/finish", "POST", "结束面试并生成报告"),
            ("/api/reports", "GET", "获取历史报告列表"),
            ("/api/reports/<id>", "GET", "获取报告详情"),
            ("/api/media/transcribe", "POST", "Whisper 语音转写"),
            ("/api/media/tts", "POST", "Edge-TTS 语音合成"),
        ),
        widths=(7.3, 2.0, 6.3),
    )

    document.add_heading("8. 隐私与数据安全", level=1)
    add_table(
        document,
        ("数据类型", "是否保存", "位置或处理方式"),
        (
            ("原始视频帧", "否", "内存中实时分析后丢弃"),
            ("人脸原始照片", "否", "提取特征后销毁"),
            ("人脸特征向量", "是", "data/auth/faces/"),
            ("回答录音", "否", "转写完成后立即删除"),
            ("问答文本", "是", "写入本地面试报告"),
            ("评分与聚合指标", "是", "runtime/reports/"),
            ("账号密码", "哈希保存", "PBKDF2 哈希写入 users.json"),
            ("会话签名密钥", "是", "data/.flask_secret"),
        ),
        widths=(3.7, 3.0, 8.9),
    )
    add_note(
        document,
        "部署建议",
        "如需对局域网或互联网开放，应启用 HTTPS，设置随机的 FLASK_SECRET_KEY，"
        "并配置 SESSION_COOKIE_SECURE=1。不要把账号、人脸向量、密钥和报告上传到公开仓库。",
    )

    document.add_heading("9. 性能说明", level=1)
    add_body(
        document,
        "系统把视频画面显示与分析分离。即使视觉模型仍在处理上一帧，"
        "浏览器中的视频画面也会持续显示。后端队列最多保留一个待处理帧，"
        "新画面会替换旧画面，因此不会因为模型速度较慢而无限堆积请求。",
    )
    add_table(
        document,
        ("优化机制", "作用"),
        (
            ("最新帧替换", "始终分析最新状态，减少过期画面"),
            ("自适应降频", "延迟升高时自动扩大采集间隔"),
            ("页面隐藏暂停", "切换到其他标签页时停止无效采集"),
            ("视觉线程隔离", "媒体分析不阻塞问答和报告接口"),
            ("前端请求互斥", "同一时间只提交一帧，避免并发堆积"),
            ("构建产物合并", "生产环境由 Flask 统一托管，减少部署复杂度"),
        ),
        widths=(5.0, 10.6),
    )

    document.add_heading("10. 故障排查", level=1)
    add_table(
        document,
        ("现象", "可能原因", "处理方式"),
        (
            (
                "摄像头不可用",
                "权限被拒绝或设备被占用",
                "允许浏览器访问，关闭占用摄像头的软件后重试",
            ),
            (
                "画面滞后又恢复",
                "CPU 或媒体模型负载较高",
                "选择 2 帧/秒，关闭其他高占用程序",
            ),
            (
                "第一题生成慢",
                "DeepSeek 网络或密钥问题",
                "检查 API Key、网络和接口地址",
            ),
            (
                "语音识别首次很慢",
                "Whisper 正在加载或下载模型",
                "等待首次加载完成，后续会复用缓存",
            ),
            (
                "页面没有声音",
                "浏览器阻止自动播放",
                "点击面试官消息旁的播放按钮",
            ),
            (
                "端口 8765 被占用",
                "已有服务或重复启动",
                "打开已有地址，或修改端口配置",
            ),
            (
                "前端显示尚未构建",
                "dist 目录缺失",
                "执行 npm run build 或运行一键启动脚本",
            ),
        ),
        widths=(4.0, 5.2, 6.4),
    )

    document.add_heading("11. 测试与维护", level=1)
    add_body(document, "运行全部后端单元测试：")
    add_code(
        document,
        r"""& "D:\Work_Software\Anaconda3\envs\nlp_env\python.exe" -m unittest discover -s backend\tests -v""",
    )
    add_body(document, "检查 Python 语法和依赖：")
    add_code(
        document,
        r"""& "D:\Work_Software\Anaconda3\envs\nlp_env\python.exe" -m compileall backend
& "D:\Work_Software\Anaconda3\envs\nlp_env\python.exe" -m pip check""",
    )
    add_body(document, "重新生成本操作手册：")
    add_code(
        document,
        r"""& "D:\Work_Software\Anaconda3\envs\nlp_env\python.exe" scripts\generate_project_manual.py""",
    )

    document.add_heading("12. 项目目录", level=1)
    add_code(
        document,
        "backend/app/api/          Flask 接口\n"
        "backend/app/services/     认证、面试、语音和视觉服务\n"
        "backend/app/storage/      会话与报告存储\n"
        "backend/tests/            后端单元测试\n"
        "frontend/src/components/  Vue 工作台组件\n"
        "frontend/src/pages/       登录、面试和报告页面\n"
        "frontend/src/stores/      Pinia 状态\n"
        "data/                     账号、人脸向量和密钥\n"
        "runtime/                  临时文件和面试报告\n"
        "scripts/                  维护脚本",
    )

    document.add_heading("13. 结语", level=1)
    add_body(
        document,
        "智能体模拟面试系统把大模型问答、语音交互和实时视觉分析整合为一个可操作的"
        "面试工作台。建议先以较低画面采样频率完成一次完整流程，再根据电脑性能和"
        "训练重点调整岗位、难度与采样设置。",
    )
    add_note(
        document,
        "使用提醒",
        "模拟评分用于训练参考，不等同于真实招聘结论。请结合报告中的具体建议，"
        "持续练习表达结构、专业深度和面试状态。",
    )

    document.save(OUTPUT_PATH)
    return OUTPUT_PATH


def add_reading_guide(document: Document) -> None:
    document.add_heading("阅读导航与五句话速览", level=1)
    add_body(
        document,
        "这是一份面向学生、答辩人员和项目使用者的白话版说明。它不只告诉你“点哪里”，"
        "还会解释界面背后发生了什么、为什么采用某项技术、数据保存在哪里，以及若要做"
        "二次开发或训练，应该从哪一部分入手。第一次接触本项目时，建议按本文档顺序阅读；"
        "如果只想尽快运行，可直接跳到第 7、8 章。",
    )
    add_table(
        document,
        ("我现在想做什么", "建议重点阅读", "读完后能得到什么"),
        (
            ("了解项目价值并准备答辩", "第 1、2、3 章", "项目定位、功能地图、总体架构"),
            ("理解模型和算法原理", "第 4、5 章", "各模型如何工作、数据从哪里来"),
            ("看懂代码并做二次开发", "第 6、11 章", "目录职责、调用链、扩展边界"),
            ("把系统在自己电脑上跑起来", "第 7、10 章", "安装步骤、配置与故障排查"),
            ("完成一场模拟面试", "第 8 章", "从注册到报告的完整操作流程"),
            ("训练或微调模型", "第 5、9 章", "当前能否训练、需要哪些数据与指标"),
        ),
        widths=(4.8, 4.7, 6.2),
    )

    document.add_heading("先记住五句话", level=2)
    add_bullets(
        document,
        (
            "第一句：这是一套本地运行的面试训练工具，浏览器负责交互，Flask 负责业务、"
            "模型调度和报告，Vue 构建后的页面由 Flask 统一提供。",
            "第二句：面试问题由 DeepSeek 根据岗位、难度和最近回答连续生成，只发送文字"
            "问答记录，不发送原始视频、人脸和录音。",
            "第三句：Whisper 把语音回答转成文字，Edge-TTS 把面试官问题读出来；语音合成"
            "失败时会自动使用 Windows SAPI 保持流程可用。",
            "第四句：人脸、表情和注意力分析主要在本机执行。MTCNN + InceptionResnetV1 "
            "负责人脸身份匹配，MediaPipe Face Mesh 负责表情与头部/视线状态。",
            "第五句：最终得分以问答为主、过程状态为辅，公式是“问答 80% + 视频综合 20%”；"
            "视频样本不足时宁可为空，也不把权重临时分给问答。",
        ),
    )
    add_note(
        document,
        "阅读标记",
        "带“一句话理解”的内容用于快速建立直觉；公式和阈值用于理解系统如何计算；"
        "“系统落点”说明代码实际怎么做；“注意边界”提醒不要把训练结果理解成招聘结论。",
    )
    add_section_break(document)
    add_toc(document)
    add_section_break(document)


def add_chapter_1_project(document: Document) -> None:
    document.add_heading("第 1 章 项目介绍", level=1)

    document.add_heading("1.1 项目背景：为什么需要它", level=2)
    add_body(
        document,
        "真实面试同时考验知识、表达、临场状态和时间控制。学生平时也许会背题，却很少"
        "有人连续追问项目细节；即使录下视频，也很难从几十分钟录像中指出“哪一句话没有"
        "数据支撑”“哪些时段没有看镜头”“哪道题回答得过于笼统”。传统练习要么缺少即时"
        "互动，要么反馈零散，难以形成可重复、可对比的训练闭环。",
    )
    add_body(
        document,
        "智能体模拟面试系统把一次练习拆成四条可观察链路：面试官是否问到关键点、回答"
        "是否被准确保留、表达过程是否稳定、结束后能否给出结构化建议。它不替代真人面试，"
        "而是把原本一次性结束的练习变成可以复盘的训练数据。",
    )

    document.add_heading("1.2 项目定位：它到底是什么", level=2)
    add_table(
        document,
        ("维度", "当前定位"),
        (
            ("产品形态", "浏览器中运行的本地 AI 面试训练平台"),
            ("核心方法", "大模型连续追问 + 本地语音识别与合成 + 本地视觉状态分析"),
            ("前端技术", "Vue 3、Vite、Pinia、Vue Router、lucide-vue-next"),
            ("后端技术", "Flask、Waitress、requests、OpenCV、MediaPipe、PyTorch"),
            ("数据形态", "账号与人脸向量、临时音频、问答文本、聚合指标、JSON 报告"),
            ("默认地址", "http://127.0.0.1:8765"),
            ("适用目标", "学生学习、课程实训、项目演示、求职前重复练习"),
        ),
        widths=(3.2, 12.5),
    )

    document.add_heading("1.3 项目目标", level=2)
    add_bullets(
        document,
        (
            "让一次模拟面试从“随机问几题”变成有岗位、有难度、有连续上下文的训练。",
            "让回答既可以用键盘输入，也可以录音后自动转写，兼顾效率与真实感。",
            "让训练过程可见：摄像头、分析状态、有效帧率、延迟和趋势都能实时查看。",
            "让结果可解释：问答分、表情分、注意力分、视频综合分和最终分分别展示。",
            "让报告可行动：给出亮点、短板、逐题建议和后续训练路径，而不是只有一个数字。",
            "让普通 Windows 电脑也能运行：默认使用 CPU、低采样率和轻量本地分析。",
        ),
    )

    document.add_heading("1.4 目标用户", level=2)
    add_table(
        document,
        ("用户", "典型诉求", "使用方式"),
        (
            ("在校学生", "首次练习技术面或综合面", "选择岗位与初级/中级难度，先完成完整流程"),
            ("求职者", "针对岗位反复训练与对比", "用同一岗位多次面试，比较历史报告"),
            ("课程学生", "理解多模态 AI 工程", "结合代码结构与调用链做实验和答辩"),
            ("教师或教练", "展示面试训练闭环", "演示语音、视觉、大模型和报告协作"),
            ("二次开发者", "替换模型或扩展评分", "从服务层和 API 层入手，保留现有接口"),
        ),
        widths=(3.2, 5.6, 6.9),
    )

    document.add_heading("1.5 典型应用场景", level=2)
    add_body(
        document,
        "个人训练场景中，用户可以先选择 Python 后端开发工程师和中级难度，用文字完成"
        "第一轮，再用语音完成第二轮，对比自己的表达清晰度和术语错误。课堂演示场景中，"
        "教师可以重点展示同一份回答如何被追问、浏览器如何采样、后端如何聚合并输出报告。"
        "项目答辩场景中，则可以用系统本身说明“为什么采用本地模型与云端 API 的混合路线”。",
    )
    add_table(
        document,
        ("场景", "推荐设置", "观察重点"),
        (
            ("首次体验", "中级、3 帧/秒、自动朗读", "完整跑通流程并确认设备和网络"),
            ("表达训练", "中级、2-3 帧/秒、语音回答", "停顿、语速转写、微笑和目光稳定性"),
            ("技术深挖", "高级或专家、文字回答", "追问是否跟随项目细节与技术取舍"),
            ("公开演示", "3 帧/秒、提前预载模型", "响应速度、画面分析、报告完整性"),
            ("性能较弱电脑", "2 帧/秒、关闭其他占用程序", "有效 FPS、延迟、是否丢帧"),
        ),
        widths=(3.4, 5.7, 6.6),
    )

    document.add_heading("1.6 系统边界与使用提醒", level=2)
    add_bullets(
        document,
        (
            "它是训练和教学工具，不是招聘决策系统，不能仅凭一次得分判断一个人的能力。",
            "表情和注意力是过程状态信号，不等于真实情绪或心理判断，也没有医学诊断用途。",
            "当前人脸登录主要验证“是不是已登记的人”，不应作为高安全等级的支付级认证。",
            "大模型评分具有随机性，评分平均偏低或偏高并不自动代表模型出现故障。",
            "Edge-TTS 与 DeepSeek 需要网络；断网时语音回退、备用首题可维持部分体验，"
            "但连续追问和最终综合评估依赖 API。",
        ),
    )
    add_note(
        document,
        "答辩建议",
        "介绍项目时先说明“解决什么问题”，再说明“为什么这样组合技术”，最后演示一条"
        "完整链路。不要把重点放在模型名称堆砌上，而要讲清本地处理与云端调用的边界。",
    )


def add_chapter_2_features(document: Document) -> None:
    add_section_break(document)
    document.add_heading("第 2 章 功能模块总览", level=1)
    document.add_heading("2.1 功能地图", level=2)
    add_body(
        document,
        "系统不是几个独立演示页面的拼接，而是一条完整业务链：用户登录后创建面试，"
        "浏览器同时处理文字/语音/视频输入，后端分别调用大模型、Whisper 和视觉分析，"
        "最后将这些结果合并为报告。阅读这一章时，可以把它理解成七个相互协作的模块。",
    )
    add_table(
        document,
        ("模块", "核心职责", "关键输出"),
        (
            ("账号与认证", "注册、密码登录、人脸绑定、人脸登录、会话管理", "当前登录用户"),
            ("面试配置", "岗位、难度、音色、自动朗读、画面采样率", "一场面试的会话参数"),
            ("面试智能体", "生成开场、第一题、连续追问和最终评估", "问题、评分与建议"),
            ("语音交互", "录音、Whisper 转写、Edge-TTS 播报与本地回退", "可提交的回答文本、可播放音频"),
            ("实时视觉", "采集最新帧，分析人脸、表情、头部姿态、视线和覆盖率", "实时状态与聚合指标"),
            ("报告中心", "保存历史、展示分数、逐题复盘、趋势和事件", "可打印或导出 PDF 的报告"),
            ("运维与隐私", "健康检查、能力探测、配置、临时文件清理、本地存储", "可诊断、可维护的运行状态"),
        ),
        widths=(3.1, 8.1, 4.6),
    )

    document.add_heading("2.2 一次完整使用旅程", level=2)
    add_steps(
        document,
        (
            "打开系统，注册账号；可以同时录入一张清晰的人脸照片，也可以稍后再绑定。",
            "登录后进入面试工作台，选择目标岗位、难度、音色和画面采样频率。",
            "点击开始并授权摄像头，系统创建面试会话，生成开场白与第一题。",
            "用户输入文字或点击语音回答；语音完成后由 Whisper 转写并提交。",
            "DeepSeek 结合岗位、难度和最近 20 条有效上下文生成下一道追问。",
            "面试期间，浏览器持续提交最新画面；后端分析人脸、表情与注意力并返回指标。",
            "结束面试后，大模型评估完整问答文本，本地汇总视觉指标，最终计算综合得分。",
            "报告自动打开；用户可回看逐题建议、状态时间线，也可进入历史报告再次复盘。",
        ),
    )

    document.add_heading("2.3 账号、密码与人脸认证", level=2)
    add_body(
        document,
        "注册时，用户名会被标准化后再比较，避免大小写或全角半角造成重复账号。密码"
        "不是明文保存，而是使用随机盐和 PBKDF2-SHA256 迭代 310000 次后保存摘要。人脸"
        "绑定则是拍摄一张清晰正面照片，系统提取 512 维特征向量并 L2 归一化，只把向量"
        "保存为 data/auth/faces/<用户编号>.npy，不保存这张原始照片。",
    )
    add_body(
        document,
        "登录有两种路径。密码登录直接验证密码摘要；人脸登录会重新提取当前人脸向量，"
        "再与本地已登记向量逐一计算欧氏距离，距离严格小于 FACE_MATCH_THRESHOLD（默认"
        " 1.0）才视为匹配。阈值越小越严格，阈值越大越容易通过，但误认风险也会增加。",
    )
    add_note(
        document,
        "安全边界",
        "系统使用本地 JSON 文件保存用户元数据和哈希，不是面向互联网的高并发账号系统。"
        "若公开部署，必须启用 HTTPS、随机会话密钥、访问控制和更完整的审计机制。",
    )

    document.add_heading("2.4 智能问答与连续追问", level=2)
    add_body(
        document,
        "系统为面试官和评估员定义了两套不同提示词。面试官模式强调“每轮只问一个问题、"
        "优先追问上一轮、回答含糊时继续深挖、过程中不直接讲标准答案”；评估员模式要求"
        "仅根据完整文字记录，从专业能力、逻辑表达、沟通清晰度和岗位匹配度给出结构化"
        " JSON。这样可以把“主持面试”和“事后评分”分开，减少角色混杂。",
    )
    add_table(
        document,
        ("环节", "发送给 DeepSeek 的内容", "温度", "结果"),
        (
            ("开场与第一题", "岗位、难度、面试官规则", "0.8", "开场白 + 第一题"),
            ("连续追问", "规则 + 最近 20 条有效上下文 + 本轮回答", "0.8", "下一道追问"),
            ("最终评估", "岗位、难度、完整面试文字记录", "0.2", "JSON 评分与建议"),
        ),
        widths=(3.2, 7.4, 1.8, 3.4),
    )
    add_body(
        document,
        "追问温度 0.8 让问题更有变化，避免每场面试完全一样；评估温度 0.2 让结果更稳定。"
        "“最近 20 条”是默认 CHAT_CONTEXT_MESSAGES，不是整场记录的上限：最终评估仍会使用"
        "完整文字记录。上下文越长越连贯，但请求更慢，也消耗更多 API token。",
    )

    document.add_heading("2.5 语音识别与语音合成", level=2)
    add_body(
        document,
        "语音回答链路是“浏览器录音 -> 上传临时文件 -> Whisper 本地转写 -> 删除临时文件"
        " -> 文字提交给面试智能体”。首次调用需要下载或加载 Whisper 模型，CPU 上 small "
        "规格通常能兼顾中文识别效果和速度；如果电脑较弱，可改为 base 或 tiny。",
    )
    add_body(
        document,
        "面试官播报链路是“问题文本 -> Edge-TTS 在线合成 -> 浏览器播放”。Edge-TTS 若连续"
        "失败，Windows 环境会自动退回 SAPI 生成 WAV。缓存在 runtime/tts，超过 24 小时的"
        "文件会被清理。浏览器可能阻止自动播放，此时点击消息旁的小喇叭即可播放。",
    )

    document.add_heading("2.6 实时视频与状态看板", level=2)
    add_body(
        document,
        "浏览器按用户选择的 2、3 或 4 帧/秒截取当前摄像头画面，最长边约压缩到 448 像素"
        "后上传。后端会继续把宽度限制在 640 像素以内。画面只用于即时分析，不写入报告。"
        "实时看板显示人脸是否存在、当前表情、视线是否居中、头部角度、注意力状态、有效"
        " FPS、分析延迟和最近趋势。",
    )
    add_table(
        document,
        ("看板指标", "白话解释", "可采取的行动"),
        (
            ("人脸检测", "当前帧有没有稳定找到一张脸", "检查光线、镜头遮挡和坐姿"),
            ("表情状态", "几何特征相对个人基线变化后的统一状态", "不要刻意长时间微笑，保持自然"),
            ("视线/姿态", "眼睛是否看向中心、头部是否明显偏转或俯仰", "尽量看向摄像头附近"),
            ("有效 FPS", "后端真正完成分析的频率", "低于目标时降低采样率或关闭高占用软件"),
            ("分析延迟", "一帧从进入后端到分析完成的耗时", "延迟升高时前端会自动降频"),
            ("丢帧数", "新画面到达时被替换掉的旧待处理帧", "有丢帧不等于故障，它用于保持实时性"),
        ),
        widths=(3.2, 7.1, 5.3),
    )

    document.add_heading("2.7 报告、历史与复盘", level=2)
    add_body(
        document,
        "报告把两类结果放在一起：DeepSeek 根据文字回答生成问答总分、能力维度、亮点、"
        "短板、逐题建议和后续训练步骤；本地视觉分析生成表情分布、注意力状态变化、"
        "离屏事件、人脸覆盖率、眨眼次数、头部角度和趋势快照。最终分数按固定公式融合，"
        "页面同时保留原始单项分。",
    )
    add_bullets(
        document,
        (
            "历史报告按当前账号隔离，只展示该用户自己的面试记录。",
            "报告保存在 runtime/reports/<用户编号>/<面试编号>.json，适合离线查看和备份。",
            "页面支持浏览器打印或导出 PDF，但其本质是网页打印，不是服务端生成 PDF。",
            "删除面试时会同时清理会话和对应报告，删除后不可撤销。",
        ),
    )

    document.add_heading("2.8 模块怎样协作", level=2)
    add_code(
        document,
        "登录/注册\n"
        "  -> 创建面试会话\n"
        "      -> 生成开场与第一题\n"
        "      -> [文字回答 / 录音转写] -> DeepSeek 追问\n"
        "      -> [摄像头最新帧] -> 视觉队列 -> 表情 + 注意力指标\n"
        "      -> 结束面试\n"
        "          -> DeepSeek 文字评估 + 本地视频汇总\n"
        "          -> 评分融合 -> JSON 报告 -> Vue 报告页",
    )
    add_note(
        document,
        "模块化价值",
        "文字问答、语音和视觉并非串行阻塞。视觉在独立工作线程中处理，大模型调用与其他"
        "接口分离；即使某一帧分析较慢，也不会让页面视频停止播放。",
    )


def add_chapter_3_architecture(document: Document) -> None:
    add_section_break(document)
    document.add_heading("第 3 章 总体架构与技术路线", level=1)
    document.add_heading("3.1 一张图看懂总体架构", level=2)
    add_code(
        document,
        "┌────────────────────── 浏览器：Vue 3 工作台 ──────────────────────┐\n"
        "│ 登录页 | 面试配置 | 摄像头与录音 | 问答与实时指标 | 报告页      │\n"
        "└───────────────────────────┬──────────────────────────────────────┘\n"
        "                            │ HTTP / JSON / Multipart\n"
        "┌───────────────────────────▼──────────────────────────────────────┐\n"
        "│ Flask API + Waitress 生产服务                                    │\n"
        "│ 认证 | 面试会话 | DeepSeek | Whisper | Edge-TTS | 视觉 | 报告   │\n"
        "└───────────────┬─────────────────────────────┬────────────────────┘\n"
        "                │                             │\n"
        "        本地模型与数据                  云端 API（按需）\n"
        "        MTCNN / FaceNet / MediaPipe      DeepSeek / Edge-TTS\n"
        "        OpenCV / 本地 JSON 报告          只发送问答文字或问题文本",
    )
    add_body(
        document,
        "浏览器主要负责“采集和展示”，后端负责“判断和编排”。摄像头视频流在浏览器中"
        "持续显示，分析用的压缩帧按设定频率发往后端；录音在浏览器完成后上传，返回的是"
        "文字；大模型只接触文字，因此视频、人脸原图和语音文件都不会直接进入云端。",
    )

    document.add_heading("3.2 技术路线：预训练模型 + 规则融合 + API 大模型", level=2)
    add_body(
        document,
        "本项目没有从零训练所有模型，而是采用更符合课程项目和普通电脑条件的三段式"
        "路线。第一段使用成熟预训练模型完成人脸特征提取、语音识别和面部关键点检测；"
        "第二段用透明、可调参的几何规则把关键点变成表情和注意力状态；第三段通过 API "
        "调用 DeepSeek 完成自然语言面试与综合评估。",
    )
    add_table(
        document,
        ("路线", "解决的问题", "优点", "局限"),
        (
            ("预训练本地模型", "看懂人脸、听写语音、定位关键点", "开箱可用、隐私可控、无需自建训练集", "模型文件较大，CPU 推理仍需时间"),
            ("几何规则融合", "把关键点转换为可解释状态和分数", "阈值直观、便于答辩与调参", "不如大规模深度模型覆盖丰富表情"),
            ("云端大模型 API", "连续追问与灵活语言评估", "语言能力强、接入快、无需本地显卡", "需要网络、密钥和调用成本"),
            ("本地结构化存储", "保存账号、向量和报告", "离线可查、部署简单", "不适合作为大型分布式数据库"),
        ),
        widths=(3.3, 4.2, 4.6, 4.1),
    )
    add_note(
        document,
        "为什么不全放云端",
        "原始视频连续上传会带来带宽、延迟和隐私压力。系统只在本地提取状态指标，把"
        "真正需要大模型理解的问答文字送到云端，这是本地与云端的合理分工。",
    )

    document.add_heading("3.3 前端技术路线", level=2)
    add_table(
        document,
        ("技术", "作用", "为什么使用"),
        (
            ("Vue 3", "构建组件化页面和响应式状态", "语法直接，适合快速开发单页应用"),
            ("Vite", "开发服务器与生产构建", "启动快，配置简单，构建产物可交给 Flask"),
            ("Pinia", "保存登录用户和全局会话状态", "模块清晰，避免组件层层传值"),
            ("Vue Router", "管理登录、面试、报告详情等页面", "支持前端路由和访问守卫"),
            ("lucide-vue-next", "统一图标表达", "图标清晰，避免手写大量 SVG"),
            ("MediaStream / Canvas", "读取摄像头并截取分析帧", "浏览器原生能力，无需额外视频插件"),
            ("MediaRecorder", "录制用户语音回答", "将录音作为文件提交给后端"),
        ),
        widths=(3.5, 5.1, 7.0),
    )

    document.add_heading("3.4 后端技术路线", level=2)
    add_table(
        document,
        ("技术/服务", "职责"),
        (
            ("Flask", "提供 REST 风格 API，校验请求并统一返回错误"),
            ("Waitress", "生产模式下承载 Flask，避免直接使用开发服务器"),
            ("AuthService", "账号校验、PBKDF2 密码哈希、人脸向量绑定"),
            ("FaceRecognizer", "MTCNN 检测对齐，InceptionResnetV1 提取 512 维向量"),
            ("InterviewService", "组织面试与评估提示词，规范大模型返回结构"),
            ("DeepSeekService", "调用 OpenAI 兼容 Chat Completions 接口"),
            ("ASRService", "懒加载 Whisper，在线程锁下执行转写"),
            ("TTSService", "Edge-TTS 合成、音频校验、缓存和 Windows SAPI 回退"),
            ("VisionManager", "解码并缩放 JPEG，交给会话级视觉分析器"),
            ("LatestFrameWorker", "用最多 1 个待处理帧的队列保留下最新画面"),
            ("SessionStore / ReportStore", "管理内存会话与本地 JSON 报告"),
        ),
        widths=(4.6, 11.0),
    )

    document.add_heading("3.5 本地模块与云端模块怎样分工", level=2)
    add_table(
        document,
        ("能力", "运行位置", "发送到云端的数据", "网络中断时"),
        (
            ("密码验证", "本地", "无", "可继续使用"),
            ("人脸识别", "本地 CPU", "无", "可继续使用"),
            ("表情与注意力", "本地 CPU", "仅上传压缩帧到本机后端，不出本机", "本地服务仍可分析"),
            ("Whisper 语音识别", "本地 CPU", "不发送到第三方", "首次未下载模型时不可用"),
            ("Edge-TTS", "微软在线服务", "面试问题文本", "自动回退 Windows SAPI"),
            ("DeepSeek 问答与评估", "DeepSeek API", "岗位、难度、问答文字", "首题可能使用备用问题，连续追问受限"),
        ),
        widths=(3.4, 3.0, 5.3, 3.9),
    )
    add_note(
        document,
        "隐私边界",
        "“本地后端”和“第三方云端”必须分开理解。浏览器把分析帧发到本机 Flask 服务，"
        "并不等于上传到互联网；代码中 DeepSeek 请求只组装 messages 文本。",
    )

    document.add_heading("3.6 五条核心数据流", level=2)
    document.add_heading("3.6.1 登录数据流", level=3)
    add_code(
        document,
        "密码登录：用户名+密码 -> AuthService -> PBKDF2 比较 -> Flask Session\n"
        "人脸登录：浏览器拍摄 -> 本机 API -> MTCNN/FaceNet -> 欧氏距离 -> Session",
    )
    document.add_heading("3.6.2 文字问答数据流", level=3)
    add_code(
        document,
        "用户文字 -> /messages -> SessionStore 校验归属\n"
        "         -> InterviewService 取最近 20 条上下文\n"
        "         -> DeepSeek -> 下一道问题 -> 前端追加消息",
    )
    document.add_heading("3.6.3 语音问答数据流", level=3)
    add_code(
        document,
        "浏览器录音 -> /media/transcribe -> runtime/asr 临时文件\n"
        "           -> Whisper 本地转写 -> 返回文字 -> 删除临时文件\n"
        "           -> 用户确认/继续提交 -> /messages -> DeepSeek",
    )
    document.add_heading("3.6.4 视频分析数据流", level=3)
    add_code(
        document,
        "摄像头视频 -> Canvas 截 JPEG -> /frames\n"
        "            -> 单帧最新队列 -> 解码/缩放 -> MediaPipe Face Mesh\n"
        "            -> 表情规则 + PnP 注意力 -> 实时指标/快照 -> 前端看板",
    )
    document.add_heading("3.6.5 报告生成数据流", level=3)
    add_code(
        document,
        "/finish -> 停止视觉会话并汇总指标\n"
        "        -> DeepSeek 评估完整文字记录\n"
        "        -> ScoreService 融合问答与视频分数\n"
        "        -> ReportStore 写 JSON -> 返回报告编号 -> Vue 报告页",
    )

    document.add_heading("3.7 为什么系统能够保持流畅", level=2)
    add_bullets(
        document,
        (
            "最新帧优先：视觉队列 maxsize=1，新帧到了就替换还没处理的旧帧。",
            "采样降频：默认 3 帧/秒，前端可选 2、3、4，后端配置上限 8。",
            "画面预压缩：前端把帧最长边压到约 448，后端再把宽度限制到 640。",
            "一次关键点检测：同一帧的 Face Mesh 结果同时供表情和注意力模块使用。",
            "请求互斥：前端同一时间只提交一帧，避免大量 HTTP 请求同时排队。",
            "页面隐藏暂停：标签页不可见时停止无效采集，减少 CPU 和摄像头占用。",
            "模型懒加载：Whisper、人脸模型和 Face Mesh 在首次需要时加载，避免拖慢所有启动。",
            "线程隔离：视觉分析运行在独立工作线程，不直接阻塞 Flask 请求线程。",
        ),
    )

    document.add_heading("3.8 架构设计中的关键取舍", level=2)
    add_body(
        document,
        "系统选择 Flask 托管前端构建产物，而不是要求用户同时运行 Nginx、Node 和 Python。"
        "生产模式只暴露 8765 一个端口，便于学生部署和演示。代价是前端更新后必须重新执行"
        " npm run build；开发模式因此保留 Vite 的 5173 端口和代理能力。",
    )
    add_body(
        document,
        "视觉部分没有把每一帧都保存下来，也没有把视频送入大模型，而是在线计算并只保存"
        "摘要事件。这个选择降低了磁盘占用和隐私风险，也意味着报告无法像普通录像那样逐帧"
        "回放；如果二次开发需要更细的回溯，应增加受控的短视频保存策略和明确的授权机制。",
    )


def add_chapter_4_principles(document: Document) -> None:
    add_section_break(document)
    document.add_heading("第 4 章 核心技术原理", level=1)
    document.add_heading("4.1 原理速查表", level=2)
    add_table(
        document,
        ("能力", "输入", "核心方法", "系统输出"),
        (
            ("密码保护", "用户密码", "随机盐 + PBKDF2-SHA256 迭代哈希", "不可逆密码摘要"),
            ("人脸识别", "一张人脸图像", "MTCNN + InceptionResnetV1 + 欧氏距离", "用户编号与距离"),
            ("语音识别", "录音文件", "Whisper 编码器/解码器", "文字、语言、时间段"),
            ("语音合成", "问题文本", "Edge-TTS，失败时 Windows SAPI", "MP3/WAV 音频"),
            ("表情识别", "478 个面部关键点", "个人基线 + 几何阈值 + 投票平滑", "6 类状态与表情分"),
            ("注意力监测", "6 个面部点 + 虹膜点", "solvePnP 头姿 + 视线偏移 + 指数平滑", "注意力状态与分数"),
            ("智能追问", "岗位、难度、最近回答", "DeepSeek Chat Completions", "下一道问题"),
            ("最终评估", "完整文字记录", "DeepSeek JSON 模式 + 本地评分融合", "报告与最终分"),
        ),
        widths=(3.0, 3.8, 5.2, 3.8),
    )

    document.add_heading("4.2 密码认证：为什么数据库里看不到原密码", level=2)
    add_body(
        document,
        "一句话理解：系统把密码“搅拌”成一段无法直接还原的摘要，登录时用同样方法再"
        "搅拌一次，只比较摘要是否一致。即使 users.json 被人看到，也无法从摘要直接读出"
        "原密码。",
    )
    add_code(
        document,
        "随机盐 16 字节 -> PBKDF2-HMAC-SHA256 -> 迭代 310000 次 -> 摘要\n"
        "验证：重新计算摘要 -> hmac.compare_digest 常数时间比较",
    )
    add_bullets(
        document,
        (
            "随机盐：每个账号的盐不同，相同密码也不会产生完全相同的记录。",
            "重复迭代：故意增加暴力猜测的计算成本，让批量破解更慢。",
            "不可逆：验证只需要重算和比较，不需要把摘要反推成密码。",
            "常数时间比较：减少通过比较耗时推测摘要内容的侧信道风险。",
        ),
    )
    add_note(
        document,
        "注意边界",
        "PBKDF2 能降低数据库泄露后的风险，但不能防止弱密码。生产环境仍应限制登录频率、"
        "使用 HTTPS、采用高熵会话密钥，并避免把 users.json 公开到网络。",
    )

    document.add_heading("4.3 人脸识别：不是记照片，而是记“脸的坐标”", level=2)
    add_body(
        document,
        "一句话理解：先找到并摆正脸，再把脸压缩成一串 512 个数字。同一个人的不同照片"
        "通常得到相近的坐标，不同人的坐标通常离得更远。系统比较两串坐标之间的距离，"
        "而不是逐像素比较照片。",
    )
    document.add_heading("4.3.1 MTCNN 负责找脸和对齐", level=3)
    add_body(
        document,
        "MTCNN 可以理解为三层逐渐细化的“找脸流水线”：先快速找出可能包含人脸的候选区域，"
        "再过滤大量错误框，最后精修人脸框和五官位置。系统还需要旋转、缩放和裁剪，使眼睛、"
        "鼻子和嘴角落在固定位置；这一步叫对齐。对齐能减少拍摄角度和构图对识别的干扰。",
    )
    document.add_heading("4.3.2 InceptionResnetV1 负责提取特征", level=3)
    add_body(
        document,
        "对齐后的人脸进入 InceptionResnetV1。该网络采用 VGGFace2 预训练权重，输出 512 维"
        "向量。可以把它想象成一位“人脸测量员”：它不保存照片，而是记录脸型、五官相对"
        "位置和纹理等抽象特征。系统再把向量长度归一化为 1，目的是让比较主要受方向差异"
        "影响，而不是受数值整体大小影响。",
    )
    document.add_heading("4.3.3 欧氏距离和阈值怎么理解", level=3)
    add_code(
        document,
        "distance = || probe_vector - gallery_vector ||₂\n"
        "matched = distance < 1.0（默认阈值）\n"
        "similarity = clip(1 - distance / 2, 0, 1)",
    )
    add_table(
        document,
        ("阈值变化", "登记体验", "误拒绝概率", "误接受风险", "建议"),
        (
            ("调小，例如 0.85", "更难通过", "升高", "降低", "高安全场景或样本质量很好"),
            ("保持 1.0", "平衡", "中等", "中等", "本项目默认起点"),
            ("调大，例如 1.10", "更容易通过", "降低", "升高", "光照波动大但风险可控的演示"),
        ),
        widths=(3.4, 3.0, 3.0, 3.0, 3.3),
    )
    add_body(
        document,
        "系统保存的是归一化后的 .npy 向量，不保存人脸原图。注册时的原始图像在请求内存中"
        "完成推理后即被丢弃；报告中只可能出现“人脸覆盖率”这类聚合数字，不会包含识别出的"
        "身份照片。",
    )
    add_note(
        document,
        "为什么不做活体检测",
        "当前版本重点是本地身份绑定与训练演示，没有增加眨眼动作挑战、深度相机或反欺诈"
        "分类器，因此截图或视频回放可能绕过普通人脸登录。涉及资金、门禁等场景不能直接"
        "使用本项目当前方案。",
    )

    document.add_heading("4.4 Whisper 语音识别：把声波翻译成文字", level=2)
    add_body(
        document,
        "一句话理解：Whisper 不是听“字”，而是把声音切成很短的时间片段，转换成频谱特征，"
        "再由编码器和解码器预测最可能的文字序列。它可以同时处理口音、背景噪声和标点，"
        "但不会像人一样理解说话者的真实意图。",
    )
    add_table(
        document,
        ("阶段", "发生的事情", "白话类比"),
        (
            ("预加重/分帧", "声音被分成很短的片段", "把连续讲话切成许多小纸条"),
            ("梅尔频谱", "每段声音变成二维声学特征", "记录每个时刻有哪些音高和能量"),
            ("编码器", "提取声学上下文", "综合前后片段，判断大约说了什么音"),
            ("解码器", "逐词生成文字", "结合语言概率写出最可能的句子"),
            ("后处理", "返回文字、语言和时间段", "形成可直接提交的回答"),
        ),
        widths=(3.2, 6.5, 5.9),
    )
    add_body(
        document,
        "当前默认模型是 small，运行设备是 CPU。模型中 tiny、base、small、medium、large 的"
        "规律是：规格越大，通常识别越准，但下载体积、内存占用和推理时间也越高。普通电脑"
        "不建议直接把 large 作为默认值；可先尝试 base 或 small，并只针对真实的错误测例"
        "调整。",
    )
    add_body(
        document,
        "回答录音进入 runtime/asr 临时目录，转写完成后删除。Whisper 在本地推理，因此"
        "录音不会因为语音识别而发送到第三方；但如果随后把转写文字提交给 DeepSeek，这段"
        "文字会进入大模型请求，这一点应在隐私说明中明确告诉用户。",
    )

    document.add_heading("4.5 Edge-TTS 语音合成：让系统会读问题", level=2)
    add_body(
        document,
        "一句话理解：TTS 与 Whisper 的方向相反。Whisper 把声音变成文字，TTS 把文字变成"
        "声音。系统默认调用 Edge-TTS 的 zh-CN-YunxiNeural 在线音色；如果输出文件缺失、"
        "内容过小或请求连续失败，就在 Windows 上用 System.Speech/SAPI 生成 WAV。",
    )
    add_code(
        document,
        "问题文本 -> Edge-TTS（最多尝试 2 次）\n"
        "          -> 有效 MP3 -> runtime/tts -> 浏览器播放\n"
        "          -> 失败 -> Windows SAPI WAV -> 浏览器播放\n"
        "          -> 缓存文件超过 24 小时自动清理",
    )
    add_body(
        document,
        "在线音色更自然，但依赖网络；本地 SAPI 音质通常较机械，却能在断网时维持流程。"
        "这是“可用性优先”的降级设计。如果浏览器禁止自动播放，用户需要主动点击播放按钮，"
        "这不是 TTS 故障。",
    )

    document.add_heading("4.6 MediaPipe Face Mesh：一张脸的点阵图", level=2)
    add_body(
        document,
        "一句话理解：Face Mesh 在脸上铺一张有数百个关键点的“网格”。系统不需要保存整张"
        "图片，只要根据这些点的相对距离和角度，就能估计眼睛开合、嘴角变化、头部姿态和"
        "视线方向。MediaPipe 基础面部网格约 468 个点，开启虹膜细化后可使用约 478 个点。",
    )
    add_table(
        document,
        ("系统用途", "关键点思路", "核心指标"),
        (
            ("眼睛", "左右眼各取 6 点", "EAR 眼睛纵横比、眨眼"),
            ("微笑/大笑", "嘴角、上下唇相对位置", "嘴角色差比、嘴部开度"),
            ("惊讶", "眉毛抬高 + 嘴巴张开", "眉毛与嘴部相对基线变化"),
            ("困惑", "左右眉高度不对称", "眉部不对称差"),
            ("说话", "嘴巴开度短时波动", "嘴部活动标准差"),
            ("注意力", "鼻尖、下巴、眼角和嘴角 6 点", "pitch、yaw、roll"),
            ("视线", "左右虹膜中心与眼睑边缘", "水平 gaze_x 偏移"),
        ),
        widths=(3.1, 6.1, 6.4),
    )
    add_body(
        document,
        "Face Mesh 设置 max_num_faces=1，也就是只分析画面中置信度最高的一张脸。"
        " min_detection_confidence 与 min_tracking_confidence 均默认为 0.5。每帧只执行"
        "一次网格推理，得到的 landmarks 同时交给表情分析器和注意力分析器，避免重复计算。",
    )
    add_note(
        document,
        "关键点不等于情绪",
        "同一种嘴角变化可能来自微笑、说话、肌肉抽动或拍摄角度。系统做的是可解释的几何"
        "状态估计，不应把“微笑少”解释为“人不积极”，更不能据此评价人格或心理健康。",
    )

    document.add_heading("4.7 面部表情识别：先认识你自己的“中性脸”", level=2)
    add_body(
        document,
        "一句话理解：每个人的脸型、嘴宽和眉毛位置不同，固定阈值容易误判。因此系统先让"
        "用户自然面对镜头，采集一段个人基线；之后不是看“嘴巴绝对有多开”，而是看它相对"
        "自己的基线变化了多少。",
    )
    document.add_heading("4.7.1 个人校准", level=3)
    add_steps(
        document,
        (
            "预热 5 个含人脸帧，跳过刚打开摄像头时的光照和姿态波动。",
            "再采集 30 个有效帧，记录眼睛、微笑、嘴部、眉毛等特征。",
            "对每个特征取中位数作为个人基线。中位数比平均值更不容易被极端帧拉偏。",
            "计算 MAD（中位数绝对偏差）作为个人噪声尺度，再结合最小阈值形成自适应阈值。",
            "校准完成前显示“校准中”，此时不把表情分当成有效结论。",
        ),
    )
    document.add_heading("4.7.2 六个状态如何判断", level=3)
    add_code(
        document,
        "惊讶：眉毛抬高 + 嘴巴明显张大\n"
        "大笑：嘴角抬高 + 嘴部开度进一步增大\n"
        "微笑：嘴角抬高超过自适应阈值\n"
        "说话：嘴部短时波动明显且相对基线张开\n"
        "困惑：眉部不对称增大、眉抬高但嘴没有大幅张开\n"
        "中性：以上条件均未满足，保持默认状态",
    )
    add_body(
        document,
        "连续帧可能一秒内向“微笑”和“中性”来回跳动。系统使用长度为 5 的候选历史做投票，"
        "只有同一状态至少得到 3 票才确认输出；眼睛和嘴部等的短窗口均值也会进一步降低"
        "抖动。状态变化持续至少 0.5 秒才写入事件，避免报告被瞬间噪声填满。",
    )
    document.add_heading("4.7.3 表情分怎么算", level=3)
    add_code(
        document,
        "基础分 = 50\n"
        "积极比例 = （微笑帧 + 大笑帧） / 已分析帧\n"
        "表情分 = 50 + 积极比例 × 50\n"
        "若中性帧占比超过 80%，再减 20\n"
        "最终限制在 0-100，保留 1 位小数",
    )
    add_body(
        document,
        "这个分数的目标是鼓励自然、积极的表达状态，而不是要求一直笑。超过八成时间都"
        "是中性时会扣分，是训练导向规则，不具备心理学诊断意义。",
    )

    document.add_heading("4.8 注意力监测：头在哪里、眼睛看哪里", level=2)
    add_body(
        document,
        "一句话理解：注意力不是单一指标，而是“头部朝向是否明显偏、眼睛是否大致看镜头、"
        "画面中是否持续有人脸”三个信号的综合。系统先判断瞬时状态，再平滑，最后用人脸"
        "覆盖率修正，避免只看几帧就给出极端结论。",
    )
    document.add_heading("4.8.1 头部姿态 pitch、yaw、roll", level=3)
    add_body(
        document,
        "系统在检测到的人脸中取鼻尖、下巴、左右眼角和左右嘴角 6 个二维点，再与一个标准"
        "三维人脸模型做 cv2.solvePnP。它本质上是求解“标准脸怎样旋转和平移，才会在相机"
        "里投影成当前这 6 个点”。得到的旋转可拆成 pitch（抬头/低头）、yaw（左右转头）"
        "和 roll（左右倾斜）。",
    )
    add_code(
        document,
        "默认判定阈值：|pitch| ≤ 15°，|yaw| ≤ 20°，|roll| ≤ 15°\n"
        "初始注意力分 = 100\n"
        "yaw 超出：最多扣 30；pitch 超出：最多扣 25\n"
        "roll 超出：最多扣 20；视线偏移：固定扣 25\n"
        "瞬时状态：分数 ≥ 70 视为集中，否则视为分散",
    )
    document.add_heading("4.8.2 视线估计", level=3)
    add_body(
        document,
        "系统利用虹膜中心相对眼睛水平中心的位置估计 gaze_x。由于左右眼会分别计算，最终"
        "取平均值；为了减少眼睛轻微移动造成的跳动，又对最近 30 个 gaze_x 求均值。平滑后"
        "偏移绝对值小于 0.15 时，认为视线大致居中。",
    )
    document.add_heading("4.8.3 两个平滑机制", level=3)
    add_body(
        document,
        "注意力分数使用指数平滑，当前新分数只占 24%，上一状态保留 76%。这样一次误检不会"
        "让看板瞬间从高分变成零分。报告中的最终集中度仍会对逐帧结果求平均，因此一段"
        "短暂低头不会支配整场结果。",
    )
    add_code(
        document,
        "平滑分 = 0.24 × 当前瞬时分 + 0.76 × 上一平滑分\n"
        "覆盖率 = min(1, 含人脸分析帧 / 总分析帧)\n"
        "覆盖率修正 = 0.75 + 0.25 × 覆盖率\n"
        "修正后注意力分 = 原始注意力分 × 覆盖率修正",
    )
    add_body(
        document,
        "覆盖率修正的设计含义是：如果人长期离开画面，即使已经分析的少量人脸帧都很集中，"
        "也不能得到满分。人脸从未出现时覆盖率为 0，修正因子为 0.75；全覆盖时为 1.0。"
        "此外，服务层要求至少 10 个含人脸分析帧，视频结果才可能有效。",
    )
    add_note(
        document,
        "为什么不是盯眼仪",
        "单目普通摄像头无法精确测量真实世界注视点，只能估计相对于脸部的视线偏移。"
        "戴眼镜、低光、侧坐和镜头位置都会影响结果，因此它适合趋势反馈，不适合精确考核。",
    )

    document.add_heading("4.9 DeepSeek 大模型接入：会追问，也会做结构化评估", level=2)
    add_body(
        document,
        "一句话理解：DeepSeek 不是本地训练的“面试模型”，而是通过 HTTP 调用的通用大语言"
        "模型。系统通过提示词规定它的身份、任务、输出格式和边界，再提供岗位、难度与对话"
        "上下文，让它生成下一题；结束时再把完整文字记录交给它评估。",
    )
    add_code(
        document,
        "POST /chat/completions\n"
        "Authorization: Bearer <DEEPSEEK_API_KEY>\n"
        "{\n"
        '  "model": "deepseek-chat",\n'
        '  "messages": [\n'
        '    {"role": "system", "content": "面试官/评估员规则"},\n'
        '    {"role": "user", "content": "候选人回答"}\n'
        "  ],\n"
        '  "temperature": 0.8,\n'
        '  "stream": false\n'
        "}",
    )
    add_bullets(
        document,
        (
            "system 消息定义角色和硬性规则，例如一次只问一个问题、不泄露提示词。",
            "user/assistant 消息保存用户回答和面试官问题，组成连续上下文。",
            "追问默认取最近 20 条有效消息，既可延续话题，又控制请求规模。",
            "最终评估要求仅返回 JSON，后端还会兼容 Markdown 代码围栏并再次规范字段。",
            "只发送文字、岗位和难度，不把视频帧、人脸向量或录音发送给 DeepSeek。",
        ),
    )
    add_body(
        document,
        "接口设置 90 秒超时。若首题生成失败，系统可以使用备用开场问题维持演示；后续"
        "追问或最终评估失败时，会在接口层返回明确错误，而不是伪造一条看似正常的 AI "
        "评估。生产使用中应记录错误类型，但要避免把 API Key 写进日志。",
    )

    document.add_heading("4.10 最终评分：为什么不是简单平均", level=2)
    add_body(
        document,
        "一句话理解：面试的核心是回答内容，面部状态只作为过程参考。因此问答占 80%，"
        "表情和注意力各占视频部分的 50%，视频整体只占最终分的 20%。",
    )
    add_code(
        document,
        "视频综合分 = 表情分 × 0.5 + 注意力分 × 0.5\n"
        "最终得分 = 问答分 × 0.8 + 视频综合分 × 0.2",
    )
    add_body(
        document,
        "计算使用 Decimal 和 ROUND_HALF_UP，最终保留 1 位小数，避免普通浮点数带来尾差。"
        "只有表情已校准且有分析帧、注意力存在有效帧、含人脸分析帧达到最低数量时，视频"
        "结果才算有效。视频无效时最终得分为空，系统不会把 20% 临时分给问答。",
    )
    add_table(
        document,
        ("示例项目", "问答分", "表情分", "注意力分", "视频综合", "最终分"),
        (
            ("真实报告样例", "68.0", "50.8", "74.6", "62.7", "66.9"),
            ("示例计算", "68.0", "50.8", "74.6", "62.7", "68.0×0.8 + 62.7×0.2"),
        ),
        widths=(3.3, 2.2, 2.2, 2.4, 2.5, 3.0),
    )
    add_note(
        document,
        "读分原则",
        "不要只比较最终分。学生应先读逐题复盘，判断是知识缺口、回答结构问题还是过程状态"
        "问题；视频分用于提醒训练状态，不能替代对回答内容的分析。",
    )


def add_chapter_5_models_and_data(document: Document) -> None:
    add_section_break(document)
    document.add_heading("第 5 章 模型、数据与隐私说明", level=1)
    document.add_heading("5.1 先分清模型、算法和规则", level=2)
    add_body(
        document,
        "项目里经常出现“模型训练”这个说法，但并不是每个能力都由项目自己训练。模型是"
        "从数据中学习到参数的程序；算法是处理问题的步骤；规则则是开发者根据业务目标"
        "设定的判断条件。本项目把三者组合起来：预训练模型负责“看懂”，规则负责“判断”，"
        "大模型负责“交流与总结”。",
    )
    add_table(
        document,
        ("名称", "类型", "运行位置", "能不能直接读懂"),
        (
            ("DeepSeek", "预训练大语言模型 API", "DeepSeek 云端", "把文字作为上下文，生成问题或评估"),
            ("Whisper", "预训练语音识别模型", "本机 CPU", "把音频波形转换为中文文字"),
            ("Edge-TTS", "在线神经网络语音合成服务", "微软在线服务", "把问题文字转换为面试官语音"),
            ("MTCNN", "预训练人脸检测网络", "本机 CPU", "找到人脸框和五官关键点"),
            ("InceptionResnetV1", "预训练人脸特征网络", "本机 CPU", "把一张脸压成 512 个数字"),
            ("MediaPipe Face Mesh", "预训练面部关键点模型", "本机 CPU", "在脸上找到约 478 个关键点"),
            ("solvePnP", "计算机视觉算法", "本机 CPU", "根据二维点和三维模型估算头部姿态"),
            ("表情与注意力规则", "可解释阈值与时序逻辑", "本机 CPU", "把关键点变化变成状态与分数"),
        ),
        widths=(4.1, 3.5, 3.1, 5.4),
    )
    add_note(
        document,
        "最重要的一句话",
        "当前系统开箱即用，正常面试不需要训练任何模型。所谓“调优”多数是改配置、提示词"
        "或规则阈值；只有准备替换成自己的深度模型时，才需要准备数据并训练。",
    )

    document.add_heading("5.2 各模型从哪里来、负责什么", level=2)
    add_body(
        document,
        "MTCNN 和 InceptionResnetV1 由 facenet-pytorch 提供。InceptionResnetV1 使用"
        " VGGFace2 人脸数据预训练参数，项目只使用它的特征提取能力，不用它做情绪判断。"
        "Whisper 由 OpenAI 发布，项目使用 openai-whisper 的 small 规格；MediaPipe Face "
        "Mesh 由 Google 发布，项目使用其中的面部网格与虹膜关键点。DeepSeek 通过公开"
        "的 OpenAI 兼容 HTTP 接口调用，不下载模型权重到本机。",
    )
    add_body(
        document,
        "这些预训练模型已经在通用大规模数据上学习过，所以项目无需从零收集海量人脸、"
        "语音或语言数据。代价是它们不一定完全适合某个学校、行业或摄像头的场景，因此"
        "仍需要现场校准、清晰采集和适量阈值调优。",
    )

    document.add_heading("5.3 本项目使用的数据是什么", level=2)
    add_table(
        document,
        ("数据", "来源", "用途", "是否进入第三方云端"),
        (
            ("账号与显示名", "用户注册时输入", "登录、报告归属、页面显示", "否"),
            ("密码摘要", "本机 PBKDF2 计算", "验证密码，不保存明文", "否"),
            ("人脸特征向量", "摄像头照片经本机网络提取", "人脸登记与登录匹配", "否"),
            ("实时视频帧", "浏览器摄像头截取", "本机表情和注意力分析", "否"),
            ("回答录音", "浏览器 MediaRecorder", "交给本机 Whisper 转写", "否"),
            ("回答文字", "键盘输入或语音转写", "连续追问和最终评估", "是，发送给 DeepSeek"),
            ("面试问题", "DeepSeek 生成", "显示、保存、语音播报", "问题文本发送给 Edge-TTS"),
            ("指标与事件", "本机视觉分析生成", "报告统计、趋势和复盘", "否"),
            ("报告 JSON", "本机后端生成", "历史记录和二次统计", "否"),
        ),
        widths=(3.1, 4.0, 5.8, 3.4),
    )
    add_body(
        document,
        "这里的“云端”要分成两类：DeepSeek 会接收岗位、难度和问答文字；Edge-TTS 会接收"
        "需要朗读的问题文本。原始视频、人脸照片、人脸向量和回答录音不会上传给它们。"
        "本机浏览器把分析帧提交到 127.0.0.1，只是在同一台电脑的 Vue 和 Flask 之间传输。",
    )

    document.add_heading("5.4 数据怎样流动和保存", level=2)
    add_code(
        document,
        "注册/登录\n"
        "  users.json：账号、显示名、密码摘要、人脸绑定状态\n"
        "  faces/<user_id>.npy：512 维人脸特征向量\n\n"
        "语音回答\n"
        "  runtime/asr/：短暂保存上传音频\n"
        "      -> Whisper 转写\n"
        "      -> 删除临时音频\n"
        "      -> 文字进入当前会话\n\n"
        "面试结束\n"
        "  runtime/reports/<user_id>/<session_id>.json\n"
        "  保存问答文字、评估结果、视频汇总、评分和隐私标记",
    )
    add_table(
        document,
        ("路径", "内容", "生命周期"),
        (
            ("data/auth/users.json", "账号元数据与密码摘要", "持久保存"),
            ("data/auth/faces/", "每个用户一个人脸向量文件", "持久保存，可重新绑定覆盖"),
            ("data/.flask_secret", "自动生成的会话签名密钥", "持久保存"),
            ("runtime/asr/", "上传的临时语音文件", "转写完成后删除"),
            ("runtime/tts/", "合成后的音频缓存", "缓存复用，超过 24 小时自动清理"),
            ("runtime/reports/", "最终面试报告 JSON", "持久保存，用户删除面试时清理"),
        ),
        widths=(5.2, 6.1, 5.0),
    )

    document.add_heading("5.5 密码和人脸数据的安全处理", level=2)
    add_body(
        document,
        "密码使用 PBKDF2-HMAC-SHA256，随机盐长度为 16 字节，迭代 310000 次，最后使用"
        "常数时间比较验证。随机盐使相同密码在不同账号中产生不同摘要；迭代会故意增加"
        "暴力破解成本。即使 users.json 被看到，也不能直接得到原始密码，但仍应妥善保护"
        "数据目录和备份文件。",
    )
    add_body(
        document,
        "人脸登记只保存特征向量。向量是一串代表脸部特征关系的数字，不是可正常查看的"
        "照片，但它仍属于生物特征数据。系统不是支付级身份认证，阈值也不是绝对安全边界。"
        "公开部署时应取得明确授权、提供注销与删除机制，并使用 HTTPS、访问控制和审计日志。",
    )
    add_note(
        document,
        "禁止把密钥写进文档或截图",
        "不要把 .env、DEEPSEEK_API_KEY、data/.flask_secret、users.json、人脸向量或报告"
        "提交到公开 Git 仓库。截图时也要遮住账号、API Key、真实姓名和摄像头隐私画面。",
    )

    document.add_heading("5.6 报告里保存了什么、不保存什么", level=2)
    add_table(
        document,
        ("内容", "报告中的处理"),
        (
            ("原始视频帧", "不保存；在内存中分析后丢弃"),
            ("人脸原始照片", "不保存；只输出向量并写 .npy 文件"),
            ("回答录音", "不保存；转写结束后删除临时文件"),
            ("问答文字", "保存，用于历史查看和面试复盘"),
            ("分数与维度", "保存，包含问答、表情、注意力、视频综合和最终分"),
            ("聚合状态", "保存，如表情分布、集中比例、眨眼次数、离屏事件"),
            ("状态时间线", "保存摘要事件，不逐帧保存录像"),
            ("隐私标记", "保存三个布尔标记，明确视频、人脸照片和音频均不持久化"),
        ),
        widths=(4.3, 12.0),
    )
    add_body(
        document,
        "报告采用 JSON 文件，不依赖外部数据库，便于复制和课程检查。目录按用户编号和"
        "面试编号隔离，接口也会检查会话所属用户。它适合本地运行和小规模使用，不适合"
        "高并发、跨服务器或需要复杂权限治理的生产账号系统。",
    )

    document.add_heading("5.7 模型能力边界：能做什么、不能做什么", level=2)
    add_table(
        document,
        ("能力", "可以得出的结论", "不能据此得出的结论"),
        (
            ("人脸识别", "当前人脸与已登记向量是否足够接近", "识别真实身份、判断人性或情绪"),
            ("表情规则", "相对个人基线的几何状态变化", "诊断焦虑、抑郁、撒谎或心理疾病"),
            ("注意力估计", "头姿、视线偏移与人脸覆盖率的综合趋势", "证明真实注意力或认知能力"),
            ("Whisper", "把可听清的中文语音转成候选文字", "在强噪声和口音下保证逐字正确"),
            ("DeepSeek 评分", "对本次文字回答的阶段性训练建议", "替代招聘决策或真实职业评价"),
        ),
        widths=(3.3, 6.6, 6.4),
    )

    document.add_heading("5.8 替换模型时先看四条规则", level=2)
    add_bullets(
        document,
        (
            "保持接口输出不变：新模型也要返回页面需要的字段，否则前端和报告会同时失效。",
            "重新校准阈值：换人脸、人脸检测或姿态模型后，旧阈值通常不能直接套用。",
            "重新生成特征：人脸特征维度或训练数据变化后，旧向量不能保证与新模型兼容。",
            "保留回归集：每次调参前后，用固定样例比较准确率、延迟、误判和报告结果。",
        ),
    )
    add_note(
        document,
        "最小改动原则",
        "如果只是改岗位、题目风格、上下文或评分偏好，优先改提示词和 .env；如果只是"
        "觉得敏感或宽松，先调阈值；只有通用模型明显无法满足目标时，再考虑训练新模型。",
    )


def add_chapter_6_code_structure(document: Document) -> None:
    add_section_break(document)
    document.add_heading("第 6 章 代码结构与调用链", level=1)
    document.add_heading("6.1 一张目录卡看懂项目", level=2)
    add_code(
        document,
        "smartspeech_optimized_new/\n"
        "├─ backend/\n"
        "│  ├─ app/\n"
        "│  │  ├─ api/              HTTP 接口与权限检查\n"
        "│  │  ├─ services/         认证、面试、语音、人脸、评分、视觉\n"
        "│  │  ├─ storage/          内存会话与 JSON 报告\n"
        "│  │  └─ config.py         .env 读取与路径定义\n"
        "│  ├─ tests/               后端单元测试\n"
        "│  └─ run.py               Waitress 启动入口\n"
        "├─ frontend/\n"
        "│  ├─ src/\n"
        "│  │  ├─ api/client.js     前端 HTTP 客户端\n"
        "│  │  ├─ components/       摄像头、对话、指标、弹窗等组件\n"
        "│  │  ├─ composables/      摄像头与录音组合逻辑\n"
        "│  │  ├─ pages/            登录、面试、报告页面\n"
        "│  │  ├─ router/           路由与登录守卫\n"
        "│  │  └─ stores/auth.js    Pinia 登录状态\n"
        "│  ├─ dist/                npm run build 生成物\n"
        "│  └─ package.json         前端依赖与命令\n"
        "├─ data/                   账号、人脸向量和持久密钥\n"
        "├─ runtime/                临时音频、TTS 缓存和报告\n"
        "├─ scripts/                文档与维护脚本\n"
        "├─ .env                    本机私有配置\n"
        "├─ requirements.txt        Python 依赖\n"
        "└─ start_system.bat        一键安装、构建和启动",
    )

    document.add_heading("6.2 前端代码各自负责什么", level=2)
    add_table(
        document,
        ("文件", "职责", "常见修改点"),
        (
            ("main.js", "创建 Vue 应用并挂载 Pinia、Router", "全局插件与错误处理"),
            ("App.vue", "应用根组件与页面容器", "全局布局和主题入口"),
            ("router/index.js", "定义页面路由、登录守卫和页面跳转", "新增页面或修改权限"),
            ("stores/auth.js", "保存当前用户、获取用户信息和退出登录", "登录状态字段"),
            ("api/client.js", "统一调用 /api、处理 JSON 与错误", "接口路径、请求头、超时"),
            ("pages/LoginPage.vue", "密码登录、人脸登录、注册与人脸绑定", "登录页交互"),
            ("pages/InterviewPage.vue", "面试配置、开始、帧采样、结束和实时数据", "采样策略、工作台流程"),
            ("pages/ReportsPage.vue", "历史报告统计和列表", "列表字段与筛选"),
            ("pages/ReportDetailPage.vue", "评分、逐题复盘、趋势、事件和打印", "报告展示结构"),
            ("components/CameraPanel.vue", "摄像头预览、连接状态和画面覆盖信息", "视频叠层与提示"),
            ("components/ConversationPanel.vue", "问答消息、文字输入、录音和 TTS 播放", "对话交互"),
            ("components/LiveMetricsPanel.vue", "实时状态、综合分、趋势和动态统计", "指标卡片与图表"),
            ("components/FaceCaptureModal.vue", "注册或绑定时拍摄清晰人脸", "构图约束和错误提示"),
            ("composables/useCamera.js", "打开、关闭、切换摄像头并管理 MediaStream", "设备与前后摄像头"),
            ("utils/media.js", "截图、压缩、时长格式化等浏览器媒体工具", "帧尺寸与 JPEG 质量"),
        ),
        widths=(4.5, 7.3, 4.5),
    )

    document.add_heading("6.3 后端接口层：只负责收请求、验身份和返回结果", level=2)
    add_table(
        document,
        ("文件", "职责"),
        (
            ("api/auth.py", "注册、密码登录、获取当前用户、人脸注册和人脸登录"),
            ("api/interview.py", "创建、开场、帧上传、指标读取、消息、结束、报告和删除"),
            ("api/media.py", "音频转写与 TTS 合成"),
            ("api/system.py", "健康检查与服务能力探测"),
            ("api/common.py", "登录校验、JSON 解析、图片校验和统一错误响应"),
        ),
        widths=(4.5, 11.8),
    )
    add_body(
        document,
        "接口层不亲自做深度学习推理，而是先检查“你是谁、请求参数是否合法、这个面试是否"
        "属于你”，再调用服务层。这样做的意义是：认证规则、业务规则和模型算法可以分别"
        "测试，替换模型时也不需要重写所有路由。",
    )

    document.add_heading("6.4 后端服务层：项目的核心业务", level=2)
    add_table(
        document,
        ("文件", "核心工作", "主要输入输出"),
        (
            ("container.py", "创建并复用各项服务，避免每次请求重复加载模型", "应用级服务容器"),
            ("auth_service.py", "用户校验、PBKDF2、人脸向量登记与登录匹配", "用户记录与向量"),
            ("face_recognition.py", "MTCNN 检测、InceptionResnetV1 特征提取、距离匹配", "512 维向量或用户编号"),
            ("deepseek_service.py", "调用 Chat Completions、解析 JSON、处理超时与错误", "messages 与文本/字典"),
            ("interview_service.py", "拼装面试官/评估员提示词，规范大模型输出", "追问与评估结构"),
            ("asr_service.py", "懒加载 Whisper 并在线程锁内转写", "音频文件与文字"),
            ("tts_service.py", "Edge-TTS 合成、WAV 校验、缓存和 SAPI 回退", "问题文本与音频响应"),
            ("score_service.py", "按固定权重融合问答、表情和注意力", "最终分与公式记录"),
            ("vision/manager.py", "按会话创建视觉分析器与工作线程", "会话级视觉资源"),
            ("vision/worker.py", "单帧最新队列和异步处理", "推入帧、读取统计"),
            ("vision/session_analyzer.py", "统一执行 Face Mesh，并汇总表情、注意力", "实时指标与视频汇总"),
            ("vision/expression_analyzer.py", "个人校准、自适应阈值、时序投票和表情分", "表情状态与分布"),
            ("vision/attention_analyzer.py", "PnP 头姿、虹膜视线、平滑和注意力分", "注意力状态与事件"),
        ),
        widths=(4.7, 7.1, 4.5),
    )

    document.add_heading("6.5 存储层：会话和报告", level=2)
    add_body(
        document,
        "SessionStore 保存正在进行的面试。它包含岗位、难度、音色、采样配置、文字记录、"
        "视觉分析器、工作线程和锁。会话主要驻留在内存中，默认有约 12 小时有效窗口；"
        "服务重启后，尚未结束的内存会话不会像数据库那样完整恢复。ReportStore 则把完成"
        "后的报告写到 JSON 文件，历史页面读取的是这些持久文件。",
    )
    add_table(
        document,
        ("对象", "保存周期", "典型内容"),
        (
            ("InterviewSession", "面试期间，主要在内存", "配置、transcript、analyzer、worker、status"),
            ("Report JSON", "面试结束后持久保存", "转录、评估、视频汇总、评分和隐私标记"),
            ("Users JSON", "账号创建后持久保存", "用户名、显示名、密码摘要、人脸绑定状态"),
            ("Face NPY", "人脸登记后持久保存", "512 维归一化向量"),
        ),
        widths=(4.3, 4.4, 7.6),
    )

    document.add_heading("6.6 一条文字问答调用链", level=2)
    add_code(
        document,
        "ConversationPanel 提交回答\n"
        "  -> api/client.js: POST /api/interviews/<id>/messages\n"
        "  -> api/interview.py: 登录校验 + 会话归属校验 + 参数校验\n"
        "  -> interview_service.py: 组织 system 提示词\n"
        "  -> session_store.py: 取最近 20 条有效消息\n"
        "  -> deepseek_service.py: POST /chat/completions, temperature=0.8\n"
        "  -> interview.py: 追加 user/assistant 消息\n"
        "  -> Vue: 显示回答与下一道问题",
    )
    add_body(
        document,
        "这条链的关键边界是“只传文字”。用户回答、面试官问题和岗位难度会进入大模型；"
        "视频指标留在本地，不会拼进 messages。最后评估会改用另一套评估提示词和 JSON "
        "输出约束，温度降到 0.2，目的是减少评分波动。",
    )

    document.add_heading("6.7 一条视频分析调用链", level=2)
    add_code(
        document,
        "CameraPanel 显示原始视频（浏览器直接播放，不经过 Python）\n"
        "  -> useCamera + utils/media.js: Canvas 截 JPEG\n"
        "  -> POST /frames\n"
        "  -> VisionManager 找到当前会话的 SessionAnalyzer\n"
        "  -> LatestFrameWorker 队列 maxsize=1\n"
        "  -> OpenCV 解码 -> 缩放 -> MediaPipe Face Mesh\n"
        "  -> ExpressionAnalyzer + AttentionAnalyzer\n"
        "  -> GET /metrics / 帧接口响应 -> LiveMetricsPanel",
    )
    add_body(
        document,
        "浏览器里的视频预览是连续的，分析帧只是按 2、3 或 4 FPS 抽取的少量截图。"
        "工作线程与 Flask 请求线程分开；如果上一帧还没分析完，新帧会替换待处理旧帧，"
        "而不是无限排队。因此“丢帧”有时是保持实时性的主动策略，不代表接口失败。",
    )

    document.add_heading("6.8 一条报告生成调用链", level=2)
    add_code(
        document,
        "点击结束面试\n"
        "  -> POST /finish\n"
        "  -> 停止 VisionWorker，得到 video_summary\n"
        "  -> InterviewService.evaluate：完整文字记录 + JSON 评估\n"
        "  -> score_service.fuse_interview_scores\n"
        "  -> ReportStore.save：写 runtime/reports/<user>/<session>.json\n"
        "  -> 返回 report\n"
        "  -> Vue 跳转报告详情",
    )
    add_body(
        document,
        "问答和视频是两条独立结果线：问答分来自大模型，表情分和注意力分来自本地规则。"
        "score_service 只负责最后融合，不重新解释原始回答。这样可以在报告里同时保留"
        "“问答分”“表情分”“注意力分”和“最终分”，避免单一数字丢失诊断信息。",
    )

    document.add_heading("6.9 配置与代码的对应关系", level=2)
    add_table(
        document,
        ("想修改什么", "优先改哪里", "是否需要改代码"),
        (
            ("API Key、模型、超时", ".env", "通常不需要"),
            ("Whisper 大小和设备", ".env", "不需要"),
            ("默认 TTS 音色", ".env", "不需要"),
            ("默认帧率和上限", ".env", "不需要"),
            ("人脸匹配严格度", ".env 中的 FACE_MATCH_THRESHOLD", "不需要"),
            ("面试官语气与规则", "interview_service.py 的提示词", "是"),
            ("评估 JSON 字段", "interview_service.py 与报告页", "是"),
            ("评分权重", "score_service.py，并同步测试和页面文案", "是"),
            ("表情阈值或状态", "expression_analyzer.py", "是"),
            ("注意力阈值或平滑", "attention_analyzer.py", "是"),
            ("页面布局与指标卡", "frontend/src/components 和 pages", "是"),
        ),
        widths=(5.1, 6.7, 4.5),
    )

    document.add_heading("6.10 读代码的推荐顺序", level=2)
    add_steps(
        document,
        (
            "先读 README 和 backend/app/config.py，知道系统目标和所有外部配置。",
            "再读 backend/run.py 与 backend/app/__init__.py，理解 Waitress 和 Flask 如何启动。",
            "按一个功能读接口：先看 interview.py，再顺着调用进入 interview_service 和 storage。",
            "读前端 InterviewPage.vue，观察开始、采样、提交、指标轮询和结束分别调用哪个 API。",
            "最后读视觉模块。先看 session_analyzer 的统一入口，再分别看 expression 和 attention。",
            "修改后用单元测试、浏览器操作和报告样例三方面一起验证，不只确认“代码不报错”。",
        ),
    )
    add_note(
        document,
        "调试技巧",
        "遇到问题时先判断故障在哪一段：页面不显示先看浏览器控制台；接口失败看 Network；"
        "服务报错看启动终端；模型异常看日志中服务名称；报告缺失看 finish 是否成功以及"
        "runtime/reports 下是否生成 JSON。",
    )


def add_chapter_7_install(document: Document) -> None:
    add_section_break(document)
    document.add_heading("第 7 章 安装、配置与启动", level=1)
    document.add_heading("7.1 先准备什么", level=2)
    add_table(
        document,
        ("项目", "推荐要求", "检查命令"),
        (
            ("操作系统", "Windows 10/11，64 位", "winver"),
            ("Python", "3.10，项目固定使用 nlp_env", "python --version"),
            ("Node.js", "20 或更高版本", "node -v"),
            ("npm", "随 Node.js 安装", "npm -v"),
            ("浏览器", "Chrome 或 Edge", "打开摄像头测试页"),
            ("网络", "能访问 DeepSeek；Edge-TTS 依赖在线服务", "浏览器访问 API 地址"),
            ("磁盘", "建议至少留出数 GB 给依赖、模型和缓存", "查看系统磁盘空间"),
        ),
        widths=(3.2, 8.1, 5.2),
    )
    add_body(
        document,
        "项目文档和启动脚本按以下解释器路径编写。若你的 Anaconda 安装位置不同，需要"
        "同时修改 start_system.bat 中的 PYTHON_EXE，以及本教程命令里的解释器路径。",
    )
    add_code(
        document,
        r"D:\Work_Software\Anaconda3\envs\nlp_env\python.exe",
    )

    document.add_heading("7.2 第一次安装：最稳妥的顺序", level=2)
    add_steps(
        document,
        (
            "打开 PowerShell，进入项目根目录。路径可用资源管理器复制，避免手工输错。",
            "确认 Python 解释器存在并查看版本，应为 Python 3.10。",
            "执行 Python 依赖安装命令。不要跳过版本约束，尤其是 NumPy 1.26.4。",
            "进入 frontend，执行 npm ci。npm ci 会严格依据 package-lock.json 安装。",
            "回到项目根目录，把 .env.example 复制为 .env，填写 DEEPSEEK_API_KEY。",
            "执行 npm run build，确认 frontend/dist 生成成功。",
            "执行后端启动命令，等待“正在启动”提示后访问健康检查地址。",
        ),
    )
    add_code(
        document,
        r"""cd "D:\Project_Code_2026\tyut_comprehensive_pratice\smartspeech_optimized_new"

& "D:\Work_Software\Anaconda3\envs\nlp_env\python.exe" --version
& "D:\Work_Software\Anaconda3\envs\nlp_env\python.exe" -m pip install -r requirements.txt

cd frontend
npm ci
npm run build
cd ..

Copy-Item .env.example .env -ErrorAction SilentlyContinue
& "D:\Work_Software\Anaconda3\envs\nlp_env\python.exe" -m backend.run""",
    )
    add_note(
        document,
        "首次安装为什么慢",
        "PyTorch、MediaPipe、OpenCV 等包体积较大；首次使用 Whisper 还会下载模型，首次"
        "使用人脸模型或 Face Mesh 也会初始化权重。只要没有报错，等待完成后后续启动会"
        "复用缓存。",
    )

    document.add_heading("7.3 一键启动：日常最推荐", level=2)
    add_steps(
        document,
        (
            "确认 .env 已存在，并已填写有效的 DEEPSEEK_API_KEY。",
            "双击项目根目录的 start_system.bat。",
            "脚本先做健康检查；若系统已在运行，会直接打开现有页面。",
            "若尚未运行，脚本检查端口、Python 依赖、前端依赖并重新构建 Vue。",
            "脚本启动 Flask + Waitress，等待 /api/health 返回成功后自动打开浏览器。",
            "使用期间不要关闭启动脚本的命令行窗口；关闭窗口即停止服务。",
        ),
    )
    add_body(
        document,
        "默认访问地址是 http://127.0.0.1:8765。一键脚本会从 .env 读取 APP_HOST 和"
        " APP_PORT，所以改了端口后不需要再改脚本中的访问逻辑。它会优先使用"
        r" D:\Work_Software\nodejs\npm.cmd，找不到时再使用 PATH 中的 npm。",
    )

    document.add_heading("7.4 手动启动生产模式", level=2)
    add_body(
        document,
        "生产模式的含义不是必须上互联网，而是使用 Waitress 提供前端构建产物和后端 API。"
        "它适合日常使用、课程验收和本地演示。每次修改 Vue 源码后，必须重新运行 npm run "
        "build，否则浏览器仍会看到旧的 dist 构建结果。",
    )
    add_code(
        document,
        r"""cd "D:\Project_Code_2026\tyut_comprehensive_pratice\smartspeech_optimized_new"

cd frontend
npm run build
cd ..

& "D:\Work_Software\Anaconda3\envs\nlp_env\python.exe" -m backend.run""",
    )

    document.add_heading("7.5 前后端开发模式", level=2)
    add_body(
        document,
        "开发模式适合改 Vue 页面。后端仍在 8765 端口运行，Vite 在 5173 端口提供热更新"
        "页面，并把 /api 请求代理给后端。开发时访问 5173，而不是 8765；演示最终效果时"
        "应重新构建并访问 8765。",
    )
    add_code(
        document,
        r"""# 终端 1：项目根目录
& "D:\Work_Software\Anaconda3\envs\nlp_env\python.exe" -m backend.run

# 终端 2：frontend 目录
cd "D:\Project_Code_2026\tyut_comprehensive_pratice\smartspeech_optimized_new\frontend"
npm run dev""",
    )
    add_note(
        document,
        "不要重复启动后端",
        "如果 8765 已被本系统占用，再启动一个后端会报端口冲突。先访问健康检查地址；"
        "若已经返回正常，直接打开页面即可。",
    )

    document.add_heading("7.6 .env 配置项逐项说明", level=2)
    add_table(
        document,
        ("变量", "默认值", "白话解释", "常见调整"),
        (
            ("DEEPSEEK_API_KEY", "空", "访问 DeepSeek 的私钥", "必须填写，不能提交到仓库"),
            ("DEEPSEEK_API_URL", "官方 chat/completions", "大模型接口地址", "代理服务确有兼容接口时再改"),
            ("DEEPSEEK_MODEL", "deepseek-chat", "当前使用的对话模型", "按账号可用模型调整"),
            ("DEEPSEEK_TIMEOUT", "90", "大模型最长等待秒数", "网络慢可适度增大"),
            ("CHAT_CONTEXT_MESSAGES", "20", "追问携带的最近消息条数", "增大更连贯但更慢更耗 token"),
            ("WHISPER_MODEL_SIZE", "small", "语音识别模型规格", "轻量电脑改 tiny 或 base"),
            ("WHISPER_DEVICE", "cpu", "Whisper 运行设备", "有兼容 CUDA 环境可改 cuda"),
            ("DEFAULT_TTS_VOICE", "zh-CN-YunxiNeural", "默认面试官音色", "使用 Edge-TTS 可用音色名"),
            ("TTS_TIMEOUT_SECONDS", "35", "在线语音合成超时", "网络不稳定时可增大"),
            ("DEFAULT_VISION_FPS", "3", "默认分析帧率", "普通电脑建议 2 或 3"),
            ("MAX_VISION_FPS", "8", "后端允许的最大帧率", "不建议随意提高"),
            ("SNAPSHOT_INTERVAL_SECONDS", "2", "趋势快照间隔", "变大可减少报告数据量"),
            ("FACE_MATCH_THRESHOLD", "1.0", "人脸欧氏距离阈值", "小更严格，建议 0.9-1.1"),
            ("APP_HOST", "127.0.0.1", "服务监听地址", "仅本机访问保持默认"),
            ("APP_PORT", "8765", "系统访问端口", "冲突时改端口"),
            ("FLASK_DEBUG", "0", "是否启用调试模式", "日常使用保持 0"),
            ("SESSION_COOKIE_SECURE", "0", "Cookie 是否仅走 HTTPS", "公网 HTTPS 部署改 1"),
            ("FLASK_SECRET_KEY", "自动生成", "会话签名密钥", "生产环境显式设置随机值"),
            ("CORS_ORIGINS", "两个 5173 地址", "允许开发前端跨域访问的来源", "仅新增可信地址"),
        ),
        widths=(4.6, 3.0, 4.7, 4.0),
    )
    add_body(
        document,
        "修改 .env 后必须重启后端才会生效。不要把密钥写在 Python 源码里，也不要为了"
        "方便把 API Key 填进 README、答辩截图或课堂群聊。若密钥曾公开，应立即在平台"
        "后台撤销并重新生成。",
    )

    document.add_heading("7.7 启动后如何确认真的正常", level=2)
    add_steps(
        document,
        (
            "浏览器打开 http://127.0.0.1:8765，应看到登录界面而不是 Flask 默认错误页。",
            "访问 /api/health，应返回可解析的 JSON，表示后端进程正在监听。",
            "注册或登录，确认会话 Cookie 建立，页面能进入面试配置。",
            "允许摄像头权限，开始面试后应看到本地视频和第一道问题。",
            "说一段短回答，确认转写文字出现；问题可播放时再检查 TTS。",
            "观察有效 FPS 和延迟是否持续更新，而不是固定在零。",
            "完成至少一次回答并结束面试，确认生成报告且历史列表可见。",
        ),
    )
    add_code(
        document,
        r"""# PowerShell 健康检查
Invoke-RestMethod -Uri "http://127.0.0.1:8765/api/health" | ConvertTo-Json -Depth 5

# 检查 Python 语法
& "D:\Work_Software\Anaconda3\envs\nlp_env\python.exe" -m compileall backend

# 检查 Python 依赖冲突
& "D:\Work_Software\Anaconda3\envs\nlp_env\python.exe" -m pip check""",
    )

    document.add_heading("7.8 怎样停止、重启和换端口", level=2)
    add_bullets(
        document,
        (
            "停止：关闭运行 backend.run 的命令行窗口；一键启动时关闭启动窗口。",
            "重启：关闭旧窗口后再启动，避免同时运行两个 Waitress 实例。",
            "换端口：修改 .env 的 APP_PORT，重启后端，并改用新端口访问。",
            "Vite 代理：开发和一键脚本都会读取端口配置，但手工启动时要确认代理地址同步。",
            "清理缓存不是常规操作：TTS 超时文件会自行清理，不要随意删除 data/auth。",
        ),
    )
    add_note(
        document,
        "备份顺序",
        "升级或改模型前先备份 .env、data/auth 和 runtime/reports。要迁移到另一台电脑时，"
        "同时备份项目源码、前端构建结果和这些本地数据；不要把私密备份上传到公开网盘。",
    )

    document.add_heading("7.9 安装时报错时的快速判断", level=2)
    add_table(
        document,
        ("报错关键词", "通常原因", "先做什么"),
        (
            ("python.exe 找不到", "解释器路径变化或环境未创建", "修改 PYTHON_EXE，或创建 Python 3.10 环境"),
            ("No module named", "Python 依赖不完整", "在 nlp_env 执行 pip install -r requirements.txt"),
            ("npm 不是内部命令", "Node.js 未安装或 PATH 未刷新", "安装 Node.js 20+ 后重开终端"),
            ("npm ci 失败", "网络、缓存或 lock 文件不匹配", "检查网络，删除损坏的 node_modules 后重试"),
            ("NumPy / ABI", "错误升级到 NumPy 2.x", "按 requirements.txt 回装 NumPy 1.26.4"),
            ("端口已被占用", "已有后端或其他程序占用", "打开现有页面，或修改 APP_PORT"),
            ("前端尚未构建", "dist 缺失或构建失败", "在 frontend 执行 npm run build"),
        ),
        widths=(4.3, 5.3, 6.7),
    )


def add_chapter_8_operation(document: Document) -> None:
    add_section_break(document)
    document.add_heading("第 8 章 操作教程与真实界面解析", level=1)
    document.add_heading("8.1 一次完整面试的最短路径", level=2)
    add_steps(
        document,
        (
            "打开系统，注册账号；可同时录入人脸，也可登录后再绑定。",
            "使用密码或人脸登录，进入“面试工作台”。",
            "选择岗位、难度、TTS 音色、2/3/4 帧每秒和是否自动朗读。",
            "点击“开始面试并打开摄像头”，在浏览器弹窗中选择“允许”。",
            "保持正对镜头，等待第一题生成；若自动播放被拦截，点击问题旁喇叭。",
            "输入文字或录音回答；面试官会结合最近回答继续追问。",
            "训练足够轮次后点击“结束面试”，等待大模型评估和视频汇总完成。",
            "在报告详情查看分数、逐题复盘、趋势和事件；也可进入历史报告再次查看。",
        ),
    )
    add_note(
        document,
        "本章图片说明",
        "下面 7 张图片均来自项目根目录“截图.docx”的真实运行界面。截图只用于展示操作"
        "路径和界面含义，不代表对截图中账号或候选人作能力评价。",
    )

    with ZipFile(SCREENSHOT_SOURCE) as archive:
        add_screenshot(document, archive, "image1.png", 1, "登录与认证")
        add_body(
            document,
            "用户看到什么：左侧介绍系统的四类核心能力，右侧可以在“密码登录、人脸登录、"
            "注册”之间切换。密码登录填写用户名和密码；人脸登录会打开摄像头；注册时还"
            "可以勾选注册后录入人脸。界面下方的技术标签与系统真实链路一致。",
        )
        add_body(
            document,
            "后台在做什么：密码登录会把输入密码与 users.json 中的 PBKDF2 摘要比较；"
            "人脸登录会提取当前 512 维向量，与已登记向量计算欧氏距离；注册会先检查"
            "用户名唯一性，再保存哈希和可选人脸向量。成功后会建立约 12 小时有效会话。",
        )
        add_body(
            document,
            "对应原理与阅读要点：这不是普通的“账号密码登录页 + 人脸按钮”，而是两套本地"
            "认证链路。人脸原始照片不会保存，只保存向量；相机光线、正脸角度和遮挡会"
            "直接影响登记与匹配效果。",
        )

        add_section_break(document)
        add_screenshot(document, archive, "image2.png", 2, "面试配置首页")
        add_body(
            document,
            "用户看到什么：左侧选择目标岗位、初级到专家难度、Edge-TTS 音色、画面采样"
            "频率和自动朗读开关；右侧直接显示最终得分公式与数据隐私说明。默认是 Python "
            "后端开发工程师、中级、3 帧每秒和自动朗读。",
        )
        add_body(
            document,
            "后台在做什么：点击开始后，前端调用 POST /api/interviews，后端把岗位、难度、"
            "音色、采样率和快照间隔写入 InterviewSession。系统先返回会话，再异步生成"
            "第一题，避免大模型响应期间页面一直空白。",
        )
        add_body(
            document,
            "对应原理与阅读要点：候选人的回答质量占 80%，表情与注意力各占视频分的 50%，"
            "视频只占最终分的 20%。帧率越高趋势越细，但 CPU 压力也越大；普通电脑先选"
            " 2 或 3 帧每秒，稳定后再提高。",
        )

        add_section_break(document)
        add_screenshot(document, archive, "image3.png", 3, "面试中实时视觉看板")
        add_body(
            document,
            "用户看到什么：左侧是本地摄像头画面，右侧显示人脸检测、当前表情、视线、"
            "头部姿态、注意力准备状态、实时视频状态分、有效 FPS 和最近趋势。画面底栏"
            "明确提示“本地画面仅用于实时分析”。",
        )
        add_body(
            document,
            "后台在做什么：浏览器把 JPEG 帧提交到本机后端，视觉工作线程调用 MediaPipe "
            "Face Mesh，一次得到的关键点同时交给表情分析器和注意力分析器。实时分数"
            "随后通过指标接口返回页面。",
        )
        add_body(
            document,
            "对应原理与阅读要点：需要先完成 5 帧预热和 30 帧个人校准，系统才能形成"
            "“中性脸”基线。界面出现“准备中”是正常过程。头姿、视线和表情都是趋势信号，"
            "不要把一次低分理解为能力或人格判断。",
        )

        add_section_break(document)
        add_screenshot(document, archive, "image4.png", 4, "问答、趋势与动态变化")
        add_body(
            document,
            "用户看到什么：左侧是面试官问题和候选人的回答；右侧持续更新表情分、注意力"
            "分、表情变化次数、注意力变化次数、分散次数、眨眼次数和三项头部角度。"
            "回答框在底部，可输入文字，也可点击麦克风录音。",
        )
        add_body(
            document,
            "后台在做什么：每次回答都会保存到当前会话 transcript，DeepSeek 只读取最近"
            " 20 条有效上下文生成下一问；右侧趋势按约 2 秒间隔生成快照，结束面试后再"
            "汇总为最终视频指标。",
        )
        add_body(
            document,
            "对应原理与阅读要点：追问不只是“换一道题”，它会优先围绕上一轮的项目、职责、"
            "技术取舍和数据结果继续深挖。回答越具体，下一问越有针对性；过于笼统时模型"
            "会要求补充事实依据。"
        )

        add_section_break(document)
        add_screenshot(document, archive, "image5.png", 5, "历史报告")
        add_body(
            document,
            "用户看到什么：页面顶部显示已完成面试数、最终得分平均分和最高分；下方每条"
            "记录包含岗位、难度、最终分、问答分、表情分、注意力分、日期和时长，并提供"
            "“查看报告”按钮。",
        )
        add_body(
            document,
            "后台在做什么：页面调用 GET /api/reports，ReportStore 只读取当前登录用户的"
            "目录。报告来自 runtime/reports/<用户编号>，因此不同账号之间不会在正常"
            "接口路径下看到彼此记录。",
        )
        add_body(
            document,
            "对应原理与阅读要点：历史分适合比较同一岗位、相近难度和相同采样设置下的"
            "变化。不同岗位或不同大模型版本之间的分数不能简单横比，多次练习的逐题建议"
            "更有训练价值。",
        )

        add_section_break(document)
        add_screenshot(document, archive, "image6.png", 6, "报告详情：评分与综合评估")
        add_body(
            document,
            "用户看到什么：顶部直接展示评分计算式，截图中问答 68.0、表情 50.8、注意力"
            " 74.6、视频综合 62.7、最终 66.9。下方还能看到大模型总结和专业能力等维度，"
            "右侧展示分析帧、人脸帧、覆盖率和采样统计。",
        )
        add_body(
            document,
            "后台在做什么：结束时先把完整问答记录交给 DeepSeek 做 JSON 评估，再把本机"
            "视觉汇总交给 score_service。最终分严格按 68.0 × 0.8 + (50.8 × 0.5 + 74.6 "
            "× 0.5) × 0.2 = 66.9 计算，使用 Decimal 和四舍五入保留 1 位小数。",
        )
        add_body(
            document,
            "对应原理与阅读要点：这份样例的问答分并不高，主要反映第三问对幂等、数据库"
            "约束和补偿方案的表达不完整；注意力分高于表情分，说明内容短板和过程状态"
            "是两回事。复盘时应先看具体回答，再看视频趋势。",
        )

        add_section_break(document)
        add_screenshot(document, archive, "image7.png", 7, "报告详情：逐题复盘与事件")
        add_body(
            document,
            "用户看到什么：左侧包含后续训练建议和逐题复盘，每题分别列出回答表现与改进"
            "建议；右侧是状态时间线，可看到注意力集中/分散、人脸离开/回到画面和时间；"
            "底部隐私卡片再次说明视频、人脸照片和录音不保存。",
        )
        add_body(
            document,
            "后台在做什么：逐题内容来自 DeepSeek 评估的 question_reviews；事件来自本机"
            "视觉分析的状态变化记录；隐私卡片来自报告 privacy 字段。页面支持浏览器"
            "打印或导出 PDF。",
        )
        add_body(
            document,
            "对应原理与阅读要点：报告的目的不是给一个永久标签，而是把问题定位到下一次"
            "训练动作。优先选择一到三条建议，在下一次面试中练习完整技术解释、数据支撑"
            "和清晰收束，再比较报告变化。",
        )

    document.add_heading("8.9 注册和人脸绑定怎样提高成功率", level=2)
    add_steps(
        document,
        (
            "选择光线均匀的位置，避免背光和强逆光，脸不要一半亮一半暗。",
            "让脸部占据画面中央，保持正脸、睁眼、无口罩和大面积头发遮挡。",
            "拍摄时保持自然表情，不要离摄像头过近或过远，也不要快速晃动。",
            "人脸绑定后先用同一环境试一次人脸登录，确认匹配阈值适合当前摄像头。",
            "如频繁失败，先改善光线并重新绑定；仍失败时再小幅调整阈值。",
        ),
    )
    add_note(
        document,
        "人脸登录不是万能钥匙",
        "双胞胎、化妆、年龄变化、极端光照和摄像头差异都可能影响特征距离。日常训练"
        "优先使用密码登录；人脸登录用于展示本地生物特征匹配链路。",
    )

    document.add_heading("8.10 面试中怎样回答更容易获得有效反馈", level=2)
    add_table(
        document,
        ("回答动作", "推荐做法", "不推荐"),
        (
            ("结构", "先结论，再背景、行动、结果；可用 STAR 组织项目题", "从无关背景讲很久才回答核心"),
            ("技术细节", "说明方案、取舍、数据、异常处理和验证方式", "只报技术名词，不解释为什么"),
            ("个人贡献", "明确“我负责什么、如何做、产生什么结果”", "一直使用“我们”，看不出个人作用"),
            ("录音", "一次说清一个完整回答，语速平稳，减少环境噪声", "频繁打断、边想边多次重录"),
            ("摄像头位置", "摄像头与视线大致同高，保持在画面中央", "低头看键盘、侧脸或频繁离屏"),
            ("结束时机", "至少完成 2-3 轮有效问答，视频覆盖稳定后结束", "刚回答一题就结束或长时间只显示画面"),
        ),
        widths=(3.1, 7.5, 5.7),
    )

    document.add_heading("8.11 演讲式操作：从开始到结束", level=2)
    add_steps(
        document,
        (
            "准备：关闭会议软件、录屏软件和会占用摄像头的程序，清理桌面隐私画面。",
            "登录：优先使用已绑定账号；出现权限提示时选择允许摄像头和麦克风。",
            "配置：首次先选 2 帧每秒，确认系统流畅后再按需要提高。",
            "开场：等待问题出现；若未自动播放，点击喇叭，不要重复点击开始面试。",
            "回答：文字回答适合技术细节，语音回答适合模拟真实交流；提交前可修正转写。",
            "观察：只把 FPS、延迟、人脸状态当作链路健康提示，不要一直盯分数影响表达。",
            "结束：点击结束并耐心等待。大模型评估可能持续数十秒，不要中途关闭窗口。",
            "复盘：先读逐题建议，再结合事件时间线和隐私卡片决定下一轮练习重点。",
        ),
    )

    document.add_heading("8.12 报告如何正确阅读", level=2)
    add_table(
        document,
        ("报告区块", "回答的问题", "阅读顺序"),
        (
            ("最终分", "这次训练的总体阶段结果是多少", "先知道结果，但不只看数字"),
            ("问答分", "文字内容与岗位要求匹配得怎样", "结合逐题复盘判断原因"),
            ("表情分", "相对个人基线的积极状态是否稳定", "看趋势，不追求一直微笑"),
            ("注意力分", "头部、视线和覆盖率的综合状态", "关注持续分散或长期离屏"),
            ("视频综合", "表情与注意力各半融合后的过程分", "理解 20% 权重"),
            ("亮点与短板", "做对了什么，下一轮优先改什么", "选择 1-3 条可执行建议"),
            ("逐题复盘", "每个问题的表现与建议", "最值得反复阅读"),
            ("事件时间线", "什么时候离屏、分散或恢复", "结合当时动作解释"),
            ("隐私说明", "哪些原始数据未保存", "确认训练边界"),
        ),
        widths=(3.4, 7.1, 5.8),
    )
    add_body(
        document,
        "点击“打印 / 导出 PDF”会调用浏览器打印界面。选择“另存为 PDF”即可得到本地"
        "报告文件。打印前建议展开需要保留的问答内容；不同浏览器的分页位置可能不同，"
        "这是网页打印特性，不是报告内容丢失。",
    )
    add_note(
        document,
        "训练建议",
        "同一岗位至少练三次：第一次熟悉流程，第二次针对内容短板，第三次针对表达和"
        "过程稳定性。只有控制岗位、难度和帧率相对一致，分数变化才更有参考意义。",
    )


def add_chapter_9_training(document: Document) -> None:
    add_section_break(document)
    document.add_heading("第 9 章 训练、微调与调优指南", level=1)
    document.add_heading("9.1 先给出结论：当前系统不需要训练", level=2)
    add_body(
        document,
        "当前项目使用已经训练好的模型，并通过本地规则和 API 提示词完成任务。只要模型"
        "文件能够加载、DeepSeek Key 有效，就可以直接运行。不要因为系统里有“识别人脸、"
        "表情和语音”就误以为必须自己准备数据集训练；从零训练这些模型通常需要大量数据、"
        "显卡和评测工作，远超普通课程项目的最小闭环。",
    )
    add_note(
        document,
        "正确调优顺序",
        "先确认输入质量，再调配置阈值，再改提示词和规则，最后才考虑训练或替换模型。"
        "大量所谓“模型不准”其实是光线差、摄像头位置偏、录音噪声大或输入格式不合规。",
    )

    document.add_heading("9.2 五种“调优”分别改什么", level=2)
    add_table(
        document,
        ("层级", "改动对象", "是否需要数据训练", "典型目的"),
        (
            ("配置调优", ".env、页面帧率、上下文条数", "否", "控制速度、成本、严格度和体验"),
            ("提示词调优", "InterviewService 两套 system prompt", "否", "改变提问风格、评分视角和输出结构"),
            ("阈值规则调优", "表情、注意力、人脸匹配参数", "否", "适配具体摄像头、人群和场景"),
            ("轻量模型训练", "基于现有关键点训练表情分类器", "是", "提高表情覆盖和抗干扰能力"),
            ("替换/微调大模型", "人脸、Whisper 或注意力模型", "是，且成本最高", "解决通用模型无法覆盖的领域问题"),
        ),
        widths=(3.3, 5.9, 3.6, 3.5),
    )

    document.add_heading("9.3 DeepSeek：不训练权重，主要改提示词", level=2)
    add_body(
        document,
        "DeepSeek 的权重运行在服务商侧，本项目只发送 messages 并接收文本。因此不能也不"
        "需要在本地“训练 DeepSeek”。要改变它，最有效的方法是修改 InterviewService 中"
        "的 INTERVIEW_SYSTEM_PROMPT 和 EVALUATION_SYSTEM_PROMPT，再修改 reply() 与 "
        "evaluate() 里的温度、上下文范围和输出规范。",
    )
    add_table(
        document,
        ("想达到的效果", "建议修改", "注意事项"),
        (
            ("问题更像某岗位", "在 system prompt 增加岗位能力清单和提问顺序", "不要堆过长规则，否则占用上下文"),
            ("追问更深入", "强调追问事实、数据、取舍、异常和个人贡献", "避免连续追问同一细节"),
            ("更符合校招", "增加项目基础、学习潜力和表达结构权重", "不要使用未经提供的个人信息"),
            ("评分更严格", "明确评分锚点、扣分条件和证据要求", "必须用固定样例回归，防止分数整体漂移"),
            ("输出更稳定", "保留 JSON 结构，继续使用 chat_json 和低温度", "前端字段名不能随意改"),
            ("减少 token", "降低 CHAT_CONTEXT_MESSAGES 和 max_tokens", "上下文太短会影响连续性"),
        ),
        widths=(4.2, 6.3, 5.8),
    )
    add_body(
        document,
        "提示词调优要遵守“一次只改一类规则”。例如先只改专业能力评分标准，固定 5-10 份"
        "历史回答作为回归集，对比修改前后的分数、理由和 JSON 完整性。一次同时修改模型、"
        "温度、上下文和评分尺度，无法判断到底是哪项导致变化。",
    )
    add_code(
        document,
        "主改文件：backend/app/services/interview_service.py\n"
        "接口客户端：backend/app/services/deepseek_service.py\n"
        "环境参数：.env 中的 DEEPSEEK_MODEL / DEEPSEEK_TIMEOUT / CHAT_CONTEXT_MESSAGES\n"
        "回归样例：建议从 runtime/reports 中挑选固定报告，先备份再测试",
    )

    document.add_heading("9.4 人脸识别：先调阈值，再考虑换模型", level=2)
    document.add_heading("9.4.1 普通调优", level=3)
    add_body(
        document,
        "人脸登录把两张脸的 512 维向量做欧氏距离比较，越小越像。FACE_MATCH_THRESHOLD"
        " 默认 1.0。若本人经常登录失败，可在 1.0 到 1.1 之间小幅增加；若担心别人误"
        "通过，可在 1.0 到 0.9 之间减小。每次只改 0.05，并用多人样例测试，不能只测"
        "自己一个人。",
    )
    add_table(
        document,
        ("现象", "阈值方向", "风险"),
        (
            ("本人偶尔匹配失败", "略微增大", "误把相似的人认成同一账号"),
            ("不同人容易被误认", "略微减小", "本人失败率上升，需要更多尝试"),
            ("换摄像头后普遍失败", "先重新登记，再考虑调阈值", "直接放宽容错可能掩盖设备问题"),
            ("人脸检测不到", "不要只调距离阈值", "应先修复光线、遮挡或模型依赖"),
        ),
        widths=(5.1, 4.2, 7.0),
    )

    document.add_heading("9.4.2 为什么更换模型后必须重建向量", level=3)
    add_body(
        document,
        "每套人脸模型都有自己的“特征语言”。同样一张脸，VGGFace2 模型得到的 512 个"
        "数字与另一模型的输出不能直接比较，即使维度碰巧相同也可能没有意义。因此只要"
        "修改 MTCNN、InceptionResnetV1 的预训练权重、网络结构或归一化方式，就必须停止"
        "使用旧向量，并让所有用户重新绑定人脸。",
    )
    add_steps(
        document,
        (
            "备份并记录旧模型、旧向量和固定测试照片，确保实验可回退。",
            "在 face_recognition.py 中替换模型加载逻辑，保持 extract_face_vector 输出约定。",
            "用同一人的多张照片确认类内距离，用不同人的照片确认类间距离。",
            "按误接受率、误拒绝率、本人通过率和平均耗时重新选择阈值。",
            "迁移到新特征空间：清空旧向量后重新登记，或将旧照片重新提取向量。",
            "接口、报告和前端字段保持不变，避免登录链路因模型实验而全部重写。",
        ),
    )

    document.add_heading("9.4.3 如果确实要微调人脸模型", level=3)
    add_body(
        document,
        "常见做法是使用三元组损失或 ArcFace 类损失，让同一个人的向量更近、不同人的"
        "向量更远。训练数据必须按“人”划分训练集和验证集，不能把同一个人的照片同时"
        "放入两边，否则结果会虚高。普通课程项目通常没有授权充分的人脸数据，因此优先"
        "使用预训练模型更合理。涉及真人数据时必须取得授权，并设置保存期限和删除机制。",
    )
    add_table(
        document,
        ("训练要素", "建议"),
        (
            ("输入", "对齐后的人脸裁剪，至少 112×112 或 160×160"),
            ("标签", "身份编号，不是姓名；训练集和验证集按身份隔离"),
            ("损失", "Triplet Loss、ArcFace 或同类度量学习损失"),
            ("评估", "TAR@FAR、类内/类间距离分布、跨摄像头测试"),
            ("接入", "替换 face_recognition.py 的特征提取器，并重建全部向量"),
        ),
        widths=(4.0, 12.1),
    )

    document.add_heading("9.5 表情识别：当前是规则模型，怎样升级", level=2)
    add_body(
        document,
        "当前表情模块不是卷积神经网络分类器，而是从 Face Mesh 计算 EAR、嘴角、嘴部"
        "开度、眉毛高度和不对称等特征，再用个人基线和阈值判断六类状态。它可解释、轻量、"
        "容易调参，但对夸张表情、特殊脸型、戴眼镜或侧脸的覆盖能力有限。",
    )
    document.add_heading("9.5.1 先做规则评测，不急着训练", level=3)
    add_steps(
        document,
        (
            "录制固定测试集：每位参与者分别做中性、微笑、大笑、惊讶、困惑和说话动作。",
            "使用统一帧率和光照，记录人工标注，不从一段自然视频中随意猜标签。",
            "统计每类识别率、混淆矩阵和误报；重点观察“说话是否被误判为惊讶”等问题。",
            "只调整 expression_analyzer.py 中对应阈值，每次只改一个参数并保留测试记录。",
            "若规则无法同时平衡准确率和误报，再进入分类器训练阶段。",
        ),
    )
    document.add_heading("9.5.2 基于人脸关键点训练轻量分类器", level=3)
    add_body(
        document,
        "这是改造成本较低的路线：输入仍然使用现有几何特征，输出从固定 if/elif 规则改为"
        "可训练分类器。它可以继续在 CPU 上运行，也便于解释。推荐先尝试逻辑回归、随机"
        "森林或小型 MLP，确认数据质量后再上深度模型。",
    )
    add_code(
        document,
        "建议新增（当前仓库尚未提供）：\n"
        "training/expression/manifest.csv\n"
        "  字段：sample_id, subject_id, image_path, label, split\n\n"
        "training/train_expression.py\n"
        "  读取人脸关键点 -> 提取 31 维几何特征 -> 标准化 -> 训练分类器\n\n"
        "models/expression_classifier.pt\n"
        "  保存权重、特征顺序、标签顺序、训练版本和校验指标\n\n"
        "接入位置：backend/app/services/vision/expression_analyzer.py\n"
        "修改原则：保留 expression、confidence、expression_score 等现有输出字段",
    )
    add_bullets(
        document,
        (
            "每人每类至少采集数百到上千个有效帧，覆盖不同光照、脸型和轻微姿态变化。",
            "验证集按参与者隔离，不能把同一人的相邻帧分到训练和测试两边。",
            "类别不平衡时使用 class_weight 或重采样，不能只看总体准确率。",
            "以宏平均 F1、每类召回率和混淆矩阵为主要指标，不能只追求“中性”准确率。",
            "训练时保存标准化参数和标签映射，推理时必须使用同一顺序。",
        ),
    )
    document.add_heading("9.5.3 基于人脸图像训练 CNN", level=3)
    add_body(
        document,
        "如果关键点本身不足以表达复杂表情，可以训练图像分类网络，例如 MobileNetV3、"
        "ResNet18 或 EfficientNet-B0。公开表情数据集可用来预训练，但必须检查授权、"
        "标签质量和采集偏差。训练前统一人脸裁剪、颜色空间和输入尺寸；推理时可以在"
        " Face Mesh 检出的区域内运行分类器，避免背景干扰。",
    )
    add_note(
        document,
        "不要直接套用数据集标签",
        "公开数据集的“愤怒、恐惧、悲伤”等标签不等同于面试训练需要的“困惑、说话”等"
        "业务状态。若要用于本项目，应先定义自己的标签规范并由多人独立标注，计算标注"
        "一致性后再训练。",
    )

    document.add_heading("9.6 注意力监测：训练视线模型需要哪些数据", level=2)
    add_body(
        document,
        "当前注意力模块由 solvePnP 头姿、虹膜视线和覆盖率规则组成，不需要训练。它适合"
        "判断明显偏转和长期离屏，但无法精确到屏幕上的某个像素点。如果要把“是否看镜头”"
        "升级为真实屏幕注视区域估计，可训练的替代方向包括视线回归和视线分类。",
    )
    add_table(
        document,
        ("方向", "输入", "标签", "输出", "参考数据集/方法思路"),
        (
            ("头姿估计", "人脸图像或关键点", "yaw、pitch、roll", "三个角度", "300W-LP 等带姿态数据"),
            ("二维视线回归", "双眼裁剪 + 人脸姿态", "归一化屏幕坐标或角度", "gaze_x、gaze_y", "MPIIGaze、Gaze360 思路"),
            ("屏幕区域分类", "整脸图像", "屏幕上的区域编号", "区域概率", "按具体设备采集校准数据"),
            ("时序注意力", "连续帧头姿和视线", "人工标注的集中/分散区间", "时序状态", "LSTM、TCN 或轻量 Transformer"),
        ),
        widths=(3.0, 3.6, 3.8, 3.0, 3.0),
    )
    add_steps(
        document,
        (
            "明确预测目标：是把人分成看/不看，还是预测屏幕分区；目标不同，数据协议不同。",
            "采集数据时记录摄像头位置、屏幕尺寸、光照、是否戴眼镜和每个人脸型。",
            "按参与者划分训练与测试集，禁止同一人的相邻帧跨集合泄漏。",
            "先建立简单基线，再训练模型；如果规则模型已达到业务要求，不必为了复杂而复杂。",
            "用角度误差、区域准确率、跨人泛化、延迟和校准耗时综合评估。",
            "接入时修改 attention_analyzer.py，但保持 head pose、gaze 和 score 字段兼容。",
        ),
    )

    document.add_heading("9.7 Whisper：通常只需换规格，领域微调成本较高", level=2)
    document.add_heading("9.7.1 最实用的调优", level=3)
    add_table(
        document,
        ("规格", "速度", "准确率倾向", "使用建议"),
        (
            ("tiny", "最快", "最低", "性能很弱或只做流程演示"),
            ("base", "较快", "一般", "普通 CPU 电脑、短句回答"),
            ("small", "中等", "较好", "当前默认，均衡选择"),
            ("medium", "慢", "更好", "有较强 CPU/GPU 且可等待"),
            ("large", "最慢", "通常最好", "需要 GPU，不适合普通演示机"),
        ),
        widths=(2.7, 2.8, 3.7, 6.9),
    )
    add_body(
        document,
        "在 .env 中修改 WHISPER_MODEL_SIZE 后重启即可切换。首次使用新规格会下载模型。"
        "如果 WHISPER_DEVICE=cpu，代码会把 fp16 设为 false；只有明确使用 CUDA 时才应"
        "改成 cuda，并确保 PyTorch 与该显卡环境兼容。",
    )

    document.add_heading("9.7.2 Whisper 领域微调需要的数据", level=3)
    add_body(
        document,
        "当通用 Whisper 在专业术语、方言、远场录音或特定课程词汇上持续出错时，才考虑"
        "微调。数据应当是干净音频与逐字文本成对出现，不能只提供会议录像和一段粗略"
        "总结。训练集、验证集和测试集要按说话人或会议分组，避免同一段语音泄漏。",
    )
    add_table(
        document,
        ("数据要求", "建议标准"),
        (
            ("音频格式", "统一为单声道、16 kHz WAV；去除长静音和爆音"),
            ("时长", "每条尽量在 Whisper 的 30 秒窗口内，长音频先切段"),
            ("文本", "逐字转写，统一标点和数字规范，专业术语建立词表"),
            ("规模", "至少数小时到数十小时高质量数据；具体取决于领域差异"),
            ("划分", "按说话人/会议划分 train、validation、test，不能随机切相邻片段"),
            ("基线", "先测原始 small 模型的 WER/CER，确认微调是否真的有收益"),
        ),
        widths=(4.1, 12.0),
    )
    add_body(
        document,
        "常见训练路线是使用 Hugging Face Transformers 的 Whisper 微调流程或社区训练"
        "工具。训练会输出声学模型和 tokenizer/processor 配置。推理时应把完整的微调"
        "目录作为一个模型加载，不能只替换一个权重文件。当前 ASRService 通过"
        " whisper.load_model(model_size, device=device) 加载；接入微调模型时，应在"
        " asr_service.py 中扩展加载路径，并在 .env 增加独占的 ASR_MODEL_PATH，而不是"
        "把大体积 checkpoint 提交到 Git。",
    )
    add_code(
        document,
        "主改文件：backend/app/services/asr_service.py\n"
        "新增配置：建议 .env 增加 ASR_MODEL_PATH=...\n"
        "输出目录：建议 models/whisper_interview_zh/（加入 .gitignore）\n"
        "验收指标：中文 CER、专业词错误率、平均转写耗时、首次加载耗时",
    )
    add_note(
        document,
        "CPU 不适合训练 Whisper",
        "Whisper 微调通常需要 GPU 显存与较长时间。普通电脑建议只切换 tiny/base/small"
        " 规格，或用领域术语表做文本后处理；不要把漫长训练当作系统正常启动的必要步骤。",
    )

    document.add_heading("9.8 Edge-TTS：可以换音色，不能训练语音权重", level=2)
    add_body(
        document,
        "Edge-TTS 是在线语音合成服务，项目通过音色名请求音频，不拥有模型权重，也没有"
        "本地训练接口。最直接的调整是把 DEFAULT_TTS_VOICE 换成其他可用中文音色。页面"
        "也会提供音色选项。若需要定制品牌音色，要使用具备授权的云语音定制服务，不能"
        "用 Edge-TTS 冒充语音克隆。",
    )
    add_table(
        document,
        ("需求", "可做方式", "不建议方式"),
        (
            ("换男声/女声", "修改 DEFAULT_TTS_VOICE 或页面选项", "改源码硬编码"),
            ("语速和停顿", "按 Edge-TTS 协议扩展 SSML 或参数", "简单全局加速导致听不清"),
            ("断网可用", "保留 Windows SAPI 回退", "假设在线服务永远可用"),
            ("定制个人音色", "使用合规语音定制服务并取得授权", "未经本人同意克隆声音"),
        ),
        widths=(3.6, 7.5, 5.1),
    )

    document.add_heading("9.9 训练前必须做好的数据准备", level=2)
    add_body(
        document,
        "数据质量比模型结构更决定结果。训练前应建立数据卡，写清采集对象、采集设备、"
        "场景、授权、标签定义、脱敏方式、训练与验证划分和已知偏差。人脸、语音和视频"
        "属于敏感个人信息，即使是课程实验，也要取得参与者同意并尽量使用合成或公开"
        "授权数据。",
    )
    add_table(
        document,
        ("检查项", "合格标准"),
        (
            ("授权", "参与者知道用途、保存范围、训练方式和删除渠道"),
            ("标签", "标签定义明确，至少抽检一致性，错误样本有修订记录"),
            ("切分", "按人或会话隔离训练/验证/测试，杜绝数据泄漏"),
            ("平衡", "各类别、性别、光照、设备尽量覆盖，不靠单一样本堆量"),
            ("隐私", "去除姓名、学号、屏幕内容、背景人脸和可识别音频"),
            ("版本", "记录数据版本、预处理脚本、随机种子和模型版本"),
            ("评测", "既有总体指标，也按关键子群和困难样本评估"),
        ),
        widths=(3.3, 12.8),
    )

    document.add_heading("9.10 修改评分权重时怎样回归", level=2)
    add_steps(
        document,
        (
            "备份 runtime/reports 中至少三份不同水平的历史报告，不覆盖原始文件。",
            "列出旧公式、新公式和预期变化，例如问答是否仍占主导。",
            "先修改 score_service.py，再同步前端公式文案、README 和本文档。",
            "运行 backend/tests/test_score_service.py，检查权重和不为 1、视频无效等边界。",
            "用历史报告重算，确认四舍五入、视频缺失、0 分与 100 分边界正确。",
            "更新报告样例，避免截图仍展示旧公式；最后至少完成一场端到端面试。",
        ),
    )
    add_code(
        document,
        "当前公式（不要只改其中一处）：\n"
        "视频综合 = 表情 × 0.5 + 注意力 × 0.5\n"
        "最终分 = 问答 × 0.8 + 视频综合 × 0.2\n\n"
        "必须同步：\n"
        "backend/app/services/score_service.py\n"
        "frontend/src/pages/InterviewPage.vue\n"
        "frontend/src/pages/ReportDetailPage.vue\n"
        "backend/tests/test_score_service.py\n"
        "README.md 与本文档",
    )
    add_note(
        document,
        "不能只看平均分",
        "调参后如果平均分上升，可能是整体变宽松，也可能是高分组被过度放大。至少检查"
        "分档分布、逐题理由和固定样例的一致性，不能以“分数更好看”作为训练成功标准。",
    )


def add_chapter_10_troubleshooting(document: Document) -> None:
    add_section_break(document)
    document.add_heading("第 10 章 故障排查手册", level=1)
    document.add_heading("10.1 先按链路定位，不要盲目重装", level=2)
    add_steps(
        document,
        (
            "确认后端是否已启动：访问 http://127.0.0.1:8765/api/health。",
            "确认页面来自正确地址：日常使用 8765，Vite 开发使用 5173。",
            "打开浏览器开发者工具，看 Console 是否有 JavaScript 错误，Network 中哪个请求失败。",
            "查看启动后端/PowerShell 窗口的日志，记录 HTTP 状态码和服务名称。",
            "判断故障属于 API、摄像头、麦克风、模型、网络还是本地数据，再按对应小节处理。",
            "只改一个变量并重启验证，不要同时改端口、模型、阈值和源码。",
        ),
    )
    add_note(
        document,
        "先保护数据",
        "排障时不要随意删除 data/auth 或 runtime/reports。绝大多数启动问题与这两类"
        "业务数据无关；删除它们只会造成账号、人脸向量或报告永久丢失。",
    )

    document.add_heading("10.2 系统启动问题", level=2)
    add_table(
        document,
        ("现象", "可能原因", "处理方式"),
        (
            ("双击脚本窗口一闪而过", "路径错误、依赖或端口检查提前失败", "在 PowerShell 手工运行 start_system.bat 查看完整信息"),
            ("Python 找不到", "解释器路径不存在", "恢复 nlp_env 或改脚本中的 PYTHON_EXE"),
            ("npm 找不到", "Node.js 未安装或 PATH 未刷新", "安装 Node.js 20+，重开终端"),
            ("依赖安装很慢", "PyTorch 等包体积大或网络慢", "保持网络稳定，必要时配置可信镜像"),
            ("端口 8765 已占用", "已有后端进程", "访问现有页面，或修改 APP_PORT 后重启"),
            ("前端尚未构建", "dist 不存在", "在 frontend 执行 npm run build"),
            ("页面资源 404", "构建失败或访问了旧缓存", "重新构建，强制刷新浏览器 Ctrl+F5"),
        ),
        widths=(4.1, 5.1, 7.0),
    )
    add_code(
        document,
        r"""# 查看端口占用
Get-NetTCPConnection -State Listen -LocalPort 8765 |
  Select-Object LocalAddress,LocalPort,OwningProcess

# 健康检查
Invoke-RestMethod "http://127.0.0.1:8765/api/health"

# 检查前端构建文件
Test-Path "frontend\dist\index.html"

# 检查关键 Python 包
& "D:\Work_Software\Anaconda3\envs\nlp_env\python.exe" -c "import flask, waitress, cv2, mediapipe, whisper; print('ok')" """,
    )

    document.add_heading("10.3 登录与人脸问题", level=2)
    add_table(
        document,
        ("现象", "原因判断", "处理方式"),
        (
            ("注册提示用户名已存在", "用户名标准化后重复", "换用户名，或使用原账号登录"),
            ("密码登录失败", "密码记错、账号不一致、Cookie 被清", "核对用户名和密码，必要时重新注册"),
            ("摄像头没有画面", "权限、设备占用或 HTTPS 限制", "允许权限，关闭会议/录屏软件，使用本机地址"),
            ("人脸登记检测不到", "光线、遮挡或图像质量差", "正脸居中，增加均匀光线，重新拍摄"),
            ("本人人脸登录失败", "登记环境与登录环境差异大", "在同一环境重绑，再小幅调阈值"),
            ("别人可能被误认", "阈值过宽或人脸库样本少", "降低阈值，使用密码，补充严格测评"),
            ("重启后登录失效", "Session 失效或 secret 变化", "重新登录；确认 data/.flask_secret 未被删除"),
        ),
        widths=(4.0, 5.0, 7.2),
    )
    add_body(
        document,
        "浏览器只有在本机地址 http://127.0.0.1、localhost 或 HTTPS 页面中通常才允许"
        "持续访问摄像头。如果通过局域网 IP 访问却是 HTTP，浏览器可能直接阻止摄像头。"
        "此时应在受控环境配置 HTTPS，不能把浏览器安全限制误认为 Python 模型故障。",
    )

    document.add_heading("10.4 DeepSeek 问题", level=2)
    add_table(
        document,
        ("提示或状态", "含义", "处理方式"),
        (
            ("未配置 API Key", ".env 为空或未加载", "填写 DEEPSEEK_API_KEY 后重启后端"),
            ("HTTP 401/403", "密钥错误、失效或无权限", "重新生成并更换密钥"),
            ("HTTP 402", "账户余额或计费额度不足", "到开放平台检查额度"),
            ("HTTP 429", "请求过于频繁或达到限额", "稍后重试，降低并发和上下文长度"),
            ("请求超时", "网络慢或服务响应慢", "检查网络，可增大 DEEPSEEK_TIMEOUT"),
            ("返回无法解析 JSON", "评估输出格式异常", "重试；检查提示词与 chat_json 解析"),
            ("首题失败但仍有问题", "备用问题已启用", "连续追问仍要求 API 恢复"),
        ),
        widths=(4.2, 5.2, 6.8),
    )
    add_note(
        document,
        "不要把网络问题误判为算法问题",
        "首题有备用开场，所以 DeepSeek 故障时页面可能仍进入面试；真正提交回答后才会"
        "出现 502。判断依据应以后端日志和 Network 响应为准。",
    )

    document.add_heading("10.5 Whisper 与 TTS 问题", level=2)
    add_table(
        document,
        ("现象", "原因", "处理方式"),
        (
            ("第一次录音等待很久", "正在下载或加载 Whisper", "等待完成，后续会复用模型缓存"),
            ("内存不足", "模型规格过大", ".env 改为 base 或 tiny 并重启"),
            ("转写为空", "录音太短、静音或格式异常", "靠近麦克风，完整说一句后停止"),
            ("专业词经常错", "通用模型词汇限制", "切换规格、修正文本或开展领域微调"),
            ("没有自动声音", "浏览器阻止自动播放", "点击问题旁的喇叭按钮"),
            ("在线 TTS 失败后仍有声", "SAPI 回退生效", "属于正常降级；音质可能不同"),
            ("TTS 一直失败", "网络与 Windows 语音都不可用", "检查 edge-tts、网络和系统中文语音包"),
        ),
        widths=(4.1, 4.8, 7.3),
    )

    document.add_heading("10.6 视频分析卡顿或指标异常", level=2)
    add_table(
        document,
        ("现象", "解释", "处理方式"),
        (
            ("有效 FPS 低于设置值", "CPU 处理能力不足", "选 2 帧每秒，关闭高占用软件"),
            ("丢帧数持续增加", "新帧不断替换待处理旧帧", "说明系统优先保实时，通常不是故障"),
            ("分析延迟很高", "Face Mesh 或图像解码较慢", "降低帧率与分辨率，检查 CPU 占用"),
            ("一直显示校准中", "有效人脸帧不足", "正对镜头、改善光线，持续数秒"),
            ("人脸状态频繁跳变", "姿态、遮挡或光线不稳定", "稳定坐姿，观察趋势而非单帧"),
            ("注意力偏分散", "阈值、摄像头位置或视线偏移", "先将镜头与眼睛保持近似同高"),
            ("视频分数为空", "有效人脸帧少于最低要求", "保持画面覆盖后重新面试"),
        ),
        widths=(4.0, 5.4, 6.8),
    )
    add_body(
        document,
        "默认会发送 2/3/4 帧每秒中的一种，后端允许上限为 8，但提高上限不等于实际一定"
        "达到。有效 FPS 受浏览器截图、HTTP、JPEG 解码、Face Mesh 和 CPU 共同影响。"
        "面试训练优先保证稳定连续，不需要追求很高的帧率。",
    )

    document.add_heading("10.7 报告生成和查看问题", level=2)
    add_table(
        document,
        ("现象", "可能原因", "处理方式"),
        (
            ("结束按钮提示至少回答一次", "没有 user 消息", "先提交一轮文字或语音回答"),
            ("一直显示评估中", "DeepSeek 请求尚未返回", "等待超时；查看网络与后端日志"),
            ("评估失败且状态停止", "API 错误或 JSON 解析失败", "修复后重试；不要伪造成绩"),
            ("历史列表没有报告", "未成功 finish 或用户不匹配", "检查 runtime/reports 和登录账号"),
            ("报告打不开 JSON 错误", "文件损坏或版本不兼容", "从备份恢复，检查 ReportStore 读取"),
            ("打印分页不理想", "网页打印布局差异", "调整浏览器缩放或分页后另存 PDF"),
            ("删除后无法恢复", "删除会清理会话和报告", "重要报告先备份后再删除"),
        ),
        widths=(4.2, 4.9, 7.1),
    )

    document.add_heading("10.8 需要提供给维护者的信息", level=2)
    add_bullets(
        document,
        (
            "发生时间、操作步骤、预期结果和实际结果，尽量写成可重复的步骤。",
            "浏览器名称与版本、访问地址、摄像头/麦克风状态、是否通过代理。",
            "后端终端的完整错误附近若干行；不要只截“服务器处理失败”这一句。",
            "浏览器 Network 中失败请求的方法、URL、HTTP 状态码和响应 JSON。",
            "Python、Node.js 版本，以及是否刚修改过模型、提示词、阈值或依赖。",
            "如与报告有关，提供脱敏后的报告片段；不要直接公开真实 API Key 或人脸向量。",
        ),
    )
    add_code(
        document,
        r"""# 方便排障的基础信息
& "D:\Work_Software\Anaconda3\envs\nlp_env\python.exe" --version
node --version
npm --version
& "D:\Work_Software\Anaconda3\envs\nlp_env\python.exe" -m pip check""",
    )


def add_chapter_11_testing(document: Document) -> None:
    add_section_break(document)
    document.add_heading("第 11 章 测试、验收与二次开发", level=1)
    document.add_heading("11.1 自动化测试覆盖了什么", level=2)
    add_table(
        document,
        ("测试文件", "重点验证"),
        (
            ("test_api_health.py", "健康接口与应用基础能力"),
            ("test_auth_service.py", "用户注册、密码摘要、人脸绑定与匹配"),
            ("test_deepseek_service.py", "请求结构、错误处理和 JSON 解析"),
            ("test_interview_service.py", "提示词上下文、评估规范化和响应字段"),
            ("test_score_service.py", "评分公式、四舍五入和视频缺失边界"),
            ("test_report_store.py", "报告保存、列表、读取和用户隔离"),
            ("test_session_store.py", "会话创建、归属、状态和过期清理"),
            ("test_worker.py", "最新帧队列、替换和视觉统计"),
        ),
        widths=(5.3, 11.0),
    )
    add_body(
        document,
        "这些测试主要覆盖后端逻辑，不需要摄像头或真实 DeepSeek 调用。它们能快速发现"
        "接口结构、评分和存储回归，但不能替代真实浏览器、麦克风、摄像头与网络环境的"
        "端到端验收。",
    )

    document.add_heading("11.2 标准验证命令", level=2)
    add_code(
        document,
        r"""cd "D:\Project_Code_2026\tyut_comprehensive_pratice\smartspeech_optimized_new"

# 1. 后端单元测试
& "D:\Work_Software\Anaconda3\envs\nlp_env\python.exe" -m unittest discover -s backend\tests -v

# 2. Python 语法检查
& "D:\Work_Software\Anaconda3\envs\nlp_env\python.exe" -m compileall backend

# 3. Python 依赖一致性
& "D:\Work_Software\Anaconda3\envs\nlp_env\python.exe" -m pip check

# 4. 前端生产构建
cd frontend
npm run build

# 5. 健康检查（后端启动后）
Invoke-RestMethod "http://127.0.0.1:8765/api/health" | ConvertTo-Json -Depth 5""",
    )

    document.add_heading("11.3 手工验收清单", level=2)
    add_table(
        document,
        ("编号", "验收项", "通过标准"),
        (
            ("A1", "启动", "一键脚本能判断已运行或正常启动，浏览器打开 8765"),
            ("A2", "账号", "注册、密码登录、退出和重新登录正常"),
            ("A3", "人脸", "清晰人脸能够绑定；无脸照片给出明确提示"),
            ("A4", "接题", "开场和第一题出现；API 失败时显示备用问题警告"),
            ("A5", "文字回答", "回答进入对话，下一问与上下文相关"),
            ("A6", "语音回答", "能录音、转写、修正并提交，临时文件会清理"),
            ("A7", "TTS", "问题可自动或手动播放；失败时有回退或提示"),
            ("A8", "视觉", "有画面、人脸状态、表情、头姿、FPS 和延迟"),
            ("A9", "结束", "有回答才能结束；等待后能生成报告"),
            ("A10", "评分", "最终分与公式、单项分一致，视频不足时为空"),
            ("A11", "历史", "用户只能看到自己的报告，列表与详情一致"),
            ("A12", "导出", "浏览器能打印或另存 PDF，关键内容不缺失"),
            ("A13", "隐私", "报告声明与代码行为一致，不保存原始视频/照片/录音"),
            ("A14", "回归", "测试、build、pip check 均通过，无新增错误日志"),
        ),
        widths=(1.7, 4.3, 10.3),
    )

    document.add_heading("11.4 端到端验收建议", level=2)
    add_steps(
        document,
        (
            "重启电脑后不依赖开发热更新，直接运行 start_system.bat。",
            "新建测试账号并绑定人脸，完成 2-3 轮文字问答，确认上下文连续。",
            "另起账号进行人脸登录，确认不会匹配到测试账号。",
            "用语音完成至少一轮回答，并在报告中核对转写文字。",
            "保证摄像头覆盖率至少达到 10 个含人脸帧，确认视频分可用。",
            "完成面试并核对最终分公式、逐题数量、事件顺序和历史列表。",
            "导出 PDF，检查中文字体、图片、分数和分页；再清理测试账号与测试报告。",
        ),
    )
    add_note(
        document,
        "验收必须使用真实设备",
        "单元测试通过只能说明代码逻辑正常。摄像头权限、麦克风、浏览器自动播放、网络"
        "API、模型首次下载和报告打印都必须在最终演示电脑上实测一次。",
    )

    document.add_heading("11.5 主要 API 二次开发入口", level=2)
    add_table(
        document,
        ("方法与路径", "用途", "核心输入/输出"),
        (
            ("GET /api/health", "检查服务与能力", "状态和能力布尔值"),
            ("POST /api/auth/register", "注册", "用户名、显示名、密码、可选人脸图"),
            ("POST /api/auth/login", "密码登录", "用户名、密码"),
            ("POST /api/auth/face-register", "绑定人脸", "人脸图片"),
            ("POST /api/auth/face-login", "人脸登录", "人脸图片"),
            ("POST /api/interviews", "创建会话", "岗位、难度、音色、帧率"),
            ("POST /api/interviews/<id>/opening", "生成首题", "开场和 question 消息"),
            ("POST /api/interviews/<id>/frames", "提交分析帧", "JPEG、最新队列状态、指标"),
            ("GET /api/interviews/<id>/metrics", "读取实时指标", "analyzer metrics 与 worker status"),
            ("POST /api/interviews/<id>/messages", "回答并追问", "回答文字、下一条问题"),
            ("POST /api/interviews/<id>/finish", "结束并报告", "完整 report JSON"),
            ("DELETE /api/interviews/<id>", "放弃面试", "删除会话，不生成报告"),
            ("GET /api/reports", "报告列表", "当前用户报告摘要"),
            ("GET /api/reports/<id>", "报告详情", "当前用户指定报告"),
            ("DELETE /api/reports/<id>", "删除报告", "删除持久 JSON"),
            ("POST /api/media/transcribe", "语音转写", "音频 multipart 与 text/language"),
            ("POST /api/media/tts", "语音合成", "问题文本与音频响应"),
        ),
        widths=(6.5, 4.1, 5.7),
    )

    document.add_heading("11.6 常见二次开发方向", level=2)
    add_table(
        document,
        ("目标", "建议改动", "保持不动的部分"),
        (
            ("新增面试岗位", "前端选项与提示词岗位说明", "会话和消息结构"),
            ("接入其他大模型", "替换 DeepSeekService 的 chat/chat_json 实现", "InterviewService 输出约定"),
            ("增加评分维度", "评估 JSON、报告页面和融合公式", "历史报告兼容策略"),
            ("替换语音识别", "ASRService.transcribe 内部实现", "返回 text/language/segments"),
            ("替换表情模型", "ExpressionAnalyzer 的特征或分类逻辑", "表情状态与分数字段"),
            ("替换注视模型", "AttentionAnalyzer 的姿态/视线部分", "pitch/yaw/roll/gaze/score 字段"),
            ("增加管理后台", "新增 Flask Blueprint 与 Vue 页面", "用户数据隔离和权限检查"),
            ("支持远程部署", "HTTPS、数据库、对象存储、队列和审计", "隐私边界与报告语义"),
        ),
        widths=(3.7, 7.0, 5.6),
    )
    add_body(
        document,
        "二次开发应保持稳定契约：API 路径、报告字段和状态枚举一旦被页面和测试使用，"
        "就不要无说明地改名。需要升级数据结构时，应增加 version 字段并提供兼容读取；"
        "不要直接覆盖旧报告。",
    )

    document.add_heading("11.7 发布或答辩前的检查清单", level=2)
    add_bullets(
        document,
        (
            "在最终演示电脑上从冷启动运行一次，确认模型缓存和依赖都已准备好。",
            "准备一个无敏感信息的演示账号和一段可展示的固定回答，避免临场网络波动。",
            "确认 DeepSeek 余额、API Key 和网络可用，准备首题备用方案说明。",
            "在演示环境完成一次人脸绑定但不要展示真实人脸照片或报告隐私信息。",
            "后端单元测试、npm build 和 pip check 全部通过，记录版本与日期。",
            "确认 README、前端公式文案、评分配置和本 Word 文档一致。",
            "完成 Word 目录更新、页码检查、截图清晰度和 PDF 导出检查。",
            "把 .env、人脸向量和报告排除在公开代码包之外，只交付必要的脱敏样例。",
        ),
    )
    add_note(
        document,
        "答辩讲法",
        "用“采集 -> 本地理解 -> 云端语言决策 -> 本地融合 -> 可解释报告”讲完整链路。"
        "重点解释为什么视频本地处理、为什么问答占 80%、为什么视频样本不足时不出最终分，"
        "这些取舍比简单罗列模型名称更能体现工程设计。",
    )


def add_chapter_12_faq_glossary(document: Document) -> None:
    add_section_break(document)
    document.add_heading("第 12 章 常见问题、术语与结语", level=1)
    document.add_heading("12.1 学生常见问题", level=2)
    add_table(
        document,
        ("问题", "简明回答"),
        (
            ("系统必须联网吗？", "人脸、表情、注意力、报告和页面可本地运行；DeepSeek 与 Edge-TTS 需要联网，TTS 有 Windows 回退。"),
            ("必须训练模型吗？", "不需要。系统使用预训练模型和规则，开箱即可完成面试。"),
            ("视频会上传给 DeepSeek 吗？", "不会。只把岗位、难度和文字问答发送给 DeepSeek。"),
            ("原始照片保存在哪里？", "不保存原始照片，只保存 data/auth/faces 下的特征向量。"),
            ("回答录音会保留吗？", "不会。上传到 runtime/asr 的临时文件在转写后删除。"),
            ("表情分低说明状态不好吗？", "只能说明相对基线的几何状态偏中性，不能诊断情绪或人格。"),
            ("为什么注意力分和主观感受不同？", "普通单目摄像头只能估计头姿和视线趋势，受坐姿、镜头与光线影响。"),
            ("为什么视频分有时为空？", "有效人脸样本不足，系统拒绝用少数帧生成不可靠分数。"),
            ("为什么切换 Whisper 后第一次很慢？", "首次使用新规格需要下载模型，之后复用缓存。"),
            ("为什么自动语音没有播放？", "浏览器可能阻止自动播放，点击问题旁喇叭即可。"),
            ("为什么修改页面后 8765 没变化？", "生产页面来自 dist，需重新执行 npm run build。"),
            ("API Key 填在哪里？", "项目根目录 .env 的 DEEPSEEK_API_KEY，修改后重启后端。"),
            ("人脸登录能替代密码吗？", "可以用于本地演示，但不建议作为高安全等级的唯一个人认证。"),
            ("删除报告能恢复吗？", "不能。删除会清理会话和报告，重要文件应先备份。"),
            ("报告能发给别人吗？", "先脱敏。报告可能包含项目经历、回答文字和个人状态趋势。"),
            ("能在手机或局域网访问吗？", "技术上可配置，但摄像头通常要求 HTTPS；公开访问还需账号、安全与隐私改造。"),
        ),
        widths=(5.2, 11.1),
    )

    document.add_heading("12.2 术语表", level=2)
    add_table(
        document,
        ("术语", "白话解释"),
        (
            ("API", "程序之间约定好的调用入口，像只接收规定格式请求的窗口。"),
            ("前端", "浏览器里看到的页面与交互，本项目使用 Vue 3。"),
            ("后端", "负责业务、模型调度和文件存储的 Python 服务，本项目使用 Flask。"),
            ("Waitress", "在 Windows 上稳定运行 Flask 的生产级服务器。"),
            ("模型", "从数据中学到参数的程序，负责检测、识别或生成。"),
            ("预训练", "开发者已在大量通用数据上训练好，项目直接加载使用。"),
            ("微调", "在预训练模型基础上，用领域数据继续训练，使其更适合特定任务。"),
            ("推理", "模型加载后对一条新数据做预测，不是训练。"),
            ("特征向量", "用一串数字表示对象的关键特征，人脸当前为 512 维。"),
            ("L2 归一化", "把向量长度统一为 1，使距离比较更稳定。"),
            ("欧氏距离", "两个向量之间的直线距离，越小通常越相似。"),
            ("阈值", "判断分类的边界值，例如距离小于 1.0 视为人脸匹配。"),
            ("Face Mesh", "在脸上建立数百个关键点的网格。"),
            ("EAR", "眼睛纵横比，用眼睛上下距离与左右距离的比值估计开合。"),
            ("pitch / yaw / roll", "点头俯仰、左右转头、头部侧倾三个姿态角。"),
            ("solvePnP", "根据二维关键点和三维标准脸估计旋转与位置的视觉算法。"),
            ("FPS", "每秒完成分析的帧数；有效 FPS 是后端实际处理速度。"),
            ("延迟", "一帧从接收到处理完成的耗时，常以毫秒表示。"),
            ("丢帧", "新帧替换旧待处理帧，用最新状态换取实时性。"),
            ("上下文", "发送给大模型的最近对话消息，用于延续话题。"),
            ("温度", "控制大模型输出随机性；越高越发散，越低越稳定。"),
            ("JSON", "结构化文本格式，适合保存报告和让程序稳定读取字段。"),
            ("WER / CER", "词错误率/字错误率，评价语音转写准确程度。"),
            ("宏平均 F1", "先计算每类指标再平均，避免大类多就把小类错误掩盖。"),
            ("TAR@FAR", "在给定误接受率下的人脸正确接受率，常与阈值一起评估。"),
        ),
        widths=(4.0, 12.3),
    )

    document.add_heading("12.3 最后一个完整检查清单", level=2)
    add_bullets(
        document,
        (
            "我已经确认当前系统正常使用不需要训练模型。",
            "我知道页面端口是 8765，开发端口是 5173，启动入口是 start_system.bat。",
            "我知道 DeepSeek 只接收文字，人脸、视频和录音主要在本机处理。",
            "我能按“注册 -> 配置 -> 开始 -> 回答 -> 结束 -> 报告”完成一次面试。",
            "我知道最终分是问答 80% 加视频综合 20%，视频不足时最终分为空。",
            "我能区分配置调优、提示词调优、阈值调优和真正训练模型。",
            "我知道换人脸模型后必须重建向量，换评分权重后必须回归历史报告。",
            "我能运行 backend 单元测试、compileall、pip check 和 npm run build。",
            "我知道上报故障时应提供可复现步骤、日志、Network 状态和环境版本。",
            "我不会公开 API Key、脸部向量、真实回答、报告和未脱敏截图。",
        ),
    )

    document.add_heading("12.4 结语", level=2)
    add_body(
        document,
        "智能体模拟面试系统的重点不只是一个得分，而是把“面试问题、回答内容、语音链路、"
        "实时状态和事后复盘”放进同一套可解释流程。学生使用时先按默认配置完整跑通，再"
        "根据岗位、设备和训练目标逐项调整；开发或答辩时则沿着第 6 章调用链解释每个模块"
        "为什么存在、数据在哪里流动、边界在哪里。",
    )
    add_note(
        document,
        "使用提醒",
        "模拟评分仅用于训练参考，不等同于真实招聘结论，也不应作为心理、人格或健康评价。"
        "请结合逐题建议反复练习，持续提升回答结构、专业深度和自然稳定的表达状态。",
    )


def build_manual() -> Path:
    document = Document()
    configure_document(document)
    add_cover(document)
    add_reading_guide(document)

    add_chapter_1_project(document)
    add_chapter_2_features(document)
    add_chapter_3_architecture(document)
    add_chapter_4_principles(document)
    add_chapter_5_models_and_data(document)
    add_chapter_6_code_structure(document)
    add_chapter_7_install(document)
    add_chapter_8_operation(document)
    add_chapter_9_training(document)
    add_chapter_10_troubleshooting(document)
    add_chapter_11_testing(document)
    add_chapter_12_faq_glossary(document)

    document.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    path = build_manual()
    print(f"Document generated: {path}")
